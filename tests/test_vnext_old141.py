"""Old141 execution isolates copied legacy assets and never treats heldout as gold."""
from __future__ import annotations

import hashlib
import importlib.util
import json
import os
import shutil
import sys
from copy import deepcopy
from pathlib import Path
from types import ModuleType

import pytest
import yaml
from openpyxl import Workbook, load_workbook
from qdrant_client import QdrantClient, models

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
OLD_FILE = "基地云机房信息调研表.xlsx"
DIMENSION = 4096
POINTS = 2
SAFE_ITEM_KEYS = {
    "form_item_id", "file_name", "sheet_name", "row_index", "target_cell", "category_path",
    "question_text", "instruction_text", "needs_evidence", "answer_example",
}


def load_script(name: str) -> ModuleType:
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def read_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text().splitlines() if line.strip()]


def file_inventory(root: Path) -> dict[str, tuple[int, str]]:
    return {
        str(path.relative_to(root)): (path.stat().st_size, hashlib.sha256(path.read_bytes()).hexdigest())
        for path in root.rglob("*") if path.is_file()
    }


@pytest.fixture
def old141_module(monkeypatch: pytest.MonkeyPatch) -> ModuleType:
    module = load_script("vnext_old141")
    monkeypatch.setattr(module, "EXPECTED_POINTS", POINTS)
    return module


@pytest.fixture
def legacy_inputs(tmp_path: Path) -> dict[str, Path]:
    original = tmp_path / "original"
    template = original / "data/工勘单" / OLD_FILE
    template.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.title = "原表"
    workbook.active["A1"] = "保留布局"
    workbook.active["H160"] = "=SUM(1,2)"
    items = []
    for row in range(4, 145):
        secret_fact = f"仅评测可见人工事实{row}"
        workbook.active[f"G{row}"] = secret_fact
        items.append({
            "form_item_id": f"old-{row}", "file_name": OLD_FILE, "sheet_name": "原表",
            "row_index": row, "target_cell": f"G{row}", "question_text": f"字段{row}的现网值",
            "instruction_text": "填写现场实际值", "category_path": ["供电"], "needs_evidence": True,
            "existing_value": secret_fact, "current_info": secret_fact, "answer_example": secret_fact,
            "heldout_answer": secret_fact, "suggested_retrieval_query": secret_fact,
        })
    cover = workbook.create_sheet("封面")
    cover["G4"] = "封面内容必须保留"
    workbook.active = workbook.index(cover)
    workbook.save(template)
    workbook.close()
    form_items = original / "artifacts/12_gongkan_form_analysis/form_items.jsonl"
    form_items.parent.mkdir(parents=True)
    form_items.write_text("".join(json.dumps(item, ensure_ascii=False) + "\n" for item in items))
    source = original / "data/知识库/匿名能力.xlsx"
    source.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.title = "能力"
    workbook.active.append(["字段", "现网值"])
    workbook.active.append(["UPS容量", "500kVA"])
    workbook.save(source)
    workbook.close()
    store = original / "artifacts/15_vector_store/qdrant"
    client = QdrantClient(path=str(store))
    try:
        client.create_collection("datacenter_chunks_v1", vectors_config=models.VectorParams(size=DIMENSION, distance=models.Distance.COSINE))
        client.upsert("datacenter_chunks_v1", points=[
            models.PointStruct(id=index + 1, vector=[1.0, *([0.0] * (DIMENSION - 1))], payload={
                "chunk_id": f"legacy-{index}", "namespace": "base_cloud", "file_name": source.name,
                "raw_source_text": "UPS容量：500kVA", "source_type": "main_excel_capability",
                "evidence_kind": "structured_field", "sheet_name": "能力", "row_index": 2,
            }) for index in range(POINTS)
        ])
    finally:
        client.close()
    prepared = tmp_path / "prepared"
    load_script("vnext_prepare_old141").prepare(original, prepared)
    models_config = tmp_path / "models.yaml"
    models_config.write_text(yaml.safe_dump({
        "services": {
            "embedding_endpoint": "http://unused.invalid/v1/embeddings", "embedding_model": "fixture-4096",
            "rerank_endpoint": "http://unused.invalid/rerank", "rerank_model": "fixture-rerank",
            "chat_endpoint": "http://unused.invalid/v1/chat/completions", "chat_model": "fixture-chat",
            "chat_api_key_env": "OLD141_FIXTURE_CREDENTIAL", "timeout_seconds": 2,
        },
    }))
    return {"original": original, "prepared": prepared, "models_config": models_config,
            "template": template, "source": source, "store": store, "out": tmp_path / "new-run"}


def prepare_run(module: ModuleType, inputs: dict[str, Path], **kwargs) -> dict:
    options = {"models_kind": "stub", "models_note": "Explicit fixture stub; no real quality claim",
               "embedding_dimension": DIMENSION, "env": {}}
    options.update(kwargs)
    return module.prepare(
        inputs["original"], inputs["prepared"], inputs["models_config"], inputs["out"],
        **options,
    )


def test_environment_sanitizer_removes_host_overrides_and_preserves_explicit_credential(old141_module: ModuleType) -> None:
    env = {
        "PATH": "/fixture/bin", "LANG": "zh_CN.UTF-8", "OLD141_FIXTURE_CREDENTIAL": "fake-test-value",
        "EMBEDDING_ENDPOINT": "http://must-not-be-used.invalid", "CHAT_MODEL": "wrong-model",
        "QDRANT_PATH": "/wrong-store", "QDRANT_URL": "http://must-not-be-used.invalid",
        "TARGET_NAMESPACE": "wrong-room", "WRITEBACK_ALLOW_UNCERTAIN": "true",
        "NDR_WRITEBACK_ALLOW_UNCERTAIN": "true", "NDR_UNRECOGNIZED_FUTURE_OPTION": "true",
        "NESTED_DOC_RAG__WRITEBACK__EXISTING_VALUE_POLICY": "overwrite_all",
        "NESTED_DOC_RAG__PATHS__PROJECT_ROOT": "/wrong-root",
    }
    before = deepcopy(env)
    sanitized = old141_module.sanitize_environment(env, "OLD141_FIXTURE_CREDENTIAL")
    assert sanitized == {
        "PATH": "/fixture/bin", "LANG": "zh_CN.UTF-8", "OLD141_FIXTURE_CREDENTIAL": "fake-test-value",
        "PYTHONPATH": str(old141_module.ROOT / "src"), "PYTHONNOUSERSITE": "1",
    }
    assert env == before


@pytest.mark.parametrize("note", [None, "", "   "])
def test_stub_requires_explicit_nonempty_note_before_preparation(
    old141_module: ModuleType, legacy_inputs: dict[str, Path], note: str | None,
) -> None:
    original_before = file_inventory(legacy_inputs["original"])
    kwargs = {"models_note": note} if note is not None else {}
    with pytest.raises(ValueError, match="stub.*explicit.*models-note"):
        old141_module.prepare(
            legacy_inputs["original"], legacy_inputs["prepared"], legacy_inputs["models_config"], legacy_inputs["out"],
            models_kind="stub", embedding_dimension=DIMENSION, env={}, **kwargs,
        )
    assert not legacy_inputs["out"].exists()
    assert file_inventory(legacy_inputs["original"]) == original_before


def test_models_dotenv_is_local_mapping_and_cannot_override_frozen_configuration(
    old141_module: ModuleType, legacy_inputs: dict[str, Path], tmp_path: Path,
) -> None:
    models_env = tmp_path / "private-models.env"
    models_env.write_text(
        "# Explicit fixture credentials, never real secrets.\n"
        "export OLD141_FIXTURE_CREDENTIAL='dotenv-fixture-credential'\n"
        "NDR_CHAT_MODEL=dotenv-wrong-model\n"
        "CHAT_MODEL=dotenv-alias-model\n"
        "EMBEDDING_MODEL=dotenv-64-dimension-model\n"
        "QDRANT_PATH=/must-not-be-used\n"
        "NESTED_DOC_RAG__WRITEBACK__EXISTING_VALUE_POLICY=overwrite_all\n"
    )
    host = {"PATH": "/fixture/bin", "LANG": "zh_CN.UTF-8", "CHAT_MODEL": "host-wrong-model"}
    host_before = deepcopy(host)
    process_before = dict(os.environ)
    local = old141_module.read_models_environment(models_env, host)
    assert local is not host and host == host_before and dict(os.environ) == process_before
    assert local["OLD141_FIXTURE_CREDENTIAL"] == "dotenv-fixture-credential"
    clean = old141_module.sanitize_environment(local, "OLD141_FIXTURE_CREDENTIAL")
    assert clean == {
        "PATH": "/fixture/bin", "LANG": "zh_CN.UTF-8", "OLD141_FIXTURE_CREDENTIAL": "dotenv-fixture-credential",
        "PYTHONPATH": str(old141_module.ROOT / "src"), "PYTHONNOUSERSITE": "1",
    }
    ledger = prepare_run(old141_module, legacy_inputs, env=local)
    assert ledger["models_note"] == "Explicit fixture stub; no real quality claim"
    assert all(group["effective_config"]["services"]["chat_model"] == "fixture-chat" for group in ledger["groups"])
    assert all(group["effective_config"]["services"]["embedding_model"] == "fixture-4096" for group in ledger["groups"])
    assert all(group["effective_config"]["writeback"]["existing_value_policy"] == "preserve" for group in ledger["groups"])
    assert "dotenv-fixture-credential" not in json.dumps(ledger)
    assert all("dotenv-fixture-credential" not in path.read_text() for path in legacy_inputs["out"].rglob("*") if path.suffix in {".json", ".yaml"})
    assert host == host_before and dict(os.environ) == process_before


def test_models_env_cli_passes_local_mapping_to_prepare_and_execute_without_host_mutation(
    old141_module: ModuleType, legacy_inputs: dict[str, Path], tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str],
) -> None:
    models_env = tmp_path / "private-cli.env"
    models_env.write_text("OLD141_FIXTURE_CREDENTIAL=cli-fixture-credential\nNDR_CHAT_MODEL=wrong-model\n")
    process_before = dict(os.environ)
    calls = []
    real_prepare = old141_module.prepare

    def tracked_prepare(*args, **kwargs):
        calls.append(("prepare", kwargs["env"]))
        return real_prepare(*args, **kwargs)

    def isolated_execute(ledger, *, env):
        calls.append(("execute", env))
        assert "cli-fixture-credential" not in json.dumps(ledger)
        ledger["status"] = "completed"
        return ledger

    monkeypatch.setattr(old141_module, "prepare", tracked_prepare)
    monkeypatch.setattr(old141_module, "execute", isolated_execute)
    monkeypatch.setattr(sys, "argv", [
        "vnext_old141", "--original-root", str(legacy_inputs["original"]), "--prepared-dir", str(legacy_inputs["prepared"]),
        "--models-config", str(legacy_inputs["models_config"]), "--models-env", str(models_env),
        "--models-kind", "stub", "--models-note", "Explicit CLI fixture stub; no real quality claim",
        "--embedding-dimension", str(DIMENSION), "--out-dir", str(legacy_inputs["out"]), "--execute",
    ])
    assert old141_module.main() == 0
    assert [name for name, _ in calls] == ["prepare", "execute"]
    assert calls[0][1] is calls[1][1] and calls[0][1] is not os.environ
    assert calls[0][1]["OLD141_FIXTURE_CREDENTIAL"] == "cli-fixture-credential"
    assert dict(os.environ) == process_before
    assert "cli-fixture-credential" not in capsys.readouterr().out
    ledger = json.loads((legacy_inputs["out"] / "ledger.json").read_text())
    assert ledger["models_env_source"]["values_archived"] is False
    assert "cli-fixture-credential" not in json.dumps(ledger)


def test_real_copied_qdrant_reports_4096_dimension_and_tiny_actual_point_count(
    old141_module: ModuleType, legacy_inputs: dict[str, Path], tmp_path: Path,
) -> None:
    original_before = file_inventory(legacy_inputs["original"])
    copied = tmp_path / "inspection-copy"
    shutil.copytree(legacy_inputs["store"], copied)
    info = old141_module.inspect_copied_store(copied)
    assert info["collection_name"] == "datacenter_chunks_v1"
    assert info["dimensions"] == DIMENSION
    assert info["points_count"] == POINTS
    assert file_inventory(legacy_inputs["original"]) == original_before


def test_prepare_copies_original_and_blank_templates_with_closed_book_141_inputs(
    old141_module: ModuleType, legacy_inputs: dict[str, Path], monkeypatch: pytest.MonkeyPatch,
) -> None:
    original_before = file_inventory(legacy_inputs["original"])
    prepared_before = file_inventory(legacy_inputs["prepared"])
    opened_stores = []

    def tracked_qdrant(**kwargs):
        path = Path(kwargs["path"])
        assert path.resolve().is_relative_to(legacy_inputs["out"].resolve())
        opened_stores.append(path)
        return QdrantClient(**kwargs)

    def forbidden_execution(*args, **kwargs):
        raise AssertionError("preparation must not execute runners or models")

    monkeypatch.setattr(old141_module, "QdrantClient", tracked_qdrant)
    monkeypatch.setattr(old141_module, "run_command", forbidden_execution)
    ledger = prepare_run(old141_module, legacy_inputs, env={
        "OLD141_FIXTURE_CREDENTIAL": "fake-test-value", "EMBEDDING_MODEL": "wrong-model",
        "QDRANT_PATH": str(legacy_inputs["original"] / "forbidden-store"),
        "NESTED_DOC_RAG__WRITEBACK__EXISTING_VALUE_POLICY": "overwrite_all",
    })
    assert ledger["status"] == "prepared" and ledger["models_kind"] == "stub"
    assert ledger["model_requests"] == 0
    assert ledger["original_input_unchanged"] is True
    assert ledger["source_originals_before"] == ledger["source_originals_after"]
    assert file_inventory(legacy_inputs["original"]) == original_before
    assert file_inventory(legacy_inputs["prepared"]) == prepared_before
    groups = {group["name"]: group for group in ledger["groups"]}
    assert set(groups) == {"blank", "preserve"}
    assert opened_stores == [Path(group["qdrant_path"]) for group in ledger["groups"]]
    heldout_facts = {row["expected_value"] for row in read_jsonl(legacy_inputs["prepared"] / "heldout_answers.jsonl")}
    for name, group in groups.items():
        for key in ("template", "form_items", "qdrant_path", "config_path", "run_dir"):
            assert Path(group[key]).resolve().is_relative_to(legacy_inputs["out"].resolve())
        items = read_jsonl(Path(group["form_items"]))
        assert len(items) == 141 and {row["row_index"] for row in items} == set(range(4, 145))
        assert len({row["form_item_id"] for row in items}) == 141
        assert all(set(row) <= SAFE_ITEM_KEYS and row["answer_example"] == "" for row in items)
        assert all(row["target_cell"] == f"'原表'!G{row['row_index']}" for row in items)
        assert not any(fact in Path(group["form_items"]).read_text() for fact in heldout_facts)
        assert group["index_copy_verified"] is True
        assert group["index_info"]["dimensions"] == DIMENSION
        assert group["index_info"]["points_count"] == POINTS
        assert file_inventory(Path(group["qdrant_path"])) == file_inventory(legacy_inputs["store"])
        config = group["effective_config"]
        assert config["writeback"]["existing_value_policy"] == "preserve"
        assert config["retrieval"]["sufficiency_enabled"] is True
        assert config["retrieval"]["schema_first_enabled"] is False
        assert config["services"]["embedding_model"] == "fixture-4096"
        assert config["qdrant"]["url"] == ""
        assert all(Path(path).resolve().is_relative_to(legacy_inputs["out"].resolve()) for path in config["paths"].values())
        command = group["command"]
        assert "run-step15-agent" in command and "--no-judge" in command and "--writeback" in command
        assert command[command.index("--form-items") + 1] == group["form_items"]
        assert command[command.index("--template") + 1] == group["template"]
        assert "heldout_answers.jsonl" not in " ".join(command)
        workbook = load_workbook(group["template"], data_only=False)
        try:
            assert workbook.active.title == "封面" and workbook["封面"]["G4"].value == "封面内容必须保留"
            assert workbook["原表"]["A1"].value == "保留布局"
            assert workbook["原表"]["H160"].value == "=SUM(1,2)"
            for row in range(4, 145):
                assert workbook["原表"][f"G{row}"].value == (None if name == "blank" else f"仅评测可见人工事实{row}")
        finally:
            workbook.close()
    assert Path(groups["preserve"]["template"]).read_bytes() == legacy_inputs["template"].read_bytes()
    assert Path(groups["blank"]["template"]).read_bytes() == (legacy_inputs["prepared"] / "blank_target_template.xlsx").read_bytes()
    # Credentials may be forwarded to execution, but are never persisted.
    assert "fake-test-value" not in json.dumps(ledger)
    assert all("fake-test-value" not in path.read_text() for path in legacy_inputs["out"].rglob("*") if path.suffix in {".json", ".yaml"})


@pytest.mark.parametrize("bad_dimension", [64, 4095])
def test_old4096_index_rejects_incompatible_model_dimension_before_execution(
    old141_module: ModuleType, legacy_inputs: dict[str, Path], bad_dimension: int,
) -> None:
    original_before = file_inventory(legacy_inputs["original"])
    with pytest.raises(ValueError, match="dimension|4096"):
        prepare_run(old141_module, legacy_inputs, embedding_dimension=bad_dimension)
    assert file_inventory(legacy_inputs["original"]) == original_before


def test_bootstrap_calls_actual_cli_with_empty_config_environment(
    old141_module: ModuleType, legacy_inputs: dict[str, Path], monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    import nested_doc_rag.config as config_module
    from nested_doc_rag import cli

    real_loader = config_module.load_app_config
    observed_envs = []

    def tracked_loader(path=None, **kwargs):
        observed_envs.append(kwargs.get("env"))
        return real_loader(path, **kwargs)

    def forbidden_dotenv(*args, **kwargs):
        raise AssertionError("frozen CLI configuration must not load host dotenv overrides")

    monkeypatch.setenv("EMBEDDING_MODEL", "host-override")
    monkeypatch.setenv("NESTED_DOC_RAG__WRITEBACK__EXISTING_VALUE_POLICY", "overwrite_all")
    monkeypatch.setattr(config_module, "load_app_config", tracked_loader)
    monkeypatch.setattr(config_module, "load_dotenv_file", forbidden_dotenv)
    # Register restoration for the bootstrap's mutation of the live CLI module.
    monkeypatch.setattr(cli, "load_app_config", cli.load_app_config)
    monkeypatch.setattr(sys, "argv", ["old141-bootstrap", "show-config", "--config", str(legacy_inputs["models_config"])])
    exec(old141_module.BOOTSTRAP, {"__name__": "__main__"})
    config = json.loads(capsys.readouterr().out)
    assert observed_envs == [{}]
    assert config["services"]["embedding_model"] == "fixture-4096"
    assert config["writeback"]["existing_value_policy"] == "preserve"


@pytest.mark.parametrize("change_preserved_target", [False, True])
def test_execute_labels_postrun_agreement_unverified_and_checks_all_preserved_targets(
    old141_module: ModuleType, legacy_inputs: dict[str, Path], monkeypatch: pytest.MonkeyPatch,
    change_preserved_target: bool,
) -> None:
    ledger = prepare_run(old141_module, legacy_inputs)
    original_before = file_inventory(legacy_inputs["original"])
    events = []
    real_read_rows = old141_module.read_rows

    def tracked_read_rows(path):
        if Path(path).name == "heldout_answers.jsonl":
            events.append("heldout_comparison")
        return real_read_rows(path)

    def simulated_runner(command, cwd, env, log):
        run_dir = Path(command[command.index("--out-dir") + 1])
        name = run_dir.parent.name
        events.append(f"execute:{name}")
        assert cwd == legacy_inputs["out"].resolve()
        assert "NDR_CHAT_MODEL" not in env and "QDRANT_PATH" not in env
        assert env["OLD141_FIXTURE_CREDENTIAL"] == "fake-test-value"
        items_path = Path(command[command.index("--form-items") + 1])
        items = read_jsonl(items_path)
        assert all(set(item) <= SAFE_ITEM_KEYS and item["answer_example"] == "" for item in items)
        assert "仅评测可见人工事实" not in items_path.read_text()
        assert "heldout_answers.jsonl" not in " ".join(command)
        run_dir.mkdir()
        shutil.copyfile(command[command.index("--template") + 1], run_dir / "filled_form.xlsx")
        if name == "preserve" and change_preserved_target:
            workbook = load_workbook(run_dir / "filled_form.xlsx")
            workbook["原表"]["G4"] = "未经许可的覆盖"
            workbook.save(run_dir / "filled_form.xlsx")
            workbook.close()
        predictions = [
            {"field_id": item["form_item_id"], "answer_value": "仅评测可见人工事实4" if item["row_index"] == 4 else "未找到"}
            for item in items
        ]
        (run_dir / "predictions_raw.jsonl").write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in predictions))
        (run_dir / "run_manifest.json").write_text(json.dumps({"counts": {"failed": 0}}))
        log.write_text("unit orchestration substitute; no model calls")
        return {"command": command, "exit_code": 0, "elapsed_seconds": 0, "log": str(log), "log_sha256": hashlib.sha256(log.read_bytes()).hexdigest()}

    def validated_artifacts(path):
        events.append(f"validated:{Path(path).parent.name}")
        return {"valid": True, "evidence_validation": "strict"}

    monkeypatch.setattr(old141_module, "run_command", simulated_runner)
    monkeypatch.setattr(old141_module, "validate_step15_artifacts", validated_artifacts)
    monkeypatch.setattr(old141_module, "read_rows", tracked_read_rows)
    result = old141_module.execute(ledger, {
        "OLD141_FIXTURE_CREDENTIAL": "fake-test-value", "NDR_CHAT_MODEL": "wrong-model",
        "QDRANT_PATH": str(legacy_inputs["original"] / "forbidden-store"),
    })
    assert events == ["execute:blank", "validated:blank", "heldout_comparison", "execute:preserve", "validated:preserve", "heldout_comparison"]
    assert result["status"] == ("failed" if change_preserved_target else "completed")
    assert result["quality_metrics"] is None and result["original_input_unchanged"] is True
    assert result["source_originals_before"] == result["source_originals_after"]
    assert file_inventory(legacy_inputs["original"]) == original_before
    for group in result["groups"]:
        comparator = group["heldout_comparison"]
        assert comparator["label"] == "unverified_heldout_agreement"
        assert comparator["gold_verified"] is False and comparator["quality_metrics"] is None
        assert comparator["compared_count"] == 141 and comparator["agreement_count"] == 1
        assert "accuracy" not in comparator
    preserve = next(group for group in result["groups"] if group["name"] == "preserve")
    assert preserve["preserve_141_targets"] == {
        "checked_count": 141, "changed_count": int(change_preserved_target), "valid": not change_preserved_target,
    }
    if change_preserved_target:
        assert "preserve changed" in preserve["validation_error"]


@pytest.mark.parametrize("damage", ["command", "index_bytes", "index_symlink", "template", "form_items", "config_path"])
def test_execute_rejects_changed_frozen_inputs_before_any_runner_call(
    old141_module: ModuleType, legacy_inputs: dict[str, Path], monkeypatch: pytest.MonkeyPatch, damage: str,
) -> None:
    ledger = prepare_run(old141_module, legacy_inputs)
    original_before = file_inventory(legacy_inputs["original"])
    group = ledger["groups"][0]
    if damage == "command":
        group["command"].append("--judge")
    elif damage == "index_symlink":
        (Path(group["qdrant_path"]) / "unsafe-source-link").symlink_to(legacy_inputs["source"])
    else:
        path = next(path for path in Path(group["qdrant_path"]).rglob("*") if path.is_file()) if damage == "index_bytes" else Path(group[damage])
        path.write_bytes(path.read_bytes() + b"tampered-after-freeze")

    def forbidden_runner(*args, **kwargs):
        raise AssertionError("changed frozen inputs must be rejected before runner/model execution")

    monkeypatch.setattr(old141_module, "run_command", forbidden_runner)
    with pytest.raises(ValueError, match="frozen|symlink"):
        old141_module.execute(ledger, {})
    assert ledger["status"] == "failed" and ledger["original_input_unchanged"] is True
    assert all(group["execution"] is None for group in ledger["groups"])
    assert file_inventory(legacy_inputs["original"]) == original_before
    assert json.loads((legacy_inputs["out"] / "ledger.json").read_text())["status"] == "failed"


@pytest.mark.parametrize("location", ["original", "prepared", "existing_empty", "existing_nonempty"])
def test_output_requires_new_directory_outside_original_and_preparation(
    old141_module: ModuleType, legacy_inputs: dict[str, Path], location: str,
) -> None:
    original_before = file_inventory(legacy_inputs["original"])
    prepared_before = file_inventory(legacy_inputs["prepared"])
    if location in {"original", "prepared"}:
        legacy_inputs["out"] = legacy_inputs[location] / "forbidden-run"
    else:
        legacy_inputs["out"].mkdir()
        if location == "existing_nonempty":
            (legacy_inputs["out"] / "existing.txt").write_text("retain existing run")
    with pytest.raises(ValueError, match="outside|new|exist|output"):
        prepare_run(old141_module, legacy_inputs)
    assert file_inventory(legacy_inputs["original"]) == original_before
    assert file_inventory(legacy_inputs["prepared"]) == prepared_before


@pytest.mark.parametrize("damage", ["missing_field", "duplicate_row", "heldout_key", "fact_example"])
def test_prepared_form_items_reject_count_identity_and_closed_book_leakage(
    old141_module: ModuleType, legacy_inputs: dict[str, Path], damage: str,
) -> None:
    path = legacy_inputs["prepared"] / "form_items_closed_book.jsonl"
    items = read_jsonl(path)
    if damage == "missing_field":
        items.pop()
    elif damage == "duplicate_row":
        items[-1] = deepcopy(items[0])
    elif damage == "heldout_key":
        items[0]["existing_value"] = "仅评测可见人工事实4"
    else:
        items[0]["answer_example"] = "仅评测可见人工事实4"
    path.write_text("".join(json.dumps(item, ensure_ascii=False) + "\n" for item in items))
    # Preserve the generated-file integrity declaration so rejection must come
    # from the strict 141/closed-book contract, rather than a generic checksum.
    preparation_path = legacy_inputs["prepared"] / "preparation.json"
    preparation = json.loads(preparation_path.read_text())
    preparation["generated_hashes"][path.name] = hashlib.sha256(path.read_bytes()).hexdigest()
    preparation_path.write_text(json.dumps(preparation))
    original_before = file_inventory(legacy_inputs["original"])
    with pytest.raises(ValueError, match="141|row|closed|whitelist|example|forbidden"):
        prepare_run(old141_module, legacy_inputs)
    assert file_inventory(legacy_inputs["original"]) == original_before
