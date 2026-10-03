"""The Excel writer independently proves references against authoritative hits."""
from __future__ import annotations

import base64
from copy import deepcopy
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from nested_doc_rag.artifacts import ArtifactValidationError, validate_step15_artifacts
from nested_doc_rag.evidence_record import normalize_evidence_record
from nested_doc_rag.evidence_resolver import resolve_evidence_refs
from nested_doc_rag.excel.writeback import patch_workbook, writeback_from_files
from nested_doc_rag.io import read_json, read_jsonl, write_json, write_jsonl
from nested_doc_rag.schemas.eval import FieldPrediction
from nested_doc_rag.schemas.evidence import EvidenceRef

PNG = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="


def source_hit(**changes) -> dict:
    return normalize_evidence_record({
        "chunk_id": "source_ups", "knowledge_base_id": "kb_power", "namespace": "room_301",
        "file_name": "匿名能力.xlsx", "relative_path": "材料/匿名能力.xlsx", "source_type": "uploaded_excel_row",
        "evidence_kind": "structured_field", "corpus_layer": "fact", "sheet_name": "能力", "row_index": 2,
        "cell_range": "B2:C2", "raw_source_text": "UPS容量：500kVA。\n人工核对😀", "raw_text": "UPS容量：500kVA。",
        **changes,
    })


def prediction_for(hit: dict) -> FieldPrediction:
    resolution = resolve_evidence_refs([hit["chunk_id"]], [hit])
    assert resolution.resolvable, resolution.errors
    return FieldPrediction(
        field_id="ups_capacity", row_index=2, target_cell="'当前表单'!D2", answer_value="500kVA",
        answer_status="answered", confidence=0.9, source_chunk_ids=[hit["chunk_id"]], evidence_refs=resolution.refs,
    )


def form_template(path: Path) -> Path:
    workbook = Workbook()
    workbook.active.title = "当前表单"
    workbook.active["B2"] = "UPS容量"
    workbook.save(path)
    workbook.close()
    return path


def write_form(tmp_path: Path, prediction: FieldPrediction, hits: list[dict], **kwargs):
    template = form_template(tmp_path / "template.xlsx")
    output = tmp_path / "filled_form.xlsx"
    summary = patch_workbook(
        template, [prediction], output, retrieval_hits_by_field_id={prediction.field_id: hits},
        overlays_by_field_id={prediction.field_id: {"writeback_allowed": True, "critic_flags": [], "review_required": False}},
        **kwargs,
    )
    return output, summary


def test_confirmed_native_ref_creates_addressable_evidence_sheet_and_location_comment(tmp_path: Path) -> None:
    hit = source_hit()
    output, summary = write_form(tmp_path, prediction_for(hit), [hit])
    workbook = load_workbook(output)
    assert workbook["当前表单"]["D2"].value == "500kVA"
    assert summary.written_count == summary.confirmed_count == 1
    sheet = workbook["Evidence"]
    assert [cell.value for cell in sheet[1]] == ["Field", "Answer", "Status", "Source", "Location", "Evidence"]
    assert [cell.value for cell in sheet[2]] == ["ups_capacity", "500kVA", "confirmed", "匿名能力.xlsx", "能力!B2:C2", hit["raw_source_text"]]
    comment = workbook["当前表单"]["D2"].comment.text
    assert "Evidence!A2" in comment and "能力!B2:C2" in comment
    assert hit["raw_source_text"] not in comment
    audit = read_jsonl(tmp_path / "writeback_audit.jsonl")[0]
    assert audit["comment_length"] == len(comment)
    assert audit["evidence_refs"] == [prediction_for(hit).evidence_refs[0].to_dict()]
    assert not workbook["当前表单"]._images
    workbook.close()


def test_chunk_ids_and_forged_overlay_cannot_authorize_confirmed_writeback(tmp_path: Path) -> None:
    prediction = FieldPrediction(
        field_id="ups_capacity", row_index=2, target_cell="'当前表单'!D2", answer_value="500kVA",
        answer_status="answered", confidence=1, source_chunk_ids=["source_ups"],
    )
    output, summary = write_form(tmp_path, prediction, [source_hit()])
    workbook = load_workbook(output)
    assert workbook["当前表单"]["D2"].value is None
    assert summary.written_count == summary.confirmed_count == 0
    audit = read_jsonl(tmp_path / "writeback_audit.jsonl")[0]
    assert audit["status"] == "flagged" and audit["error_code"] == "WB_MISSING_EVIDENCE"
    assert audit["evidence_refs"] == [], "A source ID alone must not manufacture an EvidenceRef"
    assert read_json(tmp_path / "evidence_map.json")["fields"]["ups_capacity"]["status"] == "flagged"
    workbook.close()


@pytest.mark.parametrize(
    ("changes", "code"),
    [
        ({"chunk_id": "forged"}, "EV_REF_NOT_IN_RETRIEVAL"),
        ({"knowledge_base_id": "another_kb"}, "EV_FILE_NOT_IN_KB"),
        ({"namespace": "another_room"}, "EV_FILE_NOT_IN_KB"),
        ({"file_name": "伪造文档.xlsx"}, "EV_FILE_NOT_IN_KB"),
        ({"relative_path": "伪造路径.xlsx"}, "EV_FILE_NOT_IN_KB"),
        ({"cell_range": "B8:C8"}, "EV_REF_NOT_IN_RETRIEVAL"),
        ({"cell_range": "ZZZ0"}, "EV_CELL_RANGE_INVALID"),
        ({"sheet_name": "另一张表"}, "EV_REF_NOT_IN_RETRIEVAL"),
        ({"source_text": "模型生成的错误原文"}, "EV_REF_NOT_IN_RETRIEVAL"),
        ({"source_text": ""}, "WB_MISSING_EVIDENCE"),
        ({"attachment_ids": ["absent_attachment"]}, "EV_ATTACHMENT_NOT_FOUND"),
        ({"verified": True}, "EV_REF_NOT_IN_RETRIEVAL"),
    ],
)
def test_final_writer_guard_rejects_tampered_ref_despite_allowed_overlay(tmp_path: Path, changes: dict, code: str) -> None:
    hit = source_hit()
    prediction = prediction_for(hit)
    forged = EvidenceRef.from_dict({**prediction.evidence_refs[0].to_dict(), **changes})
    output, summary = write_form(tmp_path, replace(prediction, evidence_refs=[forged]), [hit])
    workbook = load_workbook(output)
    assert workbook["当前表单"]["D2"].value is None
    assert summary.written_count == 0
    audit = read_jsonl(tmp_path / "writeback_audit.jsonl")[0]
    assert audit["status"] == "flagged" and code in audit["reason"]
    assert audit["writeback_action"] == "review_only"
    workbook.close()


def test_reference_from_another_fields_hits_cannot_authorize_writeback(tmp_path: Path) -> None:
    hit = source_hit()
    template = form_template(tmp_path / "template.xlsx")
    prediction = prediction_for(hit)
    summary = patch_workbook(
        template, [prediction], tmp_path / "filled_form.xlsx",
        retrieval_hits_by_field_id={"different_field": [hit]},
    )
    assert summary.written_count == 0
    assert read_jsonl(tmp_path / "writeback_audit.jsonl")[0]["error_code"] == "EV_REF_NOT_IN_RETRIEVAL"


@pytest.mark.parametrize(
    ("changes", "code"),
    [({"source_chunk_ids": ["unselected_source"]}, "EV_REF_NOT_IN_RETRIEVAL"),
     ({"evidence_attachment_ids": ["unselected_attachment"]}, "EV_ATTACHMENT_NOT_FOUND")],
)
def test_valid_ref_does_not_authorize_different_model_selected_sources_or_images(tmp_path: Path, changes: dict, code: str) -> None:
    hit = source_hit()
    prediction = replace(prediction_for(hit), **changes)
    _, summary = write_form(tmp_path, prediction, [hit])
    assert summary.written_count == 0
    assert read_jsonl(tmp_path / "writeback_audit.jsonl")[0]["error_code"] == code


def test_typed_image_metadata_is_rendered_at_back_of_evidence_sheet(tmp_path: Path) -> None:
    image = tmp_path / "proof.png"
    image.write_bytes(base64.b64decode(PNG))
    hit = source_hit(proof_attachment_ids=["proof_1"], proof_attachments=[{
        "attachment_id": "proof_1", "image_path": str(image), "mapping_status": "mapped", "source_cell": "B2",
    }])
    output, summary = write_form(tmp_path, prediction_for(hit), [hit], run_id="unit_run")
    workbook = load_workbook(output)
    assert summary.written_count == 1
    assert len(workbook["Evidence"]._images) == 1
    assert not workbook["当前表单"]._images
    assert workbook["Evidence"]._images[0].anchor._from.row > 1
    assert hit["raw_source_text"] not in workbook["当前表单"]["D2"].comment.text
    workbook.close()


@pytest.mark.parametrize("mode", ["append_sheet", "adjacent_columns"])
def test_user_evidence_sheet_and_selected_target_survive_generated_evidence_output(tmp_path: Path, mode: str) -> None:
    hit = source_hit()
    template = form_template(tmp_path / "template.xlsx")
    workbook = load_workbook(template)
    user_sheet = workbook.create_sheet("Evidence")
    # Even matching column labels are not proof that we own a user's sheet.
    user_sheet.append(["Field", "Answer", "Status", "Source", "Location", "Evidence"])
    user_sheet["A5"] = "用户已有内容"
    user_sheet["F5"] = "=SUM(A1:A2)"
    workbook.save(template)
    workbook.close()
    prediction = replace(prediction_for(hit), target_cell="'Evidence'!D2")
    output = tmp_path / "filled_form.xlsx"
    summary = patch_workbook(
        template, [prediction], output, retrieval_hits_by_field_id={prediction.field_id: [hit]},
        writeback_config={"evidence_image_mode": mode},
    )
    workbook = load_workbook(output)
    assert summary.written_count == 1
    assert workbook["Evidence"]["D2"].value == "500kVA"
    assert workbook["Evidence"]["A5"].value == "用户已有内容"
    assert workbook["Evidence"]["F5"].value == "=SUM(A1:A2)"
    if mode == "append_sheet":
        assert "Evidence_1" in workbook.sheetnames
        assert "Evidence_1!A2" in workbook["Evidence"]["D2"].comment.text
        assert workbook["Evidence_1"]["F2"].value == hit["raw_source_text"]
    else:
        assert "Evidence_1" not in workbook.sheetnames
    workbook.close()


def test_second_writeback_regenerates_only_positively_marked_sheet(tmp_path: Path) -> None:
    hit = source_hit()
    template = form_template(tmp_path / "template.xlsx")
    workbook = load_workbook(template)
    workbook.create_sheet("Evidence")["A1"] = "用户证据，必须保留"
    workbook.save(template)
    workbook.close()
    prediction = prediction_for(hit)
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    kwargs = {"retrieval_hits_by_field_id": {prediction.field_id: [hit]}}
    patch_workbook(template, [prediction], first, **kwargs)
    patch_workbook(first, [prediction], second, **kwargs)
    workbook = load_workbook(second)
    assert workbook.sheetnames == ["当前表单", "Evidence", "Evidence_1"]
    assert workbook["Evidence"]["A1"].value == "用户证据，必须保留"
    assert workbook["Evidence_1"].max_row == 2
    assert "Evidence_1!A2" in workbook["当前表单"]["D2"].comment.text
    workbook.close()


def write_strict_artifacts(tmp_path: Path, *, image: bool = False) -> tuple[dict, dict]:
    hit = source_hit()
    if image:
        image_path = tmp_path / "proof.png"
        image_path.write_bytes(base64.b64decode(PNG))
        hit = source_hit(proof_attachments=[{"attachment_id": "proof_1", "image_path": str(image_path), "source_cell": "B2"}])
    prediction = prediction_for(hit)
    _, summary = write_form(tmp_path, prediction, [hit], run_id="unit_run")
    raw = prediction.to_dict()
    write_jsonl(tmp_path / "predictions_raw.jsonl", [raw])
    write_jsonl(tmp_path / "predictions.jsonl", [raw])
    write_jsonl(tmp_path / "agent_overlays.jsonl", [{"field_id": prediction.field_id, "writeback_allowed": True}])
    write_jsonl(tmp_path / "predictions_agent_view.jsonl", [raw])
    write_jsonl(tmp_path / "retrieval_evidence.jsonl", [{"field_id": prediction.field_id, "top_hits": [hit]}])
    write_jsonl(tmp_path / "trace.jsonl", [])
    write_json(tmp_path / "trace_summary.json", {"total_fields": 1})
    write_json(tmp_path / "summary.json", {"fields_total": 1})
    (tmp_path / "run_summary.md").write_text("# Native evidence unit artifact\n", encoding="utf-8")
    manifest = {
        "schema_version": "1.2", "run_id": "unit_run", "status": "completed", "writeback_enabled": True,
        "artifacts": {"retrieval_evidence": "retrieval_evidence.jsonl", "filled_form": "filled_form.xlsx"},
        "writeback": {
            "summary": {"confirmed": 1, "uncertain": 0, "flagged": 0, "written": 1, "review": 0},
            "fields": summary.fields,
        },
    }
    write_json(tmp_path / "run_manifest.json", manifest)
    return raw, manifest


@pytest.mark.parametrize("image", [False, True])
def test_new_manifest_validates_real_writer_artifact_against_authority_hits(tmp_path: Path, image: bool) -> None:
    write_strict_artifacts(tmp_path, image=image)
    assert validate_step15_artifacts(tmp_path)["evidence_validation"] == "strict"


@pytest.mark.parametrize(
    ("changes", "code"),
    [
        ({"chunk_id": "forged"}, "EV_REF_NOT_IN_RETRIEVAL"),
        ({"file_name": "伪造.xlsx"}, "EV_FILE_NOT_IN_KB"),
        ({"knowledge_base_id": "伪造KB"}, "EV_FILE_NOT_IN_KB"),
        ({"relative_path": "伪造路径.xlsx"}, "EV_FILE_NOT_IN_KB"),
        ({"sheet_name": "伪造表名"}, "EV_REF_NOT_IN_RETRIEVAL"),
        ({"cell_range": "XFE1"}, "EV_CELL_RANGE_INVALID"),
        ({"cell_range": "B3:C3"}, "EV_REF_NOT_IN_RETRIEVAL"),
        ({"source_text": "伪造原文"}, "EV_REF_NOT_IN_RETRIEVAL"),
        ({"source_text": ""}, "WB_MISSING_EVIDENCE"),
        ({"attachment_ids": ["fake"]}, "EV_ATTACHMENT_NOT_FOUND"),
        ({"verified": True}, "EV_REF_NOT_IN_RETRIEVAL"),
    ],
)
def test_artifact_validator_rejects_tampering_of_raw_typed_refs(tmp_path: Path, changes: dict, code: str) -> None:
    raw, _ = write_strict_artifacts(tmp_path)
    raw["evidence_refs"][0].update(changes)
    write_jsonl(tmp_path / "predictions_raw.jsonl", [raw])
    write_jsonl(tmp_path / "predictions.jsonl", [raw])
    with pytest.raises(ArtifactValidationError, match=code):
        validate_step15_artifacts(tmp_path)


@pytest.mark.parametrize("location", ["manifest", "audit"])
def test_artifact_validator_does_not_trust_confirmed_manifest_or_audit_refs(tmp_path: Path, location: str) -> None:
    _, manifest = write_strict_artifacts(tmp_path)
    if location == "manifest":
        manifest["writeback"]["fields"][0]["evidence_refs"][0]["source_text"] = "伪造原文"
        write_json(tmp_path / "run_manifest.json", manifest)
    else:
        rows = read_jsonl(tmp_path / "writeback_audit.jsonl")
        rows[0]["evidence_refs"][0]["file_name"] = "伪造文件.xlsx"
        write_jsonl(tmp_path / "writeback_audit.jsonl", rows)
    with pytest.raises(ArtifactValidationError, match="EV_REF_NOT_IN_RETRIEVAL|EV_FILE_NOT_IN_KB"):
        validate_step15_artifacts(tmp_path)


@pytest.mark.parametrize("location", ["manifest", "audit"])
@pytest.mark.parametrize(
    ("source_damage", "error"),
    [
        ("forged", "EV_REF_NOT_IN_RETRIEVAL: reference_does_not_match_hit"),
        ("empty", "WB_MISSING_EVIDENCE: source_text is empty"),
        ("missing", "WB_MISSING_EVIDENCE: source_text is empty"),
    ],
)
def test_review_display_label_cannot_hide_tampered_typed_authority(
    tmp_path: Path, location: str, source_damage: str, error: str,
) -> None:
    _, manifest = write_strict_artifacts(tmp_path)
    audit = read_jsonl(tmp_path / "writeback_audit.jsonl")
    for row in [manifest["writeback"]["fields"][0], audit[0]]:
        row.update(status="flagged", writeback_action="review_only", old_value=None, new_value=None)
    manifest["writeback"]["summary"].update(confirmed=0, flagged=1, written=0, review=1)
    workbook = load_workbook(tmp_path / "filled_form.xlsx")
    workbook["当前表单"]["D2"].value = None
    workbook["当前表单"]["D2"].comment = None
    workbook.save(tmp_path / "filled_form.xlsx")
    workbook.close()
    write_json(tmp_path / "run_manifest.json", manifest)
    write_jsonl(tmp_path / "writeback_audit.jsonl", audit)
    assert validate_step15_artifacts(tmp_path)["evidence_validation"] == "strict"

    row = manifest["writeback"]["fields"][0] if location == "manifest" else audit[0]
    ref = row["evidence_refs"][0]
    ref["reference_contract"] = "review-display-v1"
    assert ref["knowledge_base_id"] and ref["source_text_hash"]
    if source_damage == "missing":
        del ref["source_text"]
    else:
        ref["source_text"] = "伪造原文" if source_damage == "forged" else ""
    if location == "manifest":
        write_json(tmp_path / "run_manifest.json", manifest)
    else:
        write_jsonl(tmp_path / "writeback_audit.jsonl", audit)

    # Require the authority error itself: a manifest/audit mismatch alone would
    # also fail, even if the forged display label bypassed typed validation.
    with pytest.raises(ArtifactValidationError, match=error):
        validate_step15_artifacts(tmp_path)


@pytest.mark.parametrize("location", ["manifest", "audit"])
def test_confirmed_display_reference_cannot_replace_typed_authority(tmp_path: Path, location: str) -> None:
    _, manifest = write_strict_artifacts(tmp_path)
    assert validate_step15_artifacts(tmp_path)["evidence_validation"] == "strict"
    audit = read_jsonl(tmp_path / "writeback_audit.jsonl")
    row = manifest["writeback"]["fields"][0] if location == "manifest" else audit[0]
    typed = row["evidence_refs"][0]
    row["evidence_refs"] = [{
        "reference_contract": "review-display-v1",
        "chunk_id": typed["chunk_id"], "evidence_kind": typed["evidence_kind"],
        "namespace": typed["namespace"], "file_name": typed["file_name"],
        "relative_path": typed["relative_path"], "sheet_name": typed["sheet_name"],
        "cell_range": typed["cell_range"], "text_preview": typed["source_text"],
    }]
    if location == "manifest":
        write_json(tmp_path / "run_manifest.json", manifest)
    else:
        write_jsonl(tmp_path / "writeback_audit.jsonl", audit)

    with pytest.raises(ArtifactValidationError, match="WB_MISSING_EVIDENCE: source_text is empty"):
        validate_step15_artifacts(tmp_path)


def test_confirmed_without_raw_refs_is_invalid_even_if_manifest_is_valid(tmp_path: Path) -> None:
    raw, _ = write_strict_artifacts(tmp_path)
    raw["evidence_refs"] = []
    write_jsonl(tmp_path / "predictions_raw.jsonl", [raw])
    write_jsonl(tmp_path / "predictions.jsonl", [raw])
    with pytest.raises(ArtifactValidationError, match="WB_MISSING_EVIDENCE"):
        validate_step15_artifacts(tmp_path)


@pytest.mark.parametrize(
    ("changes", "code"),
    [({"source_chunk_ids": ["unselected_source"]}, "EV_REF_NOT_IN_RETRIEVAL"),
     ({"evidence_attachment_ids": ["unselected_attachment"]}, "EV_ATTACHMENT_NOT_FOUND")],
)
def test_artifact_validator_checks_confirmed_model_selection_linkage(tmp_path: Path, changes: dict, code: str) -> None:
    raw, _ = write_strict_artifacts(tmp_path)
    raw.update(changes)
    write_jsonl(tmp_path / "predictions_raw.jsonl", [raw])
    write_jsonl(tmp_path / "predictions.jsonl", [raw])
    with pytest.raises(ArtifactValidationError, match=code):
        validate_step15_artifacts(tmp_path)


def test_new_schema_requires_field_scoped_authority_artifact(tmp_path: Path) -> None:
    _, manifest = write_strict_artifacts(tmp_path)
    del manifest["artifacts"]["retrieval_evidence"]
    write_json(tmp_path / "run_manifest.json", manifest)
    with pytest.raises(ArtifactValidationError, match="EV_REF_NOT_IN_RETRIEVAL"):
        validate_step15_artifacts(tmp_path)


def test_image_presentation_path_tampering_is_not_hidden_by_valid_canonical_ref(tmp_path: Path) -> None:
    _, manifest = write_strict_artifacts(tmp_path, image=True)
    manifest["writeback"]["fields"][0]["evidence_refs"][0]["image_path"] = "/forged/image.png"
    write_json(tmp_path / "run_manifest.json", manifest)
    with pytest.raises(ArtifactValidationError, match="EV_ATTACHMENT_NOT_FOUND"):
        validate_step15_artifacts(tmp_path)


def test_file_writer_uses_real_retrieval_artifact_and_rejects_missing_authority(tmp_path: Path) -> None:
    hit = source_hit()
    prediction = prediction_for(hit)
    template = form_template(tmp_path / "template.xlsx")
    predictions_path = tmp_path / "predictions_raw.jsonl"
    write_jsonl(predictions_path, [prediction.to_dict()])
    write_jsonl(tmp_path / "retrieval_evidence.jsonl", [{"field_id": prediction.field_id, "top_hits": [hit]}])
    summary = writeback_from_files(template_path=template, predictions_path=predictions_path, output_path=tmp_path / "filled.xlsx")
    assert summary.written_count == 1
    (tmp_path / "retrieval_evidence.jsonl").unlink()
    summary = writeback_from_files(template_path=template, predictions_path=predictions_path, output_path=tmp_path / "rejected.xlsx")
    assert summary.written_count == 0


def test_duplicate_hit_identity_with_conflicting_source_rejects_writer(tmp_path: Path) -> None:
    hit = source_hit()
    conflict = deepcopy(hit)
    conflict["raw_source_text"] = "同ID但不同来源内容"
    output, summary = write_form(tmp_path, prediction_for(hit), [hit, conflict])
    assert summary.written_count == 0
    workbook = load_workbook(output)
    assert workbook["当前表单"]["D2"].value is None
    workbook.close()
