from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from nested_doc_rag.ingestion import build_ingestion_records


def ingest(path: Path) -> list[dict]:
    records, skipped = build_ingestion_records(path.parent, namespace="xixian_4", knowledge_base_id="kb-upload")
    assert skipped == []
    return records


def test_excel_separates_structured_fields_from_headers_empty_values_and_formulas(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "动力"
    sheet.append(["类别", "实际值", "字段", "备注"])
    sheet.append(["供电", 0, "油机数量", "现场记录"])
    sheet.append([None, False, "备用电源可用", None])
    sheet.append([None, None, "UPS容量", None])
    sheet.append([None, "=SUM(B2:B3)", "计算数量", None])
    sheet.append([None, '=_xlfn.DISPIMG("ID_ONLY",1)', "铭牌照片", None])
    sheet.merge_cells("A2:A6")
    path = tmp_path / "动力.xlsx"
    workbook.save(path)

    records = ingest(path)

    assert [record["evidence_kind"] for record in records] == [
        "table_row", "structured_field", "structured_field", "table_row", "table_row", "table_row",
    ]
    assert records[1]["field_name"] == "油机数量"
    assert records[1]["field_value"] == "0"
    assert records[2]["field_value"] == "False"
    assert records[1]["address"]["cell_range"] == "A2:D2"
    assert records[2]["address"]["cell_range"] == "B3:C3"
    assert records[2]["structural_path"] == ["供电"]
    assert "供电" not in records[2]["raw_source_text"]
    context = records[2]["metadata"]["cells"][0]
    assert context["raw_value"] is None
    assert context["effective_value"] == "供电"
    assert context["value_origin"] == "A2"
    assert context["merged_range"] == "A2:A6"
    assert records[4]["metadata"]["cells"][1]["is_formula"] is True
    assert records[5]["metadata"]["cells"][1]["is_dispimg"] is True
    assert all(record["field_value"] is None for record in records[3:])
    assert all(record["source_type"] == "uploaded_excel_row" for record in records)


def test_excel_handles_merged_headers_and_preserves_generic_table_rows(tmp_path: Path) -> None:
    workbook = Workbook()
    structured = workbook.active
    structured.title = "配置"
    structured.append(["字段", "值", "图片"])
    structured.merge_cells("A1:A2")
    structured.merge_cells("B1:B2")
    structured.merge_cells("C1:C2")
    structured.append([None, None, None])
    structured.append(["供电路数", 2, None])
    generic = workbook.create_sheet("通信协议")
    generic.append(["协议", "用途"])
    generic.append(["BGP", "路由交换"])
    generic.append(["OSPF", "内部路由"])
    pairs = workbook.create_sheet("清晰字段")
    pairs.append(["UPS容量", "2400 kVA"])
    path = tmp_path / "清单.xlsx"
    workbook.save(path)

    records = ingest(path)

    values = [record for record in records if record["evidence_kind"] == "structured_field"]
    assert [(record["field_name"], record["field_value"]) for record in values] == [("供电路数", "2"), ("UPS容量", "2400 kVA")]
    assert values[0]["address"]["row_index"] == 3
    assert all(record["evidence_kind"] == "table_row" for record in records if record["sheet_name"] == "通信协议")


def test_word_preserves_block_order_heading_paths_and_original_indices(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("动力系统", level=1)
    document.add_paragraph("  市电：两路\n分别接入😀  ")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "设备"
    table.cell(0, 1).text = "容量"
    table.cell(1, 0).text = "UPS"
    table.cell(1, 1).text = "2400 kVA"
    document.add_paragraph("")
    document.add_heading("冗余", level=2)
    document.add_paragraph("N+1")
    document.add_heading("制冷系统", level=1)
    document.add_paragraph("冷通道封闭")
    path = tmp_path / "现场说明.docx"
    document.save(path)

    records = ingest(path)

    assert [record["evidence_kind"] for record in records] == [
        "paragraph", "paragraph", "table_row", "table_row", "paragraph", "paragraph", "paragraph", "paragraph",
    ]
    assert [record["metadata"]["block_index"] for record in records] == [1, 2, 3, 3, 5, 6, 7, 8]
    assert records[1]["raw_source_text"] == "  市电：两路\n分别接入😀  "
    assert records[2]["address"]["table_index"] == 1
    assert records[3]["address"]["row_index"] == 2
    assert records[4]["address"]["paragraph_index"] == 4
    assert records[5]["structural_path"] == ["动力系统", "冗余"]
    assert records[7]["structural_path"] == ["制冷系统"]
    assert records[3]["field_name"] is records[3]["field_value"] is None
    for record in records:
        assert record["source_text_hash"] == "sha256:" + hashlib.sha256(record["raw_source_text"].encode()).hexdigest()


def test_text_native_whitespace_survives_canonical_ingestion(tmp_path: Path) -> None:
    path = tmp_path / "记录.txt"
    original = "  第一行😀\r\n\t第二行\r\n"
    path.write_bytes(original.encode())

    records = ingest(path)

    assert len(records) == 1
    assert records[0]["evidence_kind"] == "document_chunk"
    assert records[0]["raw_source_text"] == original
    assert records[0]["source_text_hash"] == "sha256:" + hashlib.sha256(original.encode()).hexdigest()


def test_structured_values_keep_complete_native_rows_and_stable_ids(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["字段", "值"])
    value = "现场记录😀  " * 400
    sheet.append(["运行状态", value])
    path = tmp_path / "长字段.xlsx"
    workbook.save(path)

    records = ingest(path)
    repeated = ingest(path)

    assert len(records) == 2
    assert records[1]["field_value"] == value
    assert records[1]["raw_source_text"] == "运行状态 / " + value
    assert records[1]["address"]["source_anchor"] == "Sheet!row 2"
    assert records[1]["parser_version"] == "native-office-v2"
    assert records[1]["chunk_id"] == repeated[1]["chunk_id"]
    assert records[1]["point_id"] == repeated[1]["point_id"]


def test_unmapped_image_only_excel_row_preserves_formula_identity(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    formula = '=_xlfn.DISPIMG("ID_MISSING",1)'
    sheet["C3"] = formula
    path = tmp_path / "未映射图片.xlsx"
    workbook.save(path)

    records = ingest(path)

    assert len(records) == 1
    assert records[0]["evidence_kind"] == "table_row"
    assert records[0]["raw_source_text"] == formula
    assert records[0]["field_value"] is None
    assert records[0]["proof_attachments"] == []
    assert records[0]["metadata"]["cells"][2]["is_dispimg"] is True
    assert records[0]["address"]["cell_range"] == "C3:C3"
