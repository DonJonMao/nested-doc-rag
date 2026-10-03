from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook
from qdrant_client import QdrantClient

from nested_doc_rag.config import load_app_config
from nested_doc_rag.ingestion import IngestionOptions, build_ingestion_records, run_knowledge_ingestion
from nested_doc_rag.io import read_json, read_jsonl
from nested_doc_rag.retrieval.field_schema import (
    FIELD_SCHEMA_CONTRACT_VERSION,
    bind_field_schema_families,
    build_field_schema_records,
    field_schema_for_record,
)


def value_record(**overrides: Any) -> dict[str, Any]:
    return {
        "chunk_id": "value-1", "point_id": "physical-point", "namespace": "room301",
        "knowledge_base_id": "kb1", "document_id": "document-v1", "index_version": "v1",
        "source_type": "uploaded_excel_row", "evidence_kind": "structured_field",
        "corpus_layer": "fact", "file_name": "参数.xlsx", "relative_path": "动力/参数.xlsx",
        "sheet_name": "当前设备", "structural_path": ["动力", "UPS"],
        "field_name": "额定容量", "field_value": "VALUE_SECRET_739",
        "raw_source_text": "额定容量 / VALUE_SECRET_739", "raw_text": "VALUE_SECRET_739",
        "source_text_hash": "sha256:VALUE_SECRET_739", "source_text_hash_space": "raw_source_text",
        "address": {"file_name": "参数.xlsx", "sheet_name": "当前设备", "row_index": 4, "cell_range": "A4:C4"},
        "proof_attachments": [{"attachment_id": "proof", "image_path": "/actual/proof.png"}],
        "metadata": {
            "column_headers": {3: "单位", 1: "字段", 2: "当前值"},
            "field_schema": {"field_name": "额定容量", "unit": "kVA"},
            "heldout": "VALUE_SECRET_739", "gold": {"answer": "VALUE_SECRET_739"},
        },
        **overrides,
    }


def test_schema_payload_is_value_free_and_is_not_an_evidence_record() -> None:
    record = value_record()
    original = deepcopy(record)
    schema = build_field_schema_records(iter([record]))[0]

    assert schema["schema_contract_version"] == FIELD_SCHEMA_CONTRACT_VERSION == "field-schema-v1"
    assert schema["retrieval_object"] == schema["corpus_layer"] == "field_schema"
    assert schema["column_headers"] == ["字段", "当前值", "单位"]
    assert schema["unit"] == "kVA"
    assert schema["schema_text"] == schema["text_for_embedding"] == "字段：额定容量。类别：动力 / UPS。工作表：当前设备。列头：字段 / 当前值 / 单位。单位：kVA"
    assert schema["namespace"] == "room301" and schema["knowledge_base_id"] == "kb1"
    assert schema["index_version"] == "v1" and schema["document_id"] == "document-v1"
    assert "VALUE_SECRET_739" not in json.dumps(schema, ensure_ascii=False)
    assert {"evidence_kind", "chunk_id", "evidence_id", "field_value", "raw_text", "raw_source_text", "source_text_hash", "address", "proof_attachments"}.isdisjoint(schema)
    assert record == original


def test_family_and_schema_text_are_independent_of_values_raw_hash_and_version() -> None:
    first = value_record()
    second = value_record(field_value="OTHER_VALUE", raw_text="OTHER_VALUE", raw_source_text="OTHER_VALUE", source_text_hash="different", index_version="v2", document_id="document-v2")
    one, two = build_field_schema_records([first])[0], build_field_schema_records([second])[0]

    assert one["field_family_id"] == two["field_family_id"]
    assert one["schema_text"] == two["schema_text"]
    assert one["schema_id"] != two["schema_id"] and one["point_id"] != two["point_id"]
    assert two["index_version"] == "v2" and two["document_id"] == "document-v2"


@pytest.mark.parametrize("override", [
    {"namespace": "room401"}, {"knowledge_base_id": "kb2"},
    {"relative_path": "另一份/参数.xlsx"}, {"sheet_name": "规划设备"},
    {"structural_path": ["动力", "油机"]},
    {"metadata": {"field_schema": {"field_name": "UPS品牌", "unit": "kVA"}, "column_headers": {1: "字段", 2: "当前值", 3: "单位"}}},
    {"metadata": {"field_schema": {"field_name": "额定容量", "unit": "kW"}, "column_headers": {1: "字段", 2: "当前值", 3: "单位"}}},
    {"metadata": {"field_schema": {"field_name": "额定容量", "unit": "kVA"}, "column_headers": {1: "字段", 2: "规划值", 3: "单位"}}},
])
def test_structural_or_scope_changes_create_distinct_families(override: dict[str, Any]) -> None:
    base = build_field_schema_records([value_record()])[0]
    changed = build_field_schema_records([value_record(**override)])[0]

    assert base["field_family_id"] != changed["field_family_id"]


def test_family_binding_detaches_inputs_and_preserves_full_physical_evidence() -> None:
    record = value_record(field_family_id="forged-family")
    original = deepcopy(record)
    bound = bind_field_schema_families(iter([record]))[0]
    schema = build_field_schema_records([record])[0]

    assert bound["retrieval_object"] == "evidence"
    assert bound["field_family_id"] == schema["field_family_id"] != "forged-family"
    for key in ("chunk_id", "point_id", "field_value", "raw_source_text", "source_text_hash", "address", "proof_attachments"):
        assert bound[key] == record[key]
    bound["address"]["cell_range"] = "Z999"
    bound["proof_attachments"][0]["image_path"] = "/forged"
    assert record == original


def test_duplicate_rows_make_one_schema_point_per_family_and_version() -> None:
    first = value_record()
    second = value_record(chunk_id="value-2", field_value="different", raw_source_text="different")
    next_version = value_record(index_version="v2", document_id="document-v2")
    schemas = build_field_schema_records([second, next_version, first])

    assert len(schemas) == 2
    assert {schema["index_version"] for schema in schemas} == {"v1", "v2"}
    assert len({schema["field_family_id"] for schema in schemas}) == 1
    assert len({schema["point_id"] for schema in schemas}) == 2


@pytest.mark.parametrize("override", [
    {"field_name": None, "metadata": {}}, {"sheet_name": None}, {"namespace": None},
    {"field_name": {"gold": "SECRET"}, "metadata": {}},
    {"field_name": "=A1", "metadata": {}},
    {"source_type": "uploaded_docx_table_row", "relative_path": "参数.docx"},
])
def test_unsupported_or_ambiguous_records_are_not_guessed_into_schema(override: dict[str, Any]) -> None:
    record = value_record(**override)
    bound = bind_field_schema_families([record])[0]

    assert field_schema_for_record(record) is None
    assert build_field_schema_records([record]) == []
    assert bound["retrieval_object"] == "evidence" and "field_family_id" not in bound


def test_schema_points_cannot_be_rebound_as_physical_evidence() -> None:
    schema = build_field_schema_records([value_record()])[0]

    assert build_field_schema_records([schema]) == []
    with pytest.raises(ValueError, match="cannot be bound as evidence"):
        bind_field_schema_families([schema])


def save_fields(path: Path, values: list[Any]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "现场设备"
    sheet.append(["类别", "字段", "当前值", "单位"])
    for label, value in zip(["UPS容量", "UPS冗余模式", "UPS照片", "油机数量"], values, strict=True):
        sheet.append(["UPS" if label.startswith("UPS") else "油机", label, value, "kVA" if label == "UPS容量" else None])
    workbook.save(path)
    workbook.close()


def test_native_excel_schema_covers_empty_formula_and_image_fields_without_fabricating_values(tmp_path: Path) -> None:
    save_fields(tmp_path / "参数.xlsx", [None, "=A1", '=_xlfn.DISPIMG("PHOTO",1)', 0])
    records, skipped = build_ingestion_records(tmp_path, namespace="room301", knowledge_base_id="kb1")
    schemas = build_field_schema_records(records)

    assert skipped == [] and len(records) == 5 and len(schemas) == 4
    assert all(record["retrieval_object"] == "evidence" for record in records)
    assert "field_family_id" not in records[0]
    assert [record["field_value"] for record in records[1:]] == [None, None, None, "0"]
    assert [record["evidence_kind"] for record in records[1:]] == ["table_row", "table_row", "table_row", "structured_field"]
    assert {schema["field_name"] for schema in schemas} == {"UPS容量", "UPS冗余模式", "UPS照片", "油机数量"}
    assert all(record.get("field_family_id") for record in records[1:])
    assert next(schema for schema in schemas if schema["field_name"] == "UPS容量")["unit"] == "kVA"
    assert "DISPIMG" not in json.dumps(schemas) and "=A1" not in json.dumps(schemas)


def test_value_permutation_does_not_change_native_family_selection_or_schema_embedding(tmp_path: Path) -> None:
    path = tmp_path / "参数.xlsx"
    save_fields(path, ["VALUE_A", "VALUE_B", "VALUE_C", "VALUE_D"])
    before, _ = build_ingestion_records(tmp_path, namespace="room301", knowledge_base_id="kb1")
    save_fields(path, ["VALUE_D", "VALUE_C", "VALUE_B", "VALUE_A"])
    after, _ = build_ingestion_records(tmp_path, namespace="room301", knowledge_base_id="kb1")
    schemas_before, schemas_after = build_field_schema_records(before), build_field_schema_records(after)

    assert [(row["field_family_id"], row["text_for_embedding"]) for row in schemas_before] == [(row["field_family_id"], row["text_for_embedding"]) for row in schemas_after]
    assert "VALUE_" not in json.dumps(schemas_before) and "VALUE_" not in json.dumps(schemas_after)
    assert [row["field_value"] for row in before[1:]] != [row["field_value"] for row in after[1:]]


def test_clear_headerless_labels_keep_schema_when_actual_value_is_missing_or_formula(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["UPS容量", None])
    sheet.append(["UPS冗余模式", "=A1"])
    sheet.append(["自由说明", None])
    workbook.save(tmp_path / "现场.xlsx")
    workbook.close()
    records, _ = build_ingestion_records(tmp_path, namespace="room301", knowledge_base_id="kb1")

    assert {schema["field_name"] for schema in build_field_schema_records(records)} == {"UPS容量", "UPS冗余模式"}
    assert all(record["evidence_kind"] == "table_row" and record["field_value"] is None for record in records)
    assert "field_family_id" not in records[2]


def test_units_are_read_from_labels_headers_or_unit_column_and_not_value_text(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["字段", "值"])
    sheet.append(["UPS额定容量(kVA)", "500 kW"])
    sheet.append(["油机容量", "2000 kVA"])
    workbook.save(tmp_path / "容量.xlsx")
    workbook.close()
    records, _ = build_ingestion_records(tmp_path, namespace="room301", knowledge_base_id="kb1")
    schemas = {schema["field_name"]: schema for schema in build_field_schema_records(records)}

    assert schemas["UPS额定容量(kVA)"]["unit"] == "kVA"
    assert schemas["油机容量"]["unit"] is None
    assert all("500" not in schema["text_for_embedding"] and "2000" not in schema["text_for_embedding"] for schema in schemas.values())


def test_ingestion_indexes_auxiliary_points_and_keeps_evidence_artifact_counts_separate(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    uploads = tmp_path / "uploads"
    uploads.mkdir()
    save_fields(uploads / "参数.xlsx", ["500", "N+1", None, 0])
    client = QdrantClient(":memory:")

    class FixedEmbedding:
        def __init__(self, **kwargs: Any) -> None:
            pass

        def embed(self, texts: list[str]) -> list[list[float]]:
            return [[1.0, 0.0, 0.0] for _ in texts]

    monkeypatch.setattr("nested_doc_rag.ingestion.EmbeddingClient", FixedEmbedding)
    monkeypatch.setattr("nested_doc_rag.ingestion.build_qdrant_client", lambda **kwargs: client)
    monkeypatch.setattr(client, "close", lambda: None)
    config = load_app_config(project_root=tmp_path, default_config=tmp_path / "missing", env={})
    output = tmp_path / "out"
    try:
        summary = run_knowledge_ingestion(IngestionOptions(uploads, "room301", "kb1", output, config))
        evidence = read_jsonl(output / "ingestion_manifest.jsonl")
        schemas = read_jsonl(output / "field_schema_manifest.jsonl")
        manifest = read_json(output / "run_manifest.json")
        points, _ = client.scroll(config.qdrant.collection_name, limit=100, with_payload=True)

        assert summary["record_count"] == summary["upserted_count"] == len(evidence) == 5
        assert summary["schema_record_count"] == summary["schema_upserted_count"] == len(schemas) == 4
        assert summary["total_point_count"] == summary["total_upserted_count"] == len(points) == 9
        assert summary["counts_by_file"] == {"参数.xlsx": 5}
        assert summary["counts_by_source_type"] == {"uploaded_excel_row": 5}
        assert manifest["counts"]["total_fields"] == 5
        assert manifest["artifacts"]["field_schema_manifest"] == "field_schema_manifest.jsonl"
        assert all(row["retrieval_object"] == "evidence" for row in evidence)
        assert sum(point.payload["retrieval_object"] == "field_schema" for point in points) == 4
    finally:
        client._client.close()
