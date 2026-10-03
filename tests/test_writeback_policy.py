from __future__ import annotations

from copy import deepcopy
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from nested_doc_rag.evidence_resolver import resolve_evidence_refs
from nested_doc_rag.excel.writeback import (
    LegacyWritebackModeWarning,
    WritebackPolicy,
    patch_workbook,
    writeback_from_files,
)
from nested_doc_rag.io import read_json, read_jsonl, write_jsonl
from nested_doc_rag.schemas.eval import FieldPrediction


def prediction(field_id: str = "field-1", target_cell: str = "Sheet1!A1", status: str = "answered") -> tuple[FieldPrediction, dict[str, Any]]:
    source = {
        "chunk_id": "native-1", "namespace": "room301", "knowledge_base_id": "kb1",
        "file_name": "现场.xlsx", "sheet_name": "动力", "row_index": 3, "cell_range": "A3:B3",
        "evidence_kind": "structured_field", "raw_source_text": "UPS容量 / 500kVA", "field_name": "UPS容量", "field_value": "500kVA",
    }
    ref = resolve_evidence_refs(["native-1"], [source]).refs[0]
    return FieldPrediction(
        field_id=field_id, row_index=1, target_cell=target_cell, answer_value="500kVA", answer_status=status,
        confidence=0.99, source_chunk_ids=["native-1"], evidence_refs=[ref],
    ), source


def template(path: Path, old_value: Any, *, other_value: Any = None) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = old_value
    sheet["B1"] = other_value
    sheet["A1"].comment = Comment("原有人工备注", "user")
    sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FFF01234")
    sheet["A1"].font = Font(bold=True)
    workbook.save(path)
    workbook.close()


def write(tmp_path: Path, old_value: Any, *, policy: str = "preserve", status: str = "answered", allow_uncertain: bool = False, mode: str = "safe", authority: bool = True):
    path, output = tmp_path / "template.xlsx", tmp_path / "out.xlsx"
    template(path, old_value)
    pred, source = prediction(status=status)
    kwargs = {
        "retrieval_hits_by_field_id": {pred.field_id: [source]} if authority else {},
        "writeback_config": {"existing_value_policy": policy, "allow_uncertain": allow_uncertain},
        "overwrite_all_cli": policy == "overwrite_all", "mode": mode,
    }
    if mode == "overwrite":
        with pytest.warns(LegacyWritebackModeWarning):
            summary = patch_workbook(path, [pred], output, **kwargs)
    else:
        summary = patch_workbook(path, [pred], output, **kwargs)
    workbook = load_workbook(output)
    return summary, read_jsonl(tmp_path / "writeback_audit.jsonl")[0], workbook


@pytest.mark.parametrize("old_value", ["人工值", 0, False, " ", "\t", "\n", datetime(2025, 1, 2, 3, 4, 5)])
def test_preserve_rejects_every_nonempty_value_and_retains_original_comment_and_style(tmp_path: Path, old_value: Any) -> None:
    summary, audit, workbook = write(tmp_path, old_value)
    cell = workbook["Sheet1"]["A1"]

    assert cell.value == old_value and type(cell.value) is type(old_value)
    assert cell.comment.text == "原有人工备注" and cell.comment.author == "user"
    assert cell.fill.fgColor.rgb == "FFF01234" and cell.font.bold is True
    assert summary.written_count == 0 and summary.review_count == 1
    assert audit["reason"] == "target_non_empty" and audit["error_code"] == "WB_TARGET_NON_EMPTY"
    assert audit["writeback_action"] == "skipped_non_empty_cell"
    expected = old_value.isoformat() if isinstance(old_value, datetime) else old_value
    assert audit["old_value"] == audit["new_value"] == expected
    assert audit["answer_value"] == "500kVA" and audit["policy"] == "preserve"
    assert audit["evidence_refs"] and audit["evidence_count"] == 1
    assert read_jsonl(tmp_path / "review_items.jsonl")[0]["reason"] == "target_non_empty"
    assert {key: summary.fields[0][key] for key in ("old_value", "new_value", "policy")} == {key: audit[key] for key in ("old_value", "new_value", "policy")}
    workbook.close()


@pytest.mark.parametrize("old_value", [None, ""])
def test_preserve_writes_only_empty_cells_with_valid_authority(tmp_path: Path, old_value: Any) -> None:
    summary, audit, workbook = write(tmp_path, old_value)

    assert summary.written_count == 1 and workbook["Sheet1"]["A1"].value == "500kVA"
    assert audit["old_value"] is None and audit["new_value"] == "500kVA"
    assert audit["policy"] == "preserve"
    workbook.close()


@pytest.mark.parametrize("policy", ["overwrite_confirmed", "overwrite_all"])
def test_explicit_policies_overwrite_confirmed_ordinary_values_and_audit_the_change(tmp_path: Path, policy: str) -> None:
    summary, audit, workbook = write(tmp_path, "原有值", policy=policy)

    assert summary.written_count == 1 and workbook["Sheet1"]["A1"].value == "500kVA"
    assert audit["old_value"] == "原有值" and audit["new_value"] == "500kVA" and audit["policy"] == policy
    workbook.close()


@pytest.mark.parametrize("policy", ["preserve", "overwrite_confirmed", "overwrite_all"])
@pytest.mark.parametrize("mode", ["safe", "overwrite"])
def test_formulas_are_protected_under_all_modes_and_policies(tmp_path: Path, policy: str, mode: str) -> None:
    formula = "=SUM(C1:C4)"
    summary, audit, workbook = write(tmp_path, formula, policy=policy, mode=mode)

    assert workbook["Sheet1"]["A1"].value == formula and summary.formula_skipped_count == 1
    assert audit["old_value"] == audit["new_value"] == formula
    assert audit["reason"] == "skipped_formula" and audit["policy"] == policy
    assert workbook["Sheet1"]["A1"].comment.text == "原有人工备注"
    workbook.close()


def test_legacy_overwrite_mode_does_not_implicitly_allow_overwriting(tmp_path: Path) -> None:
    summary, audit, workbook = write(tmp_path, "人工值", mode="overwrite")

    assert summary.written_count == 0 and audit["policy"] == "preserve"
    assert workbook["Sheet1"]["A1"].value == "人工值"
    workbook.close()


@pytest.mark.parametrize("status", ["partial_clue", "not_found", "conflict_unresolved"])
def test_overwrite_confirmed_rejects_nonconfirmed_existing_cells(tmp_path: Path, status: str) -> None:
    summary, audit, workbook = write(tmp_path, "人工值", policy="overwrite_confirmed", status=status, allow_uncertain=True)

    assert summary.written_count == 0 and audit["error_code"] == "WB_OVERWRITE_POLICY"
    assert audit["reason"] == "overwrite_policy_rejected" and audit["old_value"] == audit["new_value"] == "人工值"
    assert workbook["Sheet1"]["A1"].value == "人工值"
    workbook.close()


def test_overwrite_all_still_obeys_allow_uncertain_and_typed_authority(tmp_path: Path) -> None:
    disabled = tmp_path / "disabled"
    missing = tmp_path / "missing"
    allowed = tmp_path / "allowed"
    disabled.mkdir()
    missing.mkdir()
    allowed.mkdir()
    skipped, disabled_audit, first = write(disabled, "人工值", policy="overwrite_all", status="partial_clue")
    rejected, missing_audit, second = write(missing, "人工值", policy="overwrite_all", status="partial_clue", allow_uncertain=True, authority=False)
    written, allowed_audit, third = write(allowed, "人工值", policy="overwrite_all", status="partial_clue", allow_uncertain=True)

    assert skipped.written_count == rejected.written_count == 0
    assert disabled_audit["writeback_action"] == "skipped_uncertain_policy"
    assert missing_audit["error_code"] == "EV_REF_NOT_IN_RETRIEVAL"
    assert disabled_audit["new_value"] == missing_audit["new_value"] == "人工值"
    assert written.written_count == 1 and allowed_audit["writeback_action"] == "written_red_comment"
    assert allowed_audit["old_value"] == "人工值" and allowed_audit["new_value"] == "500kVA"
    first.close()
    second.close()
    third.close()


@pytest.mark.parametrize("policy", ["overwrite_confirmed", "overwrite_all"])
def test_overwrite_policies_do_not_bypass_missing_authority_for_confirmed(tmp_path: Path, policy: str) -> None:
    summary, audit, workbook = write(tmp_path, "人工值", policy=policy, authority=False)

    assert summary.written_count == 0 and audit["error_code"] == "EV_REF_NOT_IN_RETRIEVAL"
    assert audit["old_value"] == audit["new_value"] == "人工值"
    workbook.close()


@pytest.mark.parametrize("entry", ["patch", "files"])
def test_overwrite_all_requires_explicit_cli_selection_before_reading_any_workbook_or_output(tmp_path: Path, entry: str, monkeypatch: pytest.MonkeyPatch) -> None:
    def unexpected_read(*args: Any, **kwargs: Any):
        raise AssertionError("must validate dangerous policy before file I/O")

    monkeypatch.setattr("nested_doc_rag.excel.writeback.load_workbook", unexpected_read)
    monkeypatch.setattr("nested_doc_rag.excel.writeback.read_jsonl", unexpected_read)
    output = tmp_path / "untouched" / "out.xlsx"
    kwargs = {"template_path": tmp_path / "missing.xlsx", "output_path": output, "writeback_config": {"existing_value_policy": "overwrite_all"}}
    with pytest.raises(ValueError, match="explicit CLI"):
        if entry == "patch":
            patch_workbook(predictions=[], **kwargs)
        else:
            writeback_from_files(predictions_path=tmp_path / "missing.jsonl", **kwargs)
    assert not output.parent.exists()


@pytest.mark.parametrize("invalid", [None, True, 1, "", "OVERWRITE_ALL", "overwrite"])
def test_invalid_policy_fails_closed(invalid: Any) -> None:
    with pytest.raises(ValueError, match="existing_value_policy"):
        WritebackPolicy.from_value({"existing_value_policy": invalid})


def test_invalid_duplicate_and_merged_targets_have_exact_original_and_final_value_audits(tmp_path: Path) -> None:
    path, output = tmp_path / "template.xlsx", tmp_path / "out.xlsx"
    template(path, "人工值")
    workbook = load_workbook(path)
    workbook["Sheet1"].merge_cells("C1:D1")
    workbook["Sheet1"]["C1"] = "合并人工值"
    workbook.save(path)
    workbook.close()
    predictions = [prediction("first")[0], prediction("second")[0], prediction("invalid", "Missing!Z999")[0], prediction("merged", "Sheet1!D1")[0]]
    authority = {pred.field_id: [prediction()[1]] for pred in predictions}
    summary = patch_workbook(path, predictions, output, retrieval_hits_by_field_id=authority)
    audits = {row["field_id"]: row for row in read_jsonl(tmp_path / "writeback_audit.jsonl")}

    assert summary.conflict_count == 2 and summary.invalid_count == 1 and summary.written_count == 0
    for key in ("first", "second"):
        assert audits[key]["old_value"] == audits[key]["new_value"] == "人工值"
        assert audits[key]["reason"] == "duplicate_target_cell"
    assert audits["invalid"]["old_value"] is audits["invalid"]["new_value"] is None
    assert audits["invalid"]["policy"] == "preserve"
    assert audits["merged"]["cell"] == "C1" and audits["merged"]["old_value"] == audits["merged"]["new_value"] == "合并人工值"


@pytest.mark.parametrize("other_value", [None, "人工值", "=SUM(C1:C4)"])
def test_auxiliary_adjacent_evidence_cannot_change_another_preserved_target(tmp_path: Path, other_value: Any) -> None:
    path, output = tmp_path / "template.xlsx", tmp_path / "out.xlsx"
    template(path, None)
    workbook = load_workbook(path)
    workbook["Sheet1"]["B2"] = other_value
    workbook.save(path)
    workbook.close()
    first, source = prediction(target_cell="Sheet1!A2")
    second, _ = prediction("second", "Sheet1!B2", "not_found")
    patch_workbook(path, [first, second], output, retrieval_hits_by_field_id={first.field_id: [source], second.field_id: [source]}, writeback_config={"evidence_image_mode": "adjacent_columns"})
    workbook = load_workbook(output)
    audits = {row["field_id"]: row for row in read_jsonl(tmp_path / "writeback_audit.jsonl")}

    assert workbook["Sheet1"]["B2"].value == other_value
    assert audits["second"]["old_value"] == audits["second"]["new_value"] == other_value
    workbook.close()


@pytest.mark.parametrize("embed_images", [False, True])
@pytest.mark.parametrize(("other_value", "protected_target"), [(None, True), ("人工值", True), ("=SUM(A3:A4)", True), ("人工值", False), ("=SUM(A3:A4)", False)])
def test_adjacent_image_output_cannot_change_preserved_targets_or_user_cells(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, other_value: Any, protected_target: bool, embed_images: bool) -> None:
    path, output = tmp_path / "template.xlsx", tmp_path / "out.xlsx"
    template(path, None)
    workbook = load_workbook(path)
    workbook["Sheet1"]["C2"] = other_value
    workbook["Sheet1"]["C2"].comment = Comment("人工图片列备注", "user")
    workbook.save(path)
    workbook.close()
    first, source = prediction(target_cell="Sheet1!A2")
    source.update(proof_attachment_ids=["proof-1"], proof_attachments=[{
        "attachment_id": "proof-1", "image_path": str(tmp_path / "missing.png"), "mapping_status": "mapped", "source_cell": "B3",
    }])
    first = replace(first, evidence_refs=resolve_evidence_refs(["native-1"], [source]).refs)
    predictions = [first]
    authority = {first.field_id: [source]}
    if protected_target:
        second, plain_source = prediction("second", "Sheet1!C2", "not_found")
        predictions.append(second)
        authority[second.field_id] = [plain_source]

    def unexpected_image(*args: Any, **kwargs: Any):
        raise AssertionError("must protect the image column before trying to embed a proof")

    monkeypatch.setattr("nested_doc_rag.excel.writeback.insert_evidence_image", unexpected_image)
    summary = patch_workbook(path, predictions, output, retrieval_hits_by_field_id=authority, writeback_config={
        "evidence_image_mode": "adjacent_columns", "embed_evidence_images": embed_images,
    })
    workbook = load_workbook(output)
    audits = {row["field_id"]: row for row in read_jsonl(tmp_path / "writeback_audit.jsonl")}

    assert summary.written_count == 1 and workbook["Sheet1"]["A2"].value == "500kVA"
    assert "状态:" in workbook["Sheet1"]["B2"].value
    assert audits[first.field_id]["evidence_refs"][0]["image_object_key"]
    assert workbook["Sheet1"]["C2"].value == other_value
    assert workbook["Sheet1"]["C2"].comment.text == "人工图片列备注"
    assert not workbook["Sheet1"]._images
    if protected_target:
        assert audits["second"]["old_value"] == audits["second"]["new_value"] == other_value
    workbook.close()


def test_file_entrypoint_forwards_policy_and_preserves_evidence_metadata(tmp_path: Path) -> None:
    path, output = tmp_path / "template.xlsx", tmp_path / "out.xlsx"
    template(path, "人工值")
    pred, source = prediction()
    original = deepcopy(pred.to_dict())
    write_jsonl(tmp_path / "predictions.jsonl", [original])
    write_jsonl(tmp_path / "retrieval_evidence.jsonl", [{"field_id": pred.field_id, "top_hits": [source]}])
    summary = writeback_from_files(template_path=path, predictions_path=tmp_path / "predictions.jsonl", output_path=output,
                                  writeback_config={"existing_value_policy": "overwrite_all"}, overwrite_all_cli=True)
    audit = read_jsonl(tmp_path / "writeback_audit.jsonl")[0]
    evidence = read_json(tmp_path / "evidence_map.json")["fields"][pred.field_id]

    assert summary.written_count == 1 and audit["old_value"] == "人工值" and audit["new_value"] == "500kVA"
    assert audit["policy"] == evidence["policy"] == "overwrite_all"
    assert audit["evidence_refs"] and pred.to_dict() == original
