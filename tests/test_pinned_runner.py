from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from nested_doc_rag.agent.step15_runner import Step15AgentRunner, attach_addressed_evidence
from nested_doc_rag.artifacts import ArtifactValidationError, validate_step15_artifacts
from nested_doc_rag.config import load_app_config
from nested_doc_rag.evaluation.step15_engine import Step15RetrievalResult
from nested_doc_rag.form.input_snapshot import FormInputMismatchError
from nested_doc_rag.io import read_json, read_jsonl, write_json, write_jsonl
from nested_doc_rag.retrieval.version_scope import IndexScopeError, versioned_point_id

KB = "10000000-0000-4000-8000-000000000001"
GLOBAL_KB = "10000000-0000-4000-8000-000000000002"
V1 = "20000000-0000-4000-8000-000000000001"
V2 = "20000000-0000-4000-8000-000000000002"
GLOBAL_V = "20000000-0000-4000-8000-000000000003"
COLLECTION = "pinned_runner"


def pins() -> list[dict[str, Any]]:
    return [{"collection": COLLECTION, "namespace": namespace, "knowledge_base_id": kb,
             "index_version_id": version, "storage_contract": "versioned_v1"}
            for namespace, kb, version in (("room", KB, V1), ("global", GLOBAL_KB, GLOBAL_V))]


def hit(**extra: Any) -> dict[str, Any]:
    point = versioned_point_id(V1, "native-ups")
    return {"chunk_id": "native-ups", "namespace": "room", "knowledge_base_id": KB,
            "index_version_id": V1, "index_version": V1, "point_id": point, "qdrant_point_id": point,
            "evidence_kind": "structured_field", "retrieval_object": "evidence", "source_type": "uploaded_excel_row",
            "corpus_layer": "fact", "file_name": "native.xlsx", "relative_path": "native.xlsx", "document_id": "document-native",
            "sheet_name": "参数", "row_index": 2, "cell_range": "A2:B2", "field_name": "UPS容量", "field_value": "500kVA",
            "raw_source_text": "UPS容量 / 500kVA", "raw_text": "UPS容量 / 500kVA", **extra}


def item() -> dict[str, Any]:
    return {"form_item_id": "ups", "file_name": "form.xlsx", "sheet_name": "Sheet1", "row_index": 2,
            "target_cell": "Sheet1!D2", "question_text": "UPS容量", "instruction_text": "填写容量",
            "category_path": ["动力"], "answer_example": "", "needs_evidence": True}


def runner(path: Path, *, scopes: list[dict[str, Any]] | None = None, resume: bool = False,
           pack: Step15RetrievalResult | None = None, forbidden_calls: bool = False) -> tuple[Step15AgentRunner, list[str]]:
    config = load_app_config(project_root=path, default_config=path / "missing.yaml", env={}, cli_overrides={
        "agentscope": {"enabled": False, "mode": "off"}, "retrieval": {"sufficiency_enabled": False, "expand_parent_payload": False},
        "grounding": {"evidence_strength_enabled": False, "field_binding_enabled": False, "field_binding_agent_enabled": False,
                      "slot_decomposition_enabled": False, "pre_writeback_consistency_enabled": False},
    })
    calls: list[str] = []

    def retrieve(query: str) -> Step15RetrievalResult:
        if forbidden_calls:
            pytest.fail("resume must reject or reuse checkpoints before any retrieval/model call")
        calls.append("retrieve")
        return deepcopy(pack) if pack is not None else Step15RetrievalResult([hit()], [hit()], "layered")

    def answer(**kwargs: Any) -> dict[str, Any]:
        if forbidden_calls:
            pytest.fail("resume must reject or reuse checkpoints before any model call")
        calls.append("answer")
        return {"answer_value": "500kVA", "answer_status": "answered", "confidence": 0.95,
                "source_chunk_ids": ["native-ups"], "reference_source_documents": [{
                    "chunk_id": "native-ups", "quote": "UPS容量 / 500kVA", "reason": "native evidence",
                }]}

    return Step15AgentRunner(config=config, target_namespace="room", global_namespace="global", out_dir=path,
                             collection_name=COLLECTION, index_scopes=scopes if scopes is not None else pins(),
                             retrieval_fn=retrieve, answer_caller=answer, chat_max_retries=0,
                             chat_retry_backoff_seconds=0, resume=resume), calls


def protected_files(path: Path) -> dict[str, bytes]:
    return {str(file.relative_to(path)): file.read_bytes() for file in path.rglob("*") if file.is_file()}


@pytest.mark.parametrize("part", ["reranked", "vector"])
@pytest.mark.parametrize("change", [{"index_version_id": V2}, {"knowledge_base_id": GLOBAL_KB}, {"namespace": "foreign"},
                                   {"index_version": V2}, {"source": {"index_version_id": V2}}])
@pytest.mark.parametrize("constrained", [False, True])
def test_injected_pack_entry_and_exit_must_match_pins_before_any_answer(tmp_path: Path, part: str, change: dict[str, Any], constrained: bool) -> None:
    bad, good = hit(**change), hit()
    pack = Step15RetrievalResult([bad] if part == "reranked" else [good], [bad] if part == "vector" else [good], "layered")
    agent, calls = runner(tmp_path, pack=pack)
    with pytest.raises(IndexScopeError):
        agent.retrieve("UPS容量", layer_names=["target_structured_fact"] if constrained else None)
    assert calls == ["retrieve"]


def test_pins_are_frozen_in_snapshot_manifest_and_typed_evidence(tmp_path: Path) -> None:
    agent, calls = runner(tmp_path)
    predictions = agent.run([item()])
    manifest = read_json(tmp_path / "run_manifest.json")
    snapshot = read_json(tmp_path / "form_input_snapshot.json")
    assert calls == ["retrieve", "answer"] and predictions[0].answer_status == "answered"
    assert manifest["index_scopes"] == snapshot["acquisition_contract"]["index_scope_contract"]["scopes"] == pins()
    assert predictions[0].evidence_refs[0].index_version == V1
    assert predictions[0].evidence_refs[0].qdrant_point_id == versioned_point_id(V1, "native-ups")
    assert validate_step15_artifacts(tmp_path)["valid"] is True


@pytest.mark.parametrize("damage", ["version", "kb", "source_version"])
def test_resume_authority_scope_tampering_fails_before_calls_or_artifact_changes(tmp_path: Path, damage: str) -> None:
    initial, _ = runner(tmp_path)
    initial.run([item()])
    authority = read_jsonl(tmp_path / "retrieval_evidence.checkpoint.jsonl")
    changes = {"version": {"index_version_id": V2}, "kb": {"knowledge_base_id": GLOBAL_KB},
               "source_version": {"source": {"index_version_id": V2}}}
    authority[0]["top_hits"][0].update(changes[damage])
    write_jsonl(tmp_path / "retrieval_evidence.checkpoint.jsonl", authority)
    before = protected_files(tmp_path)
    resumed, _ = runner(tmp_path, resume=True, forbidden_calls=True)
    with pytest.raises(IndexScopeError):
        resumed.run([item()])
    assert protected_files(tmp_path) == before


def test_changed_pins_reject_resume_before_reading_a_damaged_checkpoint(tmp_path: Path) -> None:
    initial, _ = runner(tmp_path)
    initial.run([item()])
    (tmp_path / "predictions.checkpoint.jsonl").write_text("{damaged json}\n", encoding="utf-8")
    changed = pins()
    changed[0]["index_version_id"] = V2
    before = protected_files(tmp_path)
    resumed, _ = runner(tmp_path, scopes=changed, resume=True, forbidden_calls=True)
    with pytest.raises(FormInputMismatchError):
        resumed.run([item()])
    assert protected_files(tmp_path) == before


def test_same_pins_resume_reuses_verified_authority_without_calls(tmp_path: Path) -> None:
    initial, _ = runner(tmp_path)
    first = initial.run([item()])[0]
    resumed, _ = runner(tmp_path, resume=True, forbidden_calls=True)
    assert resumed.run([item()])[0].to_dict() == first.to_dict()


def test_archival_validator_rejects_authority_from_another_version(tmp_path: Path) -> None:
    initial, _ = runner(tmp_path)
    initial.run([item()])
    rows = read_jsonl(tmp_path / "retrieval_evidence.jsonl")
    rows[0]["top_hits"][0]["index_version_id"] = V2
    write_jsonl(tmp_path / "retrieval_evidence.jsonl", rows)
    with pytest.raises(ArtifactValidationError, match="EV_INDEX_SCOPE_MISMATCH"):
        validate_step15_artifacts(tmp_path)


def test_runner_rejects_extra_namespace_pins(tmp_path: Path) -> None:
    extra = pins() + [{**pins()[0], "namespace": "foreign"}]
    with pytest.raises((IndexScopeError, ValueError), match="scope|namespace"):
        runner(tmp_path, scopes=extra)


def test_archival_validator_rejects_tampered_scope_artifact(tmp_path: Path) -> None:
    initial, _ = runner(tmp_path)
    initial.run([item()])
    changed = pins()
    changed[0]["index_version_id"] = V2
    write_json(tmp_path / "index_scopes.json", changed)
    with pytest.raises(ArtifactValidationError, match="EV_INDEX_SCOPE_MISMATCH"):
        validate_step15_artifacts(tmp_path)


def test_injected_retriever_must_carry_the_same_frozen_scopes(tmp_path: Path) -> None:
    config = load_app_config(project_root=tmp_path, default_config=tmp_path / "missing", env={},
                             cli_overrides={"agentscope": {"enabled": False, "mode": "off"}, "retrieval": {"sufficiency_enabled": False}})
    with pytest.raises(ValueError, match="injected retriever"):
        Step15AgentRunner(config=config, target_namespace="room", global_namespace="global", out_dir=tmp_path,
                          collection_name=COLLECTION, index_scopes=pins(), retriever=SimpleNamespace(index_scopes=None), reranker=object())


def test_rejected_injected_scope_never_reaches_answer_generation_or_archived_authority(tmp_path: Path) -> None:
    bad = hit(index_version_id=V2)
    agent, calls = runner(tmp_path, pack=Step15RetrievalResult([bad], [bad], "layered"))
    prediction = agent.run([item()])[0]
    assert calls == ["retrieve"] and prediction.method_name == "step15_agent_failed"
    assert not prediction.evidence_refs and not prediction.source_chunk_ids
    assert read_jsonl(tmp_path / "retrieval_evidence.jsonl")[0]["top_hits"] == []


@pytest.mark.parametrize("callback", ["answer", "field_binding"])
def test_callback_scope_mutation_cannot_authorize_real_writeback(tmp_path: Path, callback: str) -> None:
    template = tmp_path / "form.xlsx"
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.active["A2"] = "UPS容量"
    workbook.save(template)
    workbook.close()
    agent, _ = runner(tmp_path / "run")
    agent.template_path = template
    agent.writeback_enabled = True
    original_answer = agent.answer_caller

    def mutate_scope(**kwargs: Any) -> dict[str, Any]:
        for evidence in kwargs["hits"]:
            point = versioned_point_id(V2, evidence["chunk_id"])
            evidence.update(index_version_id=V2, index_version=V2, point_id=point, qdrant_point_id=point)
        if callback == "answer":
            assert original_answer is not None
            return original_answer(**kwargs)
        return {"passed": True, "label": "exact", "confidence": 0.99, "evidence_chunk_ids": ["native-ups"]}

    if callback == "answer":
        agent.answer_caller = mutate_scope
    else:
        agent.field_binding_agent_enabled = True
        agent.field_binding_judge_caller = mutate_scope
    prediction = agent.run([item()])[0]
    assert prediction.method_name == "step15_agent_failed"
    assert not prediction.evidence_refs and not prediction.source_chunk_ids
    assert read_jsonl(agent.out_dir / "retrieval_evidence.jsonl")[0]["top_hits"] == []
    filled = load_workbook(agent.out_dir / "filled_form.xlsx")
    try:
        assert filled["Sheet1"]["D2"].value is None
    finally:
        filled.close()
    assert validate_step15_artifacts(agent.out_dir)["valid"] is True


def test_writeback_rechecks_mutated_authority_before_touching_workbook(tmp_path: Path) -> None:
    template = tmp_path / "form.xlsx"
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.save(template)
    workbook.close()
    agent, _ = runner(tmp_path / "run")
    prediction = agent.run([item()])[0]
    agent.template_path = template
    agent.writeback_enabled = True
    authority = agent.retrieval_evidence_by_field_id[prediction.field_id]
    point = versioned_point_id(V2, "native-ups")
    authority[0].update(index_version_id=V2, index_version=V2, point_id=point, qdrant_point_id=point)
    prediction = attach_addressed_evidence(prediction, authority)
    assert prediction.evidence_refs[0].index_version == V2
    before = protected_files(tmp_path)
    with pytest.raises(IndexScopeError, match="version"):
        agent.maybe_writeback([prediction], [])
    assert protected_files(tmp_path) == before
    assert not (agent.out_dir / "filled_form.xlsx").exists()
