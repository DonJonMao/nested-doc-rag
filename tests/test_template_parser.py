from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook

from nested_doc_rag.form import TemplateParseError, parse_form_template, parse_template_to_run_dir, template_parser


def save_workbook(tmp_path: Path, rows: list[list], *, name: str = "新工勘.xlsx", sheet_name: str = "动力调研") -> Path:
    workbook = Workbook()
    workbook.active.title = sheet_name
    for row in rows:
        workbook.active.append(row)
    path = tmp_path / name
    workbook.save(path)
    workbook.close()
    return path


def test_explicit_headers_select_target_instead_of_last_column(tmp_path: Path) -> None:
    path = save_workbook(tmp_path, [
        ["类别", "实际情况", "指标名称", "填写说明", "备注"],
        ["供配电", "SECRET_MANUAL_ANSWER", "市电路数", "填写当前实际路数", "保留人工备注"],
    ])

    result = parse_form_template(path)
    item = result.items[0]

    assert item["file_name"] == path.name
    assert item["sheet_name"] == "动力调研"
    assert item["row_index"] == 2
    assert item["target_cell"] == "'动力调研'!B2"
    assert item["category_path"] == ["供配电"]
    assert item["question_text"] == "市电路数"
    assert item["instruction_text"] == "填写当前实际路数"
    assert item["answer_example"] is None
    assert "heldout_answer" not in item and "existing_value" not in item
    assert "SECRET_MANUAL_ANSWER" not in json.dumps(result.items, ensure_ascii=False)
    assert result.report["target_column_by_sheet"] == {"动力调研": "B"}
    assert result.report["template_sha256"] == "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()
    assert result.report["parser_name"] == template_parser.PARSER_NAME
    assert result.report["parser_version"] == template_parser.PARSER_VERSION


def test_heldout_answers_are_opt_in_and_never_used_by_retrieval_inputs(tmp_path: Path) -> None:
    path = save_workbook(tmp_path, [
        ["问题", "填写说明", "答复示例", "答案"],
        ["油机数量", "写数字", "例如2", "HUMAN_SECRET"],
        ["备用电源可用", None, None, False],
        ["备用油机数量", None, None, 0],
    ])

    without = parse_form_template(path)
    with_heldout = parse_form_template(path, include_heldout_answers=True)

    assert [item["heldout_answer"] for item in with_heldout.items] == ["HUMAN_SECRET", "False", "0"]
    assert [item["form_item_id"] for item in with_heldout.items] == [item["form_item_id"] for item in without.items]
    for item in with_heldout.items:
        safe_input = {key: item[key] for key in ("question_text", "instruction_text", "answer_example", "category_path", "suggested_retrieval_query")}
        assert "HUMAN_SECRET" not in json.dumps(safe_input)
        assert "existing_value" not in item


def test_multiple_sheets_same_row_have_distinct_ids_and_quoted_target_locations(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "电力系统"
    workbook.active.append(["指标名称", "实际值"])
    workbook.active.append(["容量", None])
    second = workbook.create_sheet("Room 'B'")
    second.append(["实际值", "填写说明", "指标名称", "备注"])
    second.append([None, "单位kW", "容量", None])
    workbook.create_sheet("空白参考")
    path = tmp_path / "多sheet.xlsx"
    workbook.save(path)
    workbook.close()

    result = parse_form_template(path)

    assert len(result.items) == 2
    assert len({item["form_item_id"] for item in result.items}) == 2
    assert result.items[0]["target_cell"] == "'电力系统'!B2"
    assert result.items[1]["target_cell"] == "'Room ''B'''!A2"
    assert result.report["sheet_count"] == 3
    assert result.report["sheets"][2]["reason"] == "no_explicit_form_header"


def test_ids_survive_file_rename_and_change_with_field_structure(tmp_path: Path) -> None:
    path = save_workbook(tmp_path, [["字段", "值"], ["机房名称", None]])
    renamed = tmp_path / "任意重命名.xlsx"
    shutil.copyfile(path, renamed)

    original = parse_form_template(path)
    renamed_result = parse_form_template(renamed)
    assert original.items[0]["form_item_id"] == renamed_result.items[0]["form_item_id"]
    assert renamed_result.items[0]["file_name"] == renamed.name

    reordered = save_workbook(tmp_path, [["值", "字段"], [None, "机房名称"]], name="重排列.xlsx")
    assert parse_form_template(reordered).items[0]["form_item_id"] != original.items[0]["form_item_id"]
    shifted = save_workbook(tmp_path, [["字段", "值"], [None, None], ["机房名称", None]], name="重排行.xlsx")
    assert parse_form_template(shifted).items[0]["form_item_id"] != original.items[0]["form_item_id"]


def test_merged_two_level_headers_and_categories_are_explained(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "供电"
    sheet.append(["类别", "指标名称", "机房信息", None])
    sheet.append([None, None, "示例", "实际值"])
    sheet.append(["市电", "路数", "2路", None])
    sheet.append([None, "供电容量", "10 MVA", None])
    sheet.merge_cells("A1:A2")
    sheet.merge_cells("B1:B2")
    sheet.merge_cells("C1:D1")
    sheet.merge_cells("A3:A4")
    path = tmp_path / "合并表头.xlsx"
    workbook.save(path)
    workbook.close()

    result = parse_form_template(path)

    assert len(result.items) == 2
    assert all(item["category_path"] == ["市电"] for item in result.items)
    assert result.items[0]["target_cell"] == "'供电'!D3"
    assert result.items[0]["answer_example"] == "2路"
    assert result.report["sheets"][0]["header_rows"] == [1, 2]
    assert result.report["sheets"][0]["column_roles"]["target"][0]["header_cell"] == "D2"


def test_category_only_rows_and_parent_changes_update_inherited_path(tmp_path: Path) -> None:
    path = save_workbook(tmp_path, [
        ["类别", "子类", "问题", "答案"],
        ["动力", "市电", None, None],
        [None, None, "路数", None],
        [None, "油机", "数量", None],
        ["制冷", None, "空调类型", None],
    ])

    result = parse_form_template(path)

    assert [item["category_path"] for item in result.items] == [["动力", "市电"], ["动力", "油机"], ["制冷"]]
    assert {row["row_index"] for row in result.report["skipped_rows"]} == {2}


def test_simple_title_does_not_turn_into_a_broad_target_header(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "机房信息"
    sheet.merge_cells("A1:D2")
    for column, value in enumerate(["类别", "指标名称", "实际情况", "备注"], 1):
        sheet.cell(3, column, value)
    for column, value in enumerate(["动力", "供电路数", None, None], 1):
        sheet.cell(4, column, value)
    path = tmp_path / "有标题.xlsx"
    workbook.save(path)
    workbook.close()

    result = parse_form_template(path)

    assert len(result.items) == 1
    assert result.items[0]["target_cell"] == "'Sheet'!C4"
    assert result.report["sheets"][0]["header_rows"] == [3]


@pytest.mark.parametrize(("headers", "reason"), [
    (["字段", "问题", "实际情况"], "ambiguous_question_columns"),
    (["字段", "实际值", "答案"], "ambiguous_target_columns"),
    (["字段", "未知列"], "missing_target_column"),
    (["未知列", "答案"], "missing_question_column"),
])
def test_missing_or_multiple_required_columns_stop_with_report(tmp_path: Path, headers: list[str], reason: str) -> None:
    path = save_workbook(tmp_path, [headers, ["机房名称", None, None][:len(headers)]])

    with pytest.raises(TemplateParseError) as error:
        parse_form_template(path)

    assert error.value.report["status"] == "error"
    assert reason in [issue["reason"] for issue in error.value.report["errors"]]
    assert error.value.report["ambiguous_rows"][0]["sheet_name"] == "动力调研"


def test_formula_targets_and_empty_questions_are_skipped(tmp_path: Path) -> None:
    path = save_workbook(tmp_path, [
        ["字段", "答案"],
        ["计算数量", "=SUM(B4:B5)"],
        [None, "PRIVATE_EMPTY_QUESTION_VALUE"],
        ["油机数量", None],
        ["铭牌照片", '=_xlfn.DISPIMG("ID_PHOTO",1)'],
        ["=B4", None],
    ])

    result = parse_form_template(path, include_heldout_answers=True)

    assert len(result.items) == 1
    assert result.items[0]["question_text"] == "油机数量"
    assert result.items[0]["target_cell"] == "'动力调研'!B4"
    assert [(row["row_index"], row["reason"]) for row in result.report["skipped_rows"]] == [
        (2, "formula_target"), (3, "empty_question"), (5, "formula_target"), (6, "formula_question"),
    ]
    assert "PRIVATE_EMPTY_QUESTION_VALUE" not in json.dumps(result.report)


def test_merged_horizontal_target_is_one_explicit_writable_range(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["字段", "实际情况", None])
    sheet.append(["机房名称", None, None])
    sheet.merge_cells("B1:C1")
    sheet.merge_cells("B2:C2")
    path = tmp_path / "合并目标.xlsx"
    workbook.save(path)
    workbook.close()

    result = parse_form_template(path)

    assert result.report["target_column_by_sheet"] == {"Sheet": "B:C"}
    assert result.items[0]["target_cell"] == "'Sheet'!B2"


def test_merged_target_header_cannot_choose_between_unmerged_cells(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["字段", "实际情况", None])
    sheet.append(["机房名称", None, None])
    sheet.merge_cells("B1:C1")
    path = tmp_path / "歧义合并目标.xlsx"
    workbook.save(path)
    workbook.close()

    with pytest.raises(TemplateParseError) as error:
        parse_form_template(path)

    assert error.value.report["errors"][0]["reason"] == "unmerged_header_range"
    assert error.value.report["errors"][0]["role"] == "target"


def test_shared_vertical_target_stops_instead_of_overwriting_two_questions(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["字段", "答案"])
    sheet.append(["油机数量", None])
    sheet.append(["UPS容量", None])
    sheet.merge_cells("B2:B3")
    path = tmp_path / "共享目标.xlsx"
    workbook.save(path)
    workbook.close()

    with pytest.raises(TemplateParseError) as error:
        parse_form_template(path)

    assert "shared_merged_target" in [issue["reason"] for issue in error.value.report["errors"]]


def test_formula_instruction_and_example_do_not_become_prompt_content(tmp_path: Path) -> None:
    path = save_workbook(tmp_path, [
        ["字段", "填写说明", "示例", "答案"],
        ["油机数量", "=D2", "=D2", "SECRET_MANUAL_VALUE"],
    ])

    result = parse_form_template(path)

    assert result.items[0]["instruction_text"] == ""
    assert result.items[0]["answer_example"] is None
    assert "SECRET_MANUAL_VALUE" not in json.dumps(result.items)
    assert len(result.report["warnings"]) == 2


def test_parser_reads_and_hashes_the_same_bytes_once(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = save_workbook(tmp_path, [["字段", "答案"], ["原问题", None]])
    other = save_workbook(tmp_path, [["字段", "答案"], ["变更问题", None]], name="other.xlsx")
    original = path.read_bytes()
    replacement = other.read_bytes()
    original_read = Path.read_bytes

    def changing_read(source: Path) -> bytes:
        data = original_read(source)
        if source == path:
            path.write_bytes(replacement)
        return data

    monkeypatch.setattr(Path, "read_bytes", changing_read)

    result = parse_form_template(path)

    assert result.items[0]["question_text"] == "原问题"
    assert result.report["template_sha256"] == "sha256:" + hashlib.sha256(original).hexdigest()


def test_standalone_writer_creates_parse_diagnostics_and_runtime_items(tmp_path: Path) -> None:
    path = save_workbook(tmp_path, [["字段", "答案"], ["机房名称", None]])
    out_dir = tmp_path / "run"

    items_path = parse_template_to_run_dir(path, out_dir)

    assert items_path == out_dir / "form_items.jsonl"
    assert json.loads(items_path.read_text())["question_text"] == "机房名称"
    report = json.loads((out_dir / "form_parse_report.json").read_text())
    assert report["status"] == "parsed"
    assert report["detected_fields"] == 1


def test_parser_error_writer_retains_report_without_producing_fields(tmp_path: Path) -> None:
    path = save_workbook(tmp_path, [["字段", "答案", "实际值"], ["机房名称", None, None]])
    out_dir = tmp_path / "error-run"

    with pytest.raises(TemplateParseError):
        parse_template_to_run_dir(path, out_dir)

    assert not (out_dir / "form_items.jsonl").exists()
    assert json.loads((out_dir / "form_parse_report.json").read_text())["status"] == "error"


def test_repeated_identical_header_is_skipped(tmp_path: Path) -> None:
    path = save_workbook(tmp_path, [
        ["类别", "字段", "答案"], ["动力", "机房名称", None],
        ["类别", "字段", "答案"], [None, "机房地址", None],
    ])

    result = parse_form_template(path)

    assert [item["question_text"] for item in result.items] == ["机房名称", "机房地址"]
    assert result.items[1]["category_path"] == []
    assert result.report["skipped_rows"][0]["reason"] == "repeated_header"


def test_changed_header_layout_does_not_keep_writing_the_previous_column(tmp_path: Path) -> None:
    path = save_workbook(tmp_path, [
        ["字段", "答案", "备注"], ["机房名称", None, None],
        ["字段", "备注", "答案"], ["机房地址", None, None],
    ])

    with pytest.raises(TemplateParseError) as error:
        parse_form_template(path)

    assert "inconsistent_header_layout" in [issue["reason"] for issue in error.value.report["errors"]]


def test_header_can_follow_many_blank_title_rows(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "现场填写须知"
    sheet["A40"] = "字段"
    sheet["C40"] = "实际情况"
    sheet["A41"] = "机房名称"
    path = tmp_path / "长前言.xlsx"
    workbook.save(path)
    workbook.close()

    result = parse_form_template(path)

    assert result.items[0]["row_index"] == 41
    assert result.items[0]["target_cell"] == "'Sheet'!C41"


def test_original_141_field_template_remains_readable() -> None:
    project = Path(__file__).resolve().parents[1]
    relative = Path("data/工勘单/基地云机房信息调研表.xlsx")
    template = next((candidate for candidate in (project / relative, project.parent / "datacenter" / relative) if candidate.exists()), None)
    if template is None:
        pytest.skip("historical workbook fixture is not shipped in this checkout")

    before = hashlib.sha256(template.read_bytes()).hexdigest()
    result = parse_form_template(template)

    assert len(result.items) == result.report["detected_fields"] == 141
    assert [item["row_index"] for item in result.items] == list(range(4, 145))
    assert {item["target_cell"].split("!")[-1] for item in result.items} == {f"G{row}" for row in range(4, 145)}
    assert len({item["form_item_id"] for item in result.items}) == 141
    assert result.items[0]["question_text"] == "机房名称"
    assert not any("existing_value" in item or "heldout_answer" in item for item in result.items)
    assert hashlib.sha256(template.read_bytes()).hexdigest() == before
