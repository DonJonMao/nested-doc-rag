from __future__ import annotations

import hashlib
from dataclasses import replace
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from openpyxl import Workbook
from qdrant_client import QdrantClient, models

from nested_doc_rag.agent.mas.roles import EvidenceRetrievalRole, QueryReplannerRole
from nested_doc_rag.agent.mas.schemas import ActionType, AgenticMASState, EvidenceAction
from nested_doc_rag.agent.step15_runner import Step15AgentRunner
from nested_doc_rag.config import AgentScopeConfig, load_app_config
from nested_doc_rag.evaluation.step15_engine import Step15RetrievalResult, run_step15_retrieval
from nested_doc_rag.grounding.evidence_strength import is_exact_structured_hit
from nested_doc_rag.grounding.provenance import locate_quote
from nested_doc_rag.ingestion import build_ingestion_records
from nested_doc_rag.retrieval.layered import constrain_layered_plan
from nested_doc_rag.retrieval.qdrant_retriever import QdrantRetriever

PROJECT_ROOT = Path(__file__).resolve().parents[1]
UPLOADED_TYPES = {"uploaded_excel_row", "uploaded_docx_paragraph", "uploaded_docx_table_row", "uploaded_text_chunk"}


class FixedEmbedding:
    def embed_query(self, query: str) -> list[float]:
        return [1.0, 0.0, 0.0]


class FixedReranker:
    def rerank(self, query: str, docs: list[str], *, top_n: int) -> list[dict[str, Any]]:
        return [{"index": i, "relevance_score": 1.0} for i in range(min(top_n, len(docs)))]


def write_uploads(directory: Path) -> None:
    directory.mkdir()
    workbook = Workbook()
    workbook.active.append(["UPS 单台容量", "500 kVA"])
    workbook.save(directory / "power.xlsx")
    document = Document()
    document.add_paragraph("301机房\n市电来源：东侧变电站。")
    row = document.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "UPS 单台容量"
    row.cells[1].text = "500 kVA"
    document.save(directory / "power.docx")
    (directory / "power.txt").write_text("301机房\n  供电：2路 ⚡\n", encoding="utf-8")


@pytest.fixture
def local_retriever(tmp_path: Path):
    uploads = tmp_path / "uploads"
    write_uploads(uploads)
    records, skipped = build_ingestion_records(uploads, namespace="room_301", knowledge_base_id="kb_301")
    assert skipped == []
    other_records, _ = build_ingestion_records(uploads, namespace="room_401", knowledge_base_id="kb_401")
    global_records, _ = build_ingestion_records(uploads, namespace="global", knowledge_base_id="kb_global")
    historical = {
        "chunk_id": "historical_main",
        "namespace": "room_301",
        "source_type": "main_excel_capability",
        "corpus_layer": "fact",
        "raw_text": "301机房 UPS 单台容量：500 kVA",
    }
    client = QdrantClient(":memory:")
    client.create_collection("test", vectors_config=models.VectorParams(size=3, distance=models.Distance.COSINE))
    payloads = [historical, *records, *other_records, *global_records]
    client.upsert("test", points=[models.PointStruct(id=i, vector=[1.0, 0.0, 0.0], payload=p) for i, p in enumerate(payloads)])
    retriever = QdrantRetriever.__new__(QdrantRetriever)
    retriever.client = client
    retriever.collection_name = "test"
    retriever.embedder = FixedEmbedding()
    try:
        yield retriever
    finally:
        client.close()


@pytest.mark.parametrize("profile", [None, "default.yaml", "docker.yaml", "local.example.yaml"])
def test_native_uploads_are_retrieved_through_real_qdrant_filters(tmp_path: Path, local_retriever, profile: str | None) -> None:
    config = load_app_config(project_root=tmp_path, default_config=PROJECT_ROOT / "config" / profile if profile else tmp_path / "missing")
    result = run_step15_retrieval(
        "301机房供电情况",
        retriever=local_retriever,
        reranker=FixedReranker(),
        target_namespace="room_301",
        global_namespace="global",
        allowed_layers=config.retrieval.query_layers,
        retrieval_mode="layered",
        vector_top_k=40,
        rerank_top_n=10,
        layered_plan=config.retrieval.layered_plan,
    )
    assert {hit["source_type"] for hit in result.vector_hits if hit["namespace"] == "room_301" and hit["source_type"] in UPLOADED_TYPES} == UPLOADED_TYPES
    assert {hit["source_type"] for hit in result.vector_hits if hit["retrieval_layer"] == "global_detail" and hit["source_type"] in UPLOADED_TYPES} == UPLOADED_TYPES
    assert {hit["namespace"] for hit in result.vector_hits} == {"room_301", "global"}
    assert result.reranked_hits[0]["retrieval_layer"] == "target_structured_fact"
    historical = next(hit for hit in result.reranked_hits if hit["chunk_id"] == "historical_main")
    assert historical["retrieval_layer"] == "target_structured_fact"
    assert historical["source_type"] == "main_excel_capability"
    uploaded = [hit for hit in result.vector_hits if hit["source_type"] in UPLOADED_TYPES]
    assert all(not is_exact_structured_hit(hit, "room_301") for hit in uploaded if hit["evidence_kind"] != "structured_field")
    structured = [hit for hit in uploaded if hit["namespace"] == "room_301" and hit["evidence_kind"] == "structured_field"]
    assert structured and all(hit["retrieval_layer"] == "target_structured_fact" for hit in structured)
    assert all(is_exact_structured_hit(hit, "room_301") for hit in structured)
    for hit in uploaded:
        assert hit["raw_source_text"] == hit["raw_text"]
        assert hit["source_text_hash"] == "sha256:" + hashlib.sha256(hit["raw_source_text"].encode("utf-8")).hexdigest()
        assert hit["document_id"] == hit["source"]["document_id"]


def test_action_constraints_reach_real_qdrant_and_preserve_namespace(tmp_path: Path, local_retriever) -> None:
    config = load_app_config(project_root=tmp_path, default_config=tmp_path / "missing")
    config = replace(config, agentscope=AgentScopeConfig(enabled=False, mode="off"))
    runner = Step15AgentRunner(
        config=config,
        target_namespace="room_301",
        global_namespace="global",
        out_dir=tmp_path / "run",
        retriever=local_retriever,
        reranker=FixedReranker(),
    )
    result = EvidenceRetrievalRole(runner).run_action(
        EvidenceAction(
            action_type=ActionType.SOURCE_SPECIFIC_RETRIEVAL,
            query_text="UPS 单台容量",
            target_layer="target_table_detail",
            source_type_preference="uploaded_docx_table_row",
        )
    )
    assert len(result.top_hits) == 1
    assert result.top_hits[0]["source_type"] == "uploaded_docx_table_row"
    assert result.top_hits[0]["retrieval_layer"] == "target_table_detail"
    assert result.top_hits[0]["namespace"] == "room_301"
    assert len(result.vector_hits) == 1


def test_invalid_or_empty_constraints_do_not_trigger_unfiltered_retrieval(tmp_path: Path) -> None:
    calls: list[str] = []
    config = load_app_config(project_root=tmp_path, default_config=tmp_path / "missing")
    runner = Step15AgentRunner(
        config=replace(config, agentscope=AgentScopeConfig(enabled=False, mode="off")),
        target_namespace="room_301",
        global_namespace="global",
        out_dir=tmp_path,
        retrieval_fn=lambda query: calls.append(query) or Step15RetrievalResult([], [], "layered"),
    )
    rejected_layer = runner.retrieve("query", layer_names=["other_room"])
    rejected_source = runner.retrieve("query", source_types=["arbitrary_source"])
    assert rejected_layer.vector_hits == rejected_source.reranked_hits == []
    assert rejected_layer.metadata["constraint_rejected"] is True
    assert rejected_source.metadata["constraint_rejected"] is True
    assert runner.trace.events[-1].step == "retrieval_constraint_rejected"
    assert runner.retrieve("query", source_types=[]).vector_hits == []
    assert calls == []
    with pytest.raises(ValueError, match="unconfigured layer"):
        constrain_layered_plan(config.retrieval.layered_plan, layer_names=["other_room"])
    original_sources = list(config.retrieval.layered_plan[0]["source_types"])
    assert constrain_layered_plan(config.retrieval.layered_plan, source_types=["uploaded_text_chunk"])
    assert config.retrieval.layered_plan[0]["source_types"] == original_sources
    invalid_namespace_plan = [dict(config.retrieval.layered_plan[0], namespaces="another_room")]
    with pytest.raises(ValueError, match="namespace must be target or global"):
        constrain_layered_plan(invalid_namespace_plan)
    runner.layered_plan = invalid_namespace_plan
    assert runner.retrieve("query").metadata["constraint_rejected"] is True
    assert calls == []


def test_native_text_whitespace_zero_and_long_chunks_are_preserved(tmp_path: Path) -> None:
    text = "机房 ⚡\n  供电：2路\n" + "温湿度记录\n" * 400
    (tmp_path / "power.txt").write_text(text, encoding="utf-8")
    workbook = Workbook()
    workbook.active.append([0])
    workbook.save(tmp_path / "zero.xlsx")
    records, skipped = build_ingestion_records(tmp_path, namespace="room_301", knowledge_base_id="kb")
    assert skipped == []
    text_records = [record for record in records if record["source_type"] == "uploaded_text_chunk"]
    assert "".join(record["raw_source_text"] for record in text_records) == text
    zero = next(record for record in records if record["source_type"] == "uploaded_excel_row")
    assert zero["raw_source_text"] == "0"
    assert zero["source"]["cell_range"] == "A1:A1"


def test_text_ingestion_preserves_crlf_code_points(tmp_path: Path) -> None:
    original = "甲\r\n  双路市电 ⚡\r\n"
    (tmp_path / "power.txt").write_bytes(original.encode("utf-8"))
    records, skipped = build_ingestion_records(tmp_path, namespace="room_301", knowledge_base_id="kb")
    assert skipped == []
    assert records[0]["raw_source_text"] == original
    assert records[0]["source_text_hash"] == "sha256:" + hashlib.sha256(original.encode("utf-8")).hexdigest()


def test_qdrant_preserves_declared_hash_space_for_legacy_text(local_retriever) -> None:
    local_retriever.client.upsert("test", points=[models.PointStruct(id=999, vector=[1.0, 0.0, 0.0], payload={
        "chunk_id": "legacy_corrupted", "namespace": "room_301", "source_type": "uploaded_text_chunk",
        "corpus_layer": "fact", "raw_text": "A😀引用B", "source_text_hash": "sha256:corrupted", "source_text_hash_space": "raw_text",
    })])
    hits = local_retriever.search("引用", namespaces=["room_301"], layers=["fact"], top_k=20)
    legacy = next(hit for hit in hits if hit["chunk_id"] == "legacy_corrupted")
    provenance = locate_quote(legacy, "引用", chunk_id=legacy["chunk_id"])
    assert provenance["match_status"] == "unavailable"
    assert provenance["reason"] == "source_hash_mismatch"


def test_configured_wildcard_matches_real_and_injected_backends(tmp_path: Path, local_retriever) -> None:
    config = load_app_config(project_root=tmp_path, default_config=tmp_path / "missing")
    plan = [{"layer_name": "target_custom", "namespaces": "target", "corpus_layers": ["fact"], "source_types": [], "description": "custom", "vector_top_k": 20, "rerank_top_n": 20}]
    config = replace(config, agentscope=AgentScopeConfig(enabled=False, mode="off"), retrieval=replace(config.retrieval, layered_plan=plan))
    real_runner = Step15AgentRunner(config=config, target_namespace="room_301", global_namespace="global", out_dir=tmp_path / "real", retriever=local_retriever, reranker=FixedReranker())
    baseline = real_runner.retrieve("power")
    injected_runner = Step15AgentRunner(config=config, target_namespace="room_301", global_namespace="global", out_dir=tmp_path / "injected", retrieval_fn=lambda _: baseline)
    real = real_runner.retrieve("power", layer_names=["target_custom"])
    injected = injected_runner.retrieve("power", layer_names=["target_custom"])
    assert len(real.reranked_hits) == len(injected.reranked_hits) == 5
    assert {hit["chunk_id"] for hit in real.reranked_hits} == {hit["chunk_id"] for hit in injected.reranked_hits}
    for runner in (real_runner, injected_runner):
        narrowed = runner.retrieve("power", layer_names=["target_custom"], source_types=["uploaded_docx_table_row"])
        assert len(narrowed.reranked_hits) == 1
        assert narrowed.reranked_hits[0]["source_type"] == "uploaded_docx_table_row"
        assert runner.retrieve("power", source_types=[]).reranked_hits == []
        state = AgenticMASState(item={"question_text": "power"}, base_query="power", current_query="power", round_index=0, evidence=[], vector_hits=[], generated=None, prediction=None, diagnosis=None)
        actions = QueryReplannerRole(runner).run_not_found_recovery(state)
        action = next(action for action in actions if action.action_type == ActionType.SOURCE_SPECIFIC_RETRIEVAL)
        built_in = EvidenceRetrievalRole(runner).run_action(action)
        assert len(built_in.top_hits) == 5
        assert all(hit["namespace"] == "room_301" for hit in built_in.top_hits)
