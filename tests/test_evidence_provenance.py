from __future__ import annotations

import hashlib
import json
from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from nested_doc_rag.agent.step15_runner import Step15AgentRunner
from nested_doc_rag.config import load_app_config
from nested_doc_rag.evaluation.step15_engine import Step15RetrievalResult, build_qdrant_answer_messages
from nested_doc_rag.form.input_snapshot import build_form_input_snapshot, persist_form_input_snapshot
from nested_doc_rag.grounding.provenance import build_field_evidence, finalize_evidence, locate_quote
from nested_doc_rag.ingestion import build_ingestion_records
from nested_doc_rag.io import read_jsonl, write_jsonl
from nested_doc_rag.schemas.eval import FieldPrediction


def test_exact_offsets_are_unicode_codepoints_and_hash_is_utf8() -> None:
    text = "设备😀：双路市电\n容量10kVA。"
    quote = "双路市电\n容量10kVA"
    result = locate_quote({"raw_source_text": text, "index_version": "kb-v3"}, quote, chunk_id="c1")

    assert result["match_status"] == "exact"
    assert result["start"] == 4
    assert result["end"] == 4 + len(quote)
    assert text[result["start"] : result["end"]] == quote
    assert result["source_text_hash"] == hashlib.sha256(text.encode("utf-8")).hexdigest()
    assert result["text_space"] == "raw_source_text"
    assert result["index_version"] == "kb-v3"


@pytest.mark.parametrize("text,quote", [("两路，两路", "两路"), ("aaaa", "aa")])
def test_repeated_including_overlapping_quotes_are_ambiguous(text: str, quote: str) -> None:
    result = locate_quote({"raw_text": text}, quote, chunk_id="c1")
    assert result["match_status"] == "ambiguous"
    assert result["reason"] == "quote_not_unique"
    assert result["start"] is None and result["end"] is None


def test_matching_does_not_normalize_spaces_or_unicode() -> None:
    assert locate_quote({"raw_text": "2路市电"}, "2 路市电", chunk_id="c1")["match_status"] == "unmatched"
    assert locate_quote({"raw_text": "cafe\u0301"}, "café", chunk_id="c1")["match_status"] == "unmatched"


def test_source_space_precedence_and_embedding_is_never_source() -> None:
    result = locate_quote({"raw_source_text": "原文", "raw_text": "合成前缀：原文"}, "合成前缀", chunk_id="c1")
    assert result["match_status"] == "unmatched"
    assert result["source_text"] == "原文"
    assert result["text_space"] == "raw_source_text"
    fallback = locate_quote({"raw_text": "合成前缀：原文"}, "原文", chunk_id="c1")
    assert fallback["text_space"] == "raw_text"
    assert fallback["index_version"] == "unknown"
    unavailable = locate_quote({"text_for_embedding": "原文"}, "原文", chunk_id="c1")
    assert unavailable["reason"] == "missing_source_text"
    assert unavailable["source_text"] == "" and unavailable["source_text_hash"] == ""


def test_declared_hash_corruption_is_unavailable_without_changing_text() -> None:
    text = "双路市电😀"
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
    hit = {"raw_source_text": text, "source_text_hash": f"sha256:{digest}", "source": {"index_version_id": "snapshot-7"}}
    assert locate_quote(hit, "双路市电", chunk_id="c1")["match_status"] == "exact"
    corrupted = locate_quote({**hit, "source_text_hash": "sha256:" + "0" * 64}, "双路市电", chunk_id="c1")
    assert corrupted["match_status"] == "unavailable"
    assert corrupted["reason"] == "source_hash_mismatch"
    assert corrupted["source_text"] == text
    assert corrupted["source_text_hash"] == digest
    assert corrupted["index_version"] == "snapshot-7"
    assert corrupted["start"] is None and corrupted["end"] is None
    # A native-text hash must not be applied to a legacy synthetic raw_text.
    fallback = locate_quote({"raw_text": "前缀：" + text, "source_text_hash": f"sha256:{digest}"}, text, chunk_id="c1")
    assert fallback["match_status"] == "exact"
    explicitly_same_space = locate_quote({"raw_text": text, "source_text_hash": "0" * 64, "source_text_hash_space": "raw_text"}, text, chunk_id="c1")
    assert explicitly_same_space["reason"] == "source_hash_mismatch"


@pytest.mark.parametrize(
    "hit,quote,chunk_id,reason",
    [(None, "quote", "missing", "chunk_not_retrieved"), ({"raw_text": "quote"}, "quote", "", "missing_chunk_id"),
     ({"raw_text": "quote"}, "", "c1", "missing_quote"), ({"raw_text": ""}, "quote", "c1", "missing_source_text")],
)
def test_unavailable_reasons(hit, quote: str, chunk_id: str, reason: str) -> None:
    result = locate_quote(hit, quote, chunk_id=chunk_id)
    assert result["match_status"] == "unavailable"
    assert result["reason"] == reason
    assert result["start"] is None and result["end"] is None


def test_authoritative_sources_and_no_manufactured_quote_or_gate_change() -> None:
    prediction = FieldPrediction(field_id="f1", row_index=4, target_cell="D4", answer_value="2路", source_chunk_ids=["c1"])
    hit = make_hit()
    overlay = {"writeback_allowed": False, "critic_flags": ["scope_mismatch"], "reasons": ["scope_mismatch"]}
    generated = {"reference_source_documents": [{"chunk_id": "c1", "quote": "2路市电", "file_name": "forged.xlsx", "namespace": "wrong"}]}
    before = deepcopy((prediction.to_dict(), overlay, generated, hit))
    field = build_field_evidence(item=make_item(4), prediction=prediction, generated=generated, top_hits=[hit], overlay=overlay)
    ref = field["evidence_refs"][0]
    assert ref["file_name"] == "真实能力表.xlsx"
    assert ref["namespace"] == "room301"
    assert ref["provenance"]["match_status"] == "exact"
    assert field["writeback_allowed"] is False
    assert field["writeback_status"] == "flagged"
    assert (prediction.to_dict(), overlay, generated, hit) == before

    no_quote = build_field_evidence(item=make_item(4), prediction=prediction, generated={}, top_hits=[hit], overlay=overlay)
    no_quote_ref = no_quote["evidence_refs"][0]
    assert no_quote_ref["text_preview"] == hit["raw_source_text"]
    assert no_quote_ref["provenance"]["quote"] == ""
    assert no_quote_ref["provenance"]["reason"] == "missing_quote"


def test_model_text_preview_is_a_candidate_but_invented_chunk_is_unavailable() -> None:
    prediction = FieldPrediction(field_id="f1", row_index=4, target_cell="D4", answer_value="未找到", answer_status="partial_clue")
    generated = {"reference_source_documents": [{"chunk_id": "c1", "text_preview": "2路市电"}, {"chunk_id": "invented", "quote": "2路市电", "file_name": "fake"}]}
    field = build_field_evidence(item=make_item(4), prediction=prediction, generated=generated, top_hits=[make_hit()], overlay=None)
    assert field["evidence_refs"][0]["provenance"]["match_status"] == "exact"
    missing = field["evidence_refs"][1]
    assert missing["file_name"] == ""
    assert missing["provenance"]["reason"] == "chunk_not_retrieved"


def test_audit_precedence_and_summary_count_references() -> None:
    prediction = FieldPrediction(field_id="f1", row_index=4, target_cell="D4", answer_value="2路", source_chunk_ids=["c1"])
    field = build_field_evidence(item=make_item(4), prediction=prediction, generated={"reference_source_documents": [{"chunk_id": "c1", "quote": "2路市电"}, {"chunk_id": "c1", "quote": "不存在"}]}, top_hits=[make_hit()], overlay={"writeback_allowed": True})
    evidence = finalize_evidence([field], audit_records=[{"field_id": "f1", "status": "flagged", "writeback_action": "skipped_formula", "reason": "skipped_formula"}], writeback_status="completed")
    assert evidence["summary"] == {"exact": 1, "ambiguous": 0, "unmatched": 1, "unavailable": 0}
    assert evidence["fields"][0]["writeback_status"] == "flagged"
    assert evidence["fields"][0]["writeback_action"] == "skipped_formula"
    assert field["writeback_status"] == "confirmed"
    no_writeback = finalize_evidence([field], audit_records=[], writeback_status="skipped: writeback disabled")
    assert no_writeback["fields"][0]["writeback_action"] == "review_only"


@pytest.mark.parametrize("kind", ["excel", "word_paragraph", "word_table"])
def test_display_physical_address_comes_from_native_hit_not_model(tmp_path: Path, kind: str) -> None:
    uploads = tmp_path / "uploads"
    uploads.mkdir()
    if kind == "excel":
        workbook = Workbook()
        workbook.active.title = "南501"
        workbook.active.append(["字段", "当前实际值"])
        workbook.active.append(["UPS容量", "500kVA"])
        workbook.save(uploads / "原生参数.xlsx")
        workbook.close()
        expected = {"file_name": "原生参数.xlsx", "sheet_name": "南501", "cell_range": "A2:B2", "row_index": 2}
    else:
        document = Document()
        document.add_paragraph("南501机房当前UPS容量为500kVA。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "设备", "当前容量"
        table.cell(1, 0).text, table.cell(1, 1).text = "UPS", "500kVA"
        document.save(uploads / "原生巡检.docx")
        expected = {"file_name": "原生巡检.docx", **({"paragraph_index": 1} if kind == "word_paragraph" else {"table_index": 1, "row_index": 2})}
    records, skipped = build_ingestion_records(input_dir=uploads, namespace="room301", knowledge_base_id="native-display-kb")
    assert not skipped
    hit = next(record for record in records if "500kVA" in record["raw_source_text"]
               and record["evidence_kind"] == {"excel": "structured_field", "word_paragraph": "paragraph", "word_table": "table_row"}[kind])
    prediction = FieldPrediction(field_id="f1", row_index=4, target_cell="D4", answer_value="500kVA", source_chunk_ids=[hit["chunk_id"]])
    forged = {"chunk_id": hit["chunk_id"], "quote": "500kVA", "file_name": "forged.docx", "sheet_name": "伪造",
              "cell_range": "Z99:ZZ99", "row_index": 99, "column_index": 99, "table_index": 99, "paragraph_index": 99,
              "source_anchor": "model invented location", "qdrant_point_id": "model invented point"}
    before = deepcopy(hit)
    field = build_field_evidence(item=make_item(4), prediction=prediction, generated={"reference_source_documents": [forged]}, top_hits=[hit], overlay=None)
    ref = field["evidence_refs"][0]
    for key, value in expected.items():
        assert ref[key] == value
    assert ref["namespace"] == "room301"
    assert ref["relative_path"] == hit["relative_path"]
    assert ref["qdrant_point_id"] == hit["point_id"]
    assert ref["source_anchor"] == hit["anchor"]
    assert ref["provenance"]["source_text"] == hit["raw_source_text"]
    assert ref["provenance"]["match_status"] == "exact"
    assert hit == before


@pytest.mark.parametrize("metadata_space", ["hit", "source", "address"])
def test_display_preserves_recorded_zero_indices_without_model_fallback(metadata_space: str) -> None:
    # Zero may occur in historical metadata. This is a presence/transport check,
    # not a claim that zero is a valid native Excel row or Word paragraph.
    hit = make_hit()
    dimensions = {"row_index": 0, "column_index": 0, "table_index": 0, "paragraph_index": 0, "page": 0}
    if metadata_space == "hit":
        hit.update(dimensions)
        hit["source"].update({key: 99 for key in dimensions})
    else:
        hit[metadata_space] = {**hit.get(metadata_space, {}), **dimensions}
    prediction = FieldPrediction(field_id="f1", row_index=4, target_cell="D4", answer_value="2路", source_chunk_ids=["c1"])
    generated = {"reference_source_documents": [{"chunk_id": "c1", "quote": "2路市电", **{key: 999 for key in dimensions}}]}
    field = build_field_evidence(item=make_item(4), prediction=prediction, generated=generated, top_hits=[hit], overlay=None)
    assert {key: field["evidence_refs"][0][key] for key in dimensions} == dimensions


def test_display_acquisition_uses_validation_only_and_copies_nested_trace() -> None:
    acquisition = {"strategy": "sufficiency_guided", "acquisition_rounds": 2, "qdrant_query_calls": None,
                   "rounds": [{"retrieval_round": 1, "missing_facts": ["柴油储量"], "hit_count": 2, "evidence_gain": 0}],
                   "final_sufficiency": {"sufficient": False, "missing_facts": ["柴油储量"], "reason": "仍不足"}}
    prediction = FieldPrediction(field_id="f1", row_index=4, target_cell="D4", answer_value="未找到", validation={"acquisition": acquisition})
    generated = {"reference_source_documents": [{"chunk_id": "c1", "quote": "2路市电"}], "acquisition": {"acquisition_rounds": 999}}
    overlay = {"writeback_allowed": False, "acquisition": {"final_sufficiency": {"sufficient": True}}}
    field = build_field_evidence(item=make_item(4), prediction=prediction, generated=generated, top_hits=[make_hit()], overlay=overlay)
    assert field["acquisition"] == acquisition
    assert field["acquisition"] is not acquisition
    field["acquisition"]["rounds"][0]["missing_facts"].append("显示侧改动")
    field["acquisition"]["final_sufficiency"]["sufficient"] = True
    assert acquisition["rounds"][0]["missing_facts"] == ["柴油储量"]
    assert acquisition["final_sufficiency"]["sufficient"] is False
    acquisition["rounds"][0]["evidence_gain"] = 123
    assert field["acquisition"]["rounds"][0]["evidence_gain"] == 0


@pytest.mark.parametrize("invalid_trace", [None, [], "model trace"])
def test_display_does_not_invent_trace_for_historical_or_invalid_validation(invalid_trace) -> None:
    prediction = FieldPrediction(field_id="f1", row_index=4, target_cell="D4", answer_value="2路", validation={"acquisition": invalid_trace})
    field = build_field_evidence(item=make_item(4), prediction=prediction,
                                 generated={"acquisition": {"strategy": "sufficiency_guided", "acquisition_rounds": 999}},
                                 top_hits=[make_hit()], overlay=None)
    assert "acquisition" not in field


@pytest.mark.parametrize("old,new,action", [(0, 0, "skipped_non_empty_cell"),
                                          ("人工确认1200kg", "人工确认1200kg", "skipped_non_empty_cell"),
                                          (None, "600kW", "written")])
def test_display_audit_preserves_actual_post_policy_value_and_zero(old, new, action: str) -> None:
    prediction = FieldPrediction(field_id="f1", row_index=4, target_cell="D4", answer_value="600kW", source_chunk_ids=["c1"])
    field = build_field_evidence(item=make_item(4), prediction=prediction, generated=make_answer(), top_hits=[make_hit()], overlay={"writeback_allowed": True})
    audit = {"field_id": "f1", "status": "confirmed", "writeback_action": action, "old_value": old,
             "new_value": new, "existing_value_policy": "preserve", "reason": "target_non_empty" if action != "written" else "written"}
    prior_field, prior_audit = deepcopy(field), deepcopy(audit)
    display = finalize_evidence([field], audit_records=[audit], writeback_status="completed")["fields"][0]
    assert "old_value" in display and "new_value" in display
    assert display["old_value"] == old and display["new_value"] == new
    assert display["existing_value_policy"] == "preserve"
    assert display["writeback_action"] == action
    assert display["answer_value"] == "600kW"
    assert field == prior_field and audit == prior_audit


def test_prompt_exposes_original_text_space_and_requests_quote() -> None:
    messages = build_qdrant_answer_messages(make_item(4), "市电", [make_hit()])
    content = messages[-1]["content"]
    assert '"raw_source_text"' in content and '"quote_text_space": "raw_source_text"' in content
    assert '"quote"' in content
    assert "定位成功不代表答案正确" in content


def test_runner_sidecar_covers_failed_and_not_found_without_writeback(tmp_path: Path) -> None:
    def answer(**kwargs):
        if kwargs["item"]["row_index"] == 5:
            raise RuntimeError("field failed")
        if kwargs["item"]["row_index"] == 6:
            return {"answer_value": "未找到", "answer_status": "not_found", "confidence": 0.1, "source_chunk_ids": []}
        return make_answer()

    runner = make_runner(tmp_path, answer)
    predictions = runner.run([make_item(row) for row in (4, 5, 6)])
    raw = read_jsonl(tmp_path / "predictions_raw.jsonl")
    manifest = json.loads((tmp_path / "run_manifest.json").read_text())
    assert manifest["schema_version"] == "1.3"
    fields = manifest["evidence"]["fields"]
    assert len(fields) == len(predictions) == 3
    assert [field["field_id"] for field in fields] == [prediction.field_id for prediction in predictions]
    assert fields[0]["evidence_refs"][0]["provenance"]["match_status"] == "exact"
    assert fields[1]["evidence_refs"][0]["provenance"]["match_status"] == "unavailable"
    assert fields[2]["answer_status"] == "not_found"
    assert all(field["writeback_action"] == "review_only" for field in fields)
    assert manifest["evidence"]["summary"] == {"exact": 1, "ambiguous": 0, "unmatched": 0, "unavailable": 2}
    assert read_jsonl(tmp_path / "evidence_provenance.jsonl") == fields
    assert raw == [prediction.to_dict() for prediction in predictions]
    assert (tmp_path / "predictions.jsonl").read_bytes() == (tmp_path / "predictions_raw.jsonl").read_bytes()
    assert "provenance" not in json.dumps(raw)


def test_runner_recovers_quote_source_and_hash_from_checkpoint(tmp_path: Path) -> None:
    initial = make_runner(tmp_path, lambda **kwargs: make_answer())
    initial.run([make_item(4)])
    prior = read_jsonl(tmp_path / "evidence_provenance.checkpoint.jsonl")[0]

    def forbidden(**kwargs):
        raise AssertionError("completed field must not invoke retrieval or model")

    resumed = make_runner(tmp_path, forbidden, retrieval_fn=forbidden, resume=True)
    resumed.run([make_item(4)])
    restored = read_jsonl(tmp_path / "evidence_provenance.jsonl")[0]
    assert restored == prior
    assert restored["evidence_refs"][0]["provenance"]["source_text_hash"]


def test_legacy_resume_does_not_upgrade_normalized_preview_to_quote(tmp_path: Path) -> None:
    prediction = FieldPrediction(field_id="item_4", row_index=4, target_cell="D4", answer_value="2路", reference_chunk_ids=["c1"], reference_source_documents=[{"chunk_id": "c1", "text_preview": "automatically enriched text"}])
    write_jsonl(tmp_path / "predictions.checkpoint.jsonl", [prediction.to_dict()])
    runner = make_runner(tmp_path, lambda **kwargs: make_answer(), resume=True)
    persist_form_input_snapshot(tmp_path, build_form_input_snapshot([make_item(4)], target_namespace="room301", global_namespace="global", acquisition_contract=runner.acquisition_contract()), resume=False)
    with pytest.raises(RuntimeError, match="retrieval authority missing"):
        runner.run([make_item(4)])
    assert not (tmp_path / "predictions_raw.jsonl").exists()
    field = build_field_evidence(item=make_item(4), prediction=prediction, generated={}, top_hits=[], overlay=None, unavailable_reason="provenance_checkpoint_missing")
    assert field["question_text"] == "市电进线情况"
    assert field["evidence_refs"][0]["provenance"]["reason"] == "provenance_checkpoint_missing"
    assert field["evidence_refs"][0]["provenance"]["quote"] == ""


def test_runner_uses_actual_formula_skip_audit(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.active["D4"] = "=1+1"
    workbook.save(template)
    runner = make_runner(tmp_path, lambda **kwargs: make_answer(), template_path=template, writeback_enabled=True)
    runner.run([make_item(4)])
    evidence = json.loads((tmp_path / "run_manifest.json").read_text())["evidence"]
    assert evidence["fields"][0]["writeback_status"] == "flagged"
    assert evidence["fields"][0]["writeback_action"] == "skipped_formula"
    assert evidence["fields"][0]["evidence_refs"][0]["provenance"]["match_status"] == "exact"


def make_item(row: int) -> dict:
    return {"form_item_id": f"item_{row}", "file_name": "form.xlsx", "sheet_name": "Sheet1", "row_index": row, "target_cell": f"D{row}", "question_text": "市电进线情况", "instruction_text": "市电路数", "category_path": ["电力"], "answer_example": "", "needs_evidence": False}


def make_hit() -> dict:
    return {"chunk_id": "c1", "namespace": "room301", "source_type": "main_excel_capability", "corpus_layer": "fact", "retrieval_layer": "target_main_fact", "file_name": "真实能力表.xlsx", "anchor": "能力!B4", "raw_source_text": "😀机房：2路市电。", "raw_text": "合成元数据前缀：😀机房：2路市电。", "source": {"sheet_name": "能力", "cell": "B4", "document_id": "doc1"}}


def make_answer() -> dict:
    return {"answer_status": "answered", "answer_value": "2路市电", "confidence": 0.9, "source_chunk_ids": ["c1"], "reference_source_documents": [{"chunk_id": "c1", "quote": "2路市电"}]}


def make_runner(out_dir: Path, answer_caller, **kwargs) -> Step15AgentRunner:
    config = load_app_config(project_root=out_dir, default_config=out_dir / "missing.yaml", cli_overrides={"agentscope": {"enabled": False, "mode": "off"}, "retrieval": {"sufficiency_enabled": False}, "grounding": {"evidence_strength_enabled": False, "field_binding_enabled": False, "slot_decomposition_enabled": False, "pre_writeback_consistency_enabled": False}})
    retrieval_fn = kwargs.pop("retrieval_fn", lambda query: Step15RetrievalResult([make_hit()], [make_hit()], "layered"))
    return Step15AgentRunner(config=config, target_namespace="room301", global_namespace="global", out_dir=out_dir, retrieval_fn=retrieval_fn, answer_caller=answer_caller, chat_max_retries=0, **kwargs)
