"""Frozen KB/version pins through real CLI, native Office, disk Qdrant and local HTTP."""
from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from qdrant_client import QdrantClient, models
from test_runtime_form_contract import (
    COLLECTION,
    GLOBAL,
    TARGET,
    Runtime,
    make_template,
    protected_files,
    read_json,
    read_jsonl,
)
from test_runtime_form_contract import (
    local_models as local_models,
)
from test_runtime_form_contract import (
    runtime as runtime,
)
from test_sufficiency_runtime_contract import final_evidence

WORKSPACE = "10000000-0000-4000-8000-000000000001"
KB = "10000000-0000-4000-8000-000000000002"
GLOBAL_KB = "10000000-0000-4000-8000-000000000003"
DOCUMENT = "10000000-0000-4000-8000-000000000004"
FILE = "10000000-0000-4000-8000-000000000005"
V1 = "20000000-0000-4000-8000-000000000001"
V2 = "20000000-0000-4000-8000-000000000002"
GLOBAL_V = "20000000-0000-4000-8000-000000000003"


def cli(runtime: Runtime, *args: str | Path) -> subprocess.CompletedProcess[str]:
    return subprocess.run([sys.executable, "-m", "nested_doc_rag.cli", *map(str, args)], cwd=runtime.root,
                          env=runtime.env, capture_output=True, text=True, timeout=45, check=False)


def pins() -> list[dict[str, Any]]:
    return [{"collection": COLLECTION, "namespace": namespace, "knowledge_base_id": kb,
             "index_version_id": version, "storage_contract": "versioned_v1"}
            for namespace, kb, version in ((TARGET, KB, V1), (GLOBAL, GLOBAL_KB, GLOBAL_V))]


def ingest(runtime: Runtime, *, version: str, namespace: str, kb: str, value: str) -> tuple[Path, Path, str, dict[str, Any]]:
    root = runtime.root / "builds" / version
    relative = f"{DOCUMENT}/{FILE}/冻结知识.xlsx"
    input_dir = root / "input"
    source = input_dir / relative
    source.parent.mkdir(parents=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "当前UPS参数"
    sheet.append(["字段", "当前值"])
    sheet.append(["UPS容量", value])
    workbook.save(source)
    workbook.close()
    snapshot = {"schema_version": "kb-build-input-v1", "workspace_id": WORKSPACE, "knowledge_base_id": kb,
                "index_version_id": version, "collection": COLLECTION, "namespace": namespace,
                "documents": [{"document_id": DOCUMENT, "file_id": FILE, "filename": source.name,
                               "relative_path": relative, "object_key": f"objects/{FILE}", "document_role": "fact",
                               "sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "size_bytes": source.stat().st_size}]}
    snapshot_path = root / "input_snapshot.json"
    data = json.dumps(snapshot, sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode()
    snapshot_path.write_bytes(data)
    digest = hashlib.sha256(data).hexdigest()
    out_dir = root / "out"
    result = cli(runtime, "ingest-knowledge", "--config", runtime.config, "--input-dir", input_dir,
                 "--namespace", namespace, "--knowledge-base-id", kb, "--qdrant-collection", COLLECTION,
                 "--out-dir", out_dir, "--index-version-id", version, "--input-snapshot", snapshot_path,
                 "--input-snapshot-hash", digest)
    assert result.returncode == 0, result.stdout + result.stderr
    receipt = read_json(out_dir / "validation_receipt.json")
    assert receipt["schema_version"] == "kb-index-validation-v1" and receipt["input_snapshot_hash"] == digest
    assert receipt["index_version_id"] == version and receipt["knowledge_base_id"] == kb
    assert receipt["expected_evidence_count"] == receipt["actual_evidence_count"] == 2
    assert receipt["expected_schema_count"] == receipt["actual_schema_count"] == 1
    assert receipt["source_hashes_verified"] and receipt["smoke_passed"]
    assert (out_dir / "input_snapshot.json").read_bytes() == data
    assert read_json(out_dir / "run_manifest.json")["artifacts"]["input_snapshot"] == "input_snapshot.json"
    return input_dir, snapshot_path, digest, next(record for record in read_jsonl(out_dir / "ingestion_manifest.jsonl") if record.get("field_value") == value)


@dataclass
class PinnedRuntime:
    runtime: Runtime
    target: dict[str, Any]
    newer: dict[str, Any]
    global_fact: dict[str, Any]
    scope_path: Path
    template: Path


@pytest.fixture
def pinned_runtime(runtime: Runtime) -> PinnedRuntime:
    _, _, _, target = ingest(runtime, version=V1, namespace=TARGET, kb=KB, value="500kVA")
    _, _, _, newer = ingest(runtime, version=V2, namespace=TARGET, kb=KB, value="900kVA")
    _, _, _, global_fact = ingest(runtime, version=GLOBAL_V, namespace=GLOBAL, kb=GLOBAL_KB, value="1000kVA")
    scope_path = runtime.root / "frozen_index_scopes.json"
    scope_path.write_text(json.dumps(pins(), ensure_ascii=False), encoding="utf-8")
    template = make_template(runtime.root / "固定旧版本工勘.xlsx", sheets={"现场工勘": [(2, "UPS容量")]})
    runtime.models.answer = {"answer_value": "500kVA", "answer_status": "answered", "confidence": 0.95,
                             "source_chunk_ids": [target["chunk_id"]], "reference_source_documents": [{
                                 "chunk_id": target["chunk_id"], "quote": target["raw_source_text"], "reason": "frozen V1 native source",
                             }]}
    return PinnedRuntime(runtime, target, newer, global_fact, scope_path, template)


def run_pinned(fixture: PinnedRuntime, *extra: str | Path) -> subprocess.CompletedProcess[str]:
    return fixture.runtime.run("--template", fixture.template, "--index-scopes", fixture.scope_path,
                               "--schema-first-enabled", "--writeback", *extra)


def test_old_pins_survive_v2_native_ingestion_and_scope_schema_values_writeback_and_archive(pinned_runtime: PinnedRuntime) -> None:
    fixture, runtime = pinned_runtime, pinned_runtime.runtime
    result = run_pinned(fixture)
    runtime.assert_completed(result, count=1)
    output = runtime.root / "run"
    prediction = read_jsonl(output / "predictions_raw.jsonl")[0]
    assert prediction["source_chunk_ids"] == [fixture.target["chunk_id"]]
    ref = prediction["evidence_refs"][0]
    assert ref["index_version"] == V1 and ref["qdrant_point_id"] == fixture.target["point_id"]
    pack = final_evidence(runtime.models.chat_requests(sufficiency=False)[0])
    assert {hit["index_version"] for hit in pack} <= {V1, GLOBAL_V}
    assert V2 not in {hit["index_version"] for hit in pack}
    assert fixture.newer["chunk_id"] not in {hit["chunk_id"] for hit in pack}
    authority = read_jsonl(output / "retrieval_evidence.jsonl")[0]["top_hits"]
    assert all((hit["namespace"], hit["knowledge_base_id"], hit["index_version_id"]) in {
        (TARGET, KB, V1), (GLOBAL, GLOBAL_KB, GLOBAL_V),
    } for hit in authority)
    assert "anonymous_fact" not in {hit["chunk_id"] for hit in authority}
    trace = read_jsonl(output / "trace.jsonl")
    newer_schema = read_jsonl(runtime.root / "builds" / V2 / "out" / "field_schema_manifest.jsonl")[0]["schema_id"]
    assert newer_schema not in json.dumps(trace, ensure_ascii=False)
    manifest = read_json(output / "run_manifest.json")
    assert manifest["index_scopes"] == pins()
    assert read_json(output / manifest["artifacts"]["index_scopes"]) == pins()
    assert manifest["form_input"]["acquisition_contract"]["index_scope_contract"]["scopes"] == pins()
    audit = read_jsonl(output / "writeback_audit.jsonl")[0]
    assert audit["new_value"] == "500kVA" and audit["evidence_refs"][0]["index_version"] == V1
    workbook = load_workbook(output / "filled_form.xlsx")
    assert workbook["现场工勘"]["D2"].value == "500kVA"
    workbook.close()
    valid = cli(runtime, "validate-artifacts", "--run-dir", output)
    assert valid.returncode == 0, valid.stdout + valid.stderr


def test_changed_pin_rejects_actual_cli_resume_before_models_or_artifact_changes(pinned_runtime: PinnedRuntime) -> None:
    fixture, runtime = pinned_runtime, pinned_runtime.runtime
    runtime.assert_completed(run_pinned(fixture), count=1)
    output = runtime.root / "run"
    (output / "predictions.checkpoint.jsonl").write_text("{damaged checkpoint}\n", encoding="utf-8")
    changed = pins()
    changed[0]["index_version_id"] = V2
    fixture.scope_path.write_text(json.dumps(changed), encoding="utf-8")
    before, calls = protected_files(output), runtime.models.recorded()
    result = run_pinned(fixture, "--resume")
    assert result.returncode != 0 and "cannot resume" in (result.stdout + result.stderr).lower()
    assert "jsondecodeerror" not in (result.stdout + result.stderr).lower()
    assert runtime.models.recorded() == calls and protected_files(output) == before


def test_tampered_checkpoint_authority_cannot_resume_or_change_archived_outputs(pinned_runtime: PinnedRuntime) -> None:
    fixture, runtime = pinned_runtime, pinned_runtime.runtime
    runtime.assert_completed(run_pinned(fixture), count=1)
    output = runtime.root / "run"
    authority = read_jsonl(output / "retrieval_evidence.checkpoint.jsonl")
    authority[0]["top_hits"][0]["index_version_id"] = V2
    (output / "retrieval_evidence.checkpoint.jsonl").write_text("".join(json.dumps(row) + "\n" for row in authority), encoding="utf-8")
    before, calls = protected_files(output), runtime.models.recorded()
    result = run_pinned(fixture, "--resume")
    assert result.returncode != 0 and "indexscopeerror" in (result.stdout + result.stderr).lower()
    assert runtime.models.recorded() == calls and protected_files(output) == before


def test_validator_rejects_actual_archived_authority_from_a_foreign_version(pinned_runtime: PinnedRuntime) -> None:
    fixture, runtime = pinned_runtime, pinned_runtime.runtime
    runtime.assert_completed(run_pinned(fixture), count=1)
    output = runtime.root / "run"
    authority = read_jsonl(output / "retrieval_evidence.jsonl")
    authority[0]["top_hits"][0]["index_version_id"] = V2
    (output / "retrieval_evidence.jsonl").write_text("".join(json.dumps(row) + "\n" for row in authority), encoding="utf-8")
    result = cli(runtime, "validate-artifacts", "--run-dir", output)
    assert result.returncode != 0 and "EV_INDEX_SCOPE_MISMATCH" in result.stderr


def test_scope_file_with_extra_namespace_is_rejected_before_any_model_or_output(pinned_runtime: PinnedRuntime) -> None:
    fixture, runtime = pinned_runtime, pinned_runtime.runtime
    fixture.scope_path.write_text(json.dumps([*pins(), {**pins()[0], "namespace": "extra-namespace"}]), encoding="utf-8")
    calls = runtime.models.recorded()
    result = run_pinned(fixture)
    assert result.returncode != 0 and "namespace" in result.stderr.lower()
    assert runtime.models.recorded() == calls and not (runtime.root / "run").exists()


def test_explicit_legacy_pins_read_only_unversioned_schema_values_despite_versioned_decoys(pinned_runtime: PinnedRuntime) -> None:
    fixture, runtime = pinned_runtime, pinned_runtime.runtime
    source = deepcopy(fixture.target)
    source["chunk_id"] = source["evidence_id"] = "legacy-native-ups"
    for container in (source, source["metadata"], source["source"]):
        for key in ("index_version_id", "index_version", "storage_contract", "point_id", "qdrant_point_id"):
            container.pop(key, None)
    schema = read_jsonl(runtime.root / "builds" / V1 / "out" / "field_schema_manifest.jsonl")[0]
    for key in ("index_version_id", "index_version", "point_id"):
        schema.pop(key, None)
    client = QdrantClient(path=str(runtime.qdrant))
    try:
        client.upsert(COLLECTION, points=[models.PointStruct(id=991, vector=[1.0, 0.0, 0.0], payload=source),
                                          models.PointStruct(id=992, vector=[1.0, 0.0, 0.0], payload=schema)], wait=True)
    finally:
        client.close()
    legacy = pins()
    legacy[0]["storage_contract"] = "legacy_unversioned"
    fixture.scope_path.write_text(json.dumps(legacy), encoding="utf-8")
    runtime.models.answer["source_chunk_ids"] = ["legacy-native-ups"]
    runtime.models.answer["reference_source_documents"][0]["chunk_id"] = "legacy-native-ups"
    result = run_pinned(fixture)
    runtime.assert_completed(result, count=1)
    assert "LegacyIndexScopeWarning" in result.stderr
    authority = read_jsonl(runtime.root / "run" / "retrieval_evidence.jsonl")[0]["top_hits"]
    target_hits = [hit for hit in authority if hit["namespace"] == TARGET]
    assert {hit["chunk_id"] for hit in target_hits} == {"legacy-native-ups"}
    assert all(hit.get("index_version_id") is None for hit in target_hits)
    assert target_hits[0]["point_id"] == target_hits[0]["qdrant_point_id"] == "991"


def test_ingestion_snapshot_raw_hash_failure_via_cli_does_not_contact_models_or_create_receipt(runtime: Runtime) -> None:
    input_dir, snapshot, digest, _ = ingest(runtime, version=V1, namespace=TARGET, kb=KB, value="500kVA")
    snapshot.write_bytes(snapshot.read_bytes() + b"\n")
    calls = runtime.models.recorded()
    out_dir = runtime.root / "rejected-ingestion"
    result = cli(runtime, "ingest-knowledge", "--config", runtime.config, "--input-dir", input_dir, "--namespace", TARGET,
                 "--knowledge-base-id", KB, "--out-dir", out_dir, "--qdrant-collection", COLLECTION,
                 "--index-version-id", V1, "--input-snapshot", snapshot, "--input-snapshot-hash", digest)
    assert result.returncode != 0 and "raw bytes hash mismatch" in result.stderr
    assert runtime.models.recorded() == calls and not out_dir.exists()
