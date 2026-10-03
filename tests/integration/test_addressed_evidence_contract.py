"""Native upload bytes -> disk Qdrant -> CLI -> independently gated writer.

Only HTTP models are deterministic substitutes. Native ingestion, retrieval,
form parser, Runner, workbook writer and artifact validator are production code.
"""
from __future__ import annotations

import json
from copy import deepcopy

import pytest
from openpyxl import Workbook, load_workbook
from qdrant_client import QdrantClient, models
from test_runtime_form_contract import (
    COLLECTION,
    TARGET,
    Runtime,
    make_template,
    protected_files,
    read_json,
    read_jsonl,
)
from test_runtime_form_contract import (
    local_models as local_models,
)
from test_runtime_form_contract import (
    runtime as runtime,
)

from nested_doc_rag.artifacts import validate_step15_artifacts
from nested_doc_rag.evidence_resolver import REVIEW_DISPLAY_REFERENCE_CONTRACT
from nested_doc_rag.ingestion import build_ingestion_records, upsert_records


class FixedEmbedding:
    def embed(self, texts):
        return [[1.0, 0.0, 0.0] for _ in texts]


def native_upload(runtime: Runtime, *, missing_source: bool = False) -> dict:
    uploads = runtime.root / "uploads"
    uploads.mkdir()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "能力"
    worksheet.append(["字段", "当前实际值"])
    worksheet.append(["UPS容量", "500kVA"])
    workbook.save(uploads / "新知识.xlsx")
    workbook.close()
    records, skipped = build_ingestion_records(uploads, namespace=TARGET, knowledge_base_id="native-contract-kb")
    assert not skipped
    fact = next(record for record in records if record["evidence_kind"] == "structured_field")
    if missing_source:
        fact.pop("raw_source_text", None)
        fact.pop("source_text_hash_space", None)
    client = QdrantClient(path=str(runtime.qdrant))
    try:
        client.delete(COLLECTION, models.PointIdsList(points=[1]), wait=True)
        count, dimension = upsert_records(
            qdrant=client, collection_name=COLLECTION, records=records,
            embedder=FixedEmbedding(), batch_size=8, namespace=TARGET,
        )
        assert count == len(records) and dimension == 3
    finally:
        client.close()
    return fact


@pytest.mark.parametrize("case", ["valid", "forged_model_address", "forged_id", "missing_source", "forged_attachment"])
def test_actual_cli_only_confirms_resolved_source_authority(runtime: Runtime, case: str) -> None:
    fact = native_upload(runtime, missing_source=case == "missing_source")
    selected = "invented-evidence" if case == "forged_id" else fact["chunk_id"]
    runtime.models.answer = {
        "answer_value": "500kVA", "answer_status": "answered", "confidence": 0.95,
        "source_chunk_ids": [selected],
        "evidence_attachment_ids": ["invented-attachment"] if case == "forged_attachment" else [],
        "reference_source_documents": [{
            "chunk_id": selected, "quote": "500kVA", "reason": "direct source",
            "file_name": "forged.xlsx", "sheet_name": "invented", "cell": "X999",
            "object_key": "other-owner/private.xlsx", "proof_attachment_ids": ["invented-attachment"],
        }],
    }
    template = make_template(runtime.root / "新表单.xlsx", sheets={"调研": [(2, "UPS容量")]})
    original = template.read_bytes()
    result = runtime.run("--template", template, "--writeback")
    runtime.assert_completed(result, count=1)
    output = runtime.root / "run"
    raw = read_jsonl(output / "predictions_raw.jsonl")[0]
    overlay = read_jsonl(output / "agent_overlays.jsonl")[0]
    audit = read_jsonl(output / "writeback_audit.jsonl")[0]
    assert raw["answer_status"] == "answered" and raw["answer_value"] == "500kVA"
    assert raw["validation"]["step15_generated"] == runtime.models.answer
    assert template.read_bytes() == original
    assert (output / "predictions_raw.jsonl").read_bytes() == (output / "predictions.jsonl").read_bytes()
    safe = case in {"valid", "forged_model_address"}
    workbook = load_workbook(output / "filled_form.xlsx")
    try:
        assert (workbook["调研"]["D2"].value == "500kVA") is safe
        assert overlay["writeback_allowed"] is safe
        assert (audit["writeback_action"] == "written") is safe
        if safe:
            ref = raw["evidence_refs"][0]
            assert ref["source_text"] == fact["raw_source_text"]
            assert ref["file_name"] == "新知识.xlsx"
            assert ref["sheet_name"] == "能力" and ref["row_index"] == 2
            assert ref["knowledge_base_id"] == "native-contract-kb"
            assert "Evidence" in workbook.sheetnames
            assert "Evidence" in workbook["调研"]["D2"].comment.text
            assert all(doc["file_name"] != "forged.xlsx" and doc["object_key"] != "other-owner/private.xlsx" for doc in raw["reference_source_documents"])
        else:
            assert raw["validation"]["addressable_evidence"]["errors"]
            assert overlay["review_required"] and audit["status"] != "confirmed"
    finally:
        workbook.close()
    validated = validate_step15_artifacts(output)
    assert validated["valid"] and validated["evidence_validation"] == "strict"
    assert read_json(output / "run_manifest.json")["schema_version"] == "1.3"


@pytest.mark.parametrize("quote", ["500kVA", "", "模型编造的引用"])
def test_native_partial_review_metadata_never_claims_typed_write_authority(runtime: Runtime, quote: str) -> None:
    fact = native_upload(runtime)
    runtime.models.answer = {
        "answer_value": "检索到相关线索，但证据不足", "answer_status": "partial_clue", "confidence": 0.2,
        "source_chunk_ids": [], "evidence_attachment_ids": [],
        "reference_source_documents": [{"chunk_id": fact["chunk_id"], "quote": quote, "reason": "review clue"}],
    }
    template = make_template(runtime.root / "部分线索表单.xlsx", sheets={"调研": [(2, "UPS容量")]})
    original = template.read_bytes()
    result = runtime.run("--template", template, "--writeback")
    runtime.assert_completed(result, count=1)
    output = runtime.root / "run"
    raw = read_jsonl(output / "predictions_raw.jsonl")[0]
    audit = read_jsonl(output / "writeback_audit.jsonl")[0]
    overlay = read_jsonl(output / "agent_overlays.jsonl")[0]
    assert raw["answer_status"] == "partial_clue"
    assert raw["evidence_refs"] == raw["source_chunk_ids"] == []
    assert not overlay["writeback_allowed"]
    assert audit["status"] == "flagged" and audit["writeback_action"] == "review_only"
    reference = next(ref for ref in audit["evidence_refs"] if ref["chunk_id"] == fact["chunk_id"])
    assert reference["reference_contract"] == REVIEW_DISPLAY_REFERENCE_CONTRACT
    assert reference["evidence_kind"] == "structured_field"
    assert reference["file_name"] == "新知识.xlsx" and reference["cell_range"] == fact["cell_range"]
    assert reference["quote"] == quote
    assert "source_text" not in reference and "start" not in reference and "end" not in reference
    assert template.read_bytes() == original
    workbook = load_workbook(output / "filled_form.xlsx")
    try:
        assert workbook["调研"]["D2"].value is None
    finally:
        workbook.close()
    validated = validate_step15_artifacts(output)
    assert validated["valid"] and validated["evidence_validation"] == "strict"
    assert validated["review_evidence_diagnostics"]


@pytest.mark.parametrize("damage", ["missing_authority", "tampered_ref"])
def test_resume_rejects_missing_or_tampered_evidence_before_model_calls(runtime: Runtime, damage: str) -> None:
    fact = native_upload(runtime)
    runtime.models.answer = {"answer_value": "500kVA", "answer_status": "answered", "confidence": 0.9, "source_chunk_ids": [fact["chunk_id"]]}
    template = make_template(runtime.root / "resume.xlsx", sheets={"调研": [(2, "UPS容量")]})
    result = runtime.run("--template", template)
    runtime.assert_completed(result, count=1)
    output = runtime.root / "run"
    if damage == "missing_authority":
        (output / "retrieval_evidence.checkpoint.jsonl").unlink()
    else:
        rows = deepcopy(read_jsonl(output / "predictions.checkpoint.jsonl"))
        rows[0]["evidence_refs"][0]["file_name"] = "tampered.xlsx"
        (output / "predictions.checkpoint.jsonl").write_text(json.dumps(rows[0], ensure_ascii=False) + "\n")
    before = protected_files(output)
    calls = runtime.models.recorded()
    resumed = runtime.run("--template", template, "--resume")
    assert resumed.returncode != 0 and "cannot resume" in resumed.stderr + resumed.stdout
    assert runtime.models.recorded() == calls
    assert protected_files(output) == before
