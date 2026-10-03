"""Real CLI, native uploads and disk Qdrant with explicit localhost models.

Deterministic HTTP replies test acquisition, citation and writeback contracts;
they do not measure a real model's semantic judgement or answer accuracy.
No production parser, retriever, runner, writer or validator is replaced.
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from qdrant_client import QdrantClient, models
from test_runtime_form_contract import (
    COLLECTION,
    GLOBAL,
    TARGET,
    Runtime,
    assert_rejected_unchanged,
    make_template,
    protected_files,
    read_json,
    read_jsonl,
)
from test_runtime_form_contract import (
    local_models as local_models,
)
from test_runtime_form_contract import (
    runtime as runtime,
)


def cli(runtime: Runtime, *args: str | Path, wrapper: str | None = None) -> subprocess.CompletedProcess[str]:
    command = [sys.executable, "-m", "nested_doc_rag.cli"] if wrapper is None else [sys.executable, "-c", wrapper]
    return subprocess.run(
        [*command, *map(str, args)], cwd=runtime.root, env=runtime.env,
        capture_output=True, text=True, timeout=45, check=False,
    )


def native_upload(runtime: Runtime, *, capacity_format: str | None = None, primary_capacity: bool = False) -> list[dict[str, Any]]:
    uploads = runtime.root / "uploads"
    uploads.mkdir()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "UPS现网配置"
    worksheet.append(["字段", "当前实际值"])
    worksheet.append(["UPS容量", "500kVA"] if primary_capacity else ["UPS品牌", "施耐德"])
    workbook.save(uploads / "UPS当前配置.xlsx")
    workbook.close()
    if capacity_format is not None:
        text = "301机房当前已投运UPS单台容量为500kVA。"
        path = uploads / f"UPS容量实测.{capacity_format}"
        if capacity_format == "docx":
            document = Document()
            document.add_paragraph(text)
            document.save(path)
        else:
            path.write_text(text, encoding="utf-8")
    client = QdrantClient(path=str(runtime.qdrant))
    try:
        client.delete(COLLECTION, models.PointIdsList(points=[1]), wait=True)
    finally:
        client.close()
    result = cli(
        runtime, "ingest-knowledge", "--config", runtime.config, "--input-dir", uploads,
        "--namespace", TARGET, "--knowledge-base-id", "sufficiency-native-kb",
        "--qdrant-collection", COLLECTION, "--out-dir", runtime.root / "ingested",
    )
    assert result.returncode == 0, result.stdout + result.stderr
    records = read_jsonl(runtime.root / "ingested" / "ingestion_manifest.jsonl")
    assert len(records) == (3 if capacity_format else 2)
    assert read_json(runtime.root / "ingested" / "summary.json")["upserted_count"] == len(records)
    return records


def user_text(request: dict[str, Any]) -> str:
    return next(message["content"] for message in request["messages"] if message["role"] == "user")


def sufficiency_payload(request: dict[str, Any]) -> dict[str, Any]:
    payload, _ = json.JSONDecoder().raw_decode(user_text(request).split("\n", 1)[1])
    return payload


def final_evidence(request: dict[str, Any]) -> list[dict[str, Any]]:
    evidence, _ = json.JSONDecoder().raw_decode(user_text(request).split("retrieved_chunks:\n", 1)[1])
    return evidence


def semantic_response(request: dict[str, Any]) -> dict[str, Any]:
    payload = sufficiency_payload(request)
    question = payload["form_item"]["question_text"]
    required = {"UPS品牌": "施耐德", "UPS容量": "500kVA"} if "品牌" in question else {"UPS容量": "500kVA"}
    evidence = payload["retrieved_evidence"]
    missing = [fact for fact, value in required.items() if not any(value in str(hit.get("raw_source_text") or "") for hit in evidence)]
    supporting = [hit["evidence_id"] for hit in evidence if any(value in str(hit.get("raw_source_text") or "") for value in required.values())]
    return {
        "sufficient": not missing, "missing_facts": missing,
        "supporting_evidence_ids": supporting,
        "reason": "Localhost substitute checks only fixture facts present in the current pack",
    }


def configure_answer(runtime: Runtime, records: list[dict[str, Any]], *, composite: bool = False) -> None:
    sources = [record for record in records if record["field_value"] == "500kVA" or record["field_value"] == "施耐德" or "500kVA" in record["raw_source_text"]]
    runtime.models.answer = {
        "answer_value": "施耐德，500kVA" if composite else "500kVA",
        "answer_status": "answered", "confidence": 0.95,
        "source_chunk_ids": [source["chunk_id"] for source in sources], "evidence_attachment_ids": [],
        "reference_source_documents": [
            {"chunk_id": source["chunk_id"], "quote": source["raw_source_text"], "reason": "native source",
             "file_name": "模型伪造地址.xlsx", "sheet_name": "模型伪造工作表", "cell": "X999"}
            for source in sources
        ],
    }


def completed(runtime: Runtime) -> tuple[dict[str, Any], dict[str, Any]]:
    output = runtime.root / "run"
    manifest = read_json(output / "run_manifest.json")
    assert manifest["status"] == "completed" and manifest["counts"]["failed"] == 0
    assert manifest["counts"]["total_fields"] == 1
    predictions = read_jsonl(output / "predictions_raw.jsonl")
    assert len(predictions) == 1 and predictions[0]["method_name"] != "step15_agent_failed"
    return predictions[0], read_jsonl(output / "agent_overlays.jsonl")[0]


def assert_acquisition(prediction: dict[str, Any], *, rounds: int, queries: int, sufficient: bool) -> dict[str, Any]:
    acquisition = prediction["validation"]["acquisition"]
    assert acquisition["strategy"] == "sufficiency_guided"
    assert acquisition["acquisition_rounds"] == rounds
    assert acquisition["qdrant_query_calls"] == queries
    assert [entry["qdrant_query_calls"] for entry in acquisition["rounds"]] == ([2] if rounds == 1 else [2, 5])
    assert acquisition["final_sufficiency"]["sufficient"] is sufficient
    return acquisition


def validate_actual_cli(runtime: Runtime) -> None:
    result = cli(runtime, "validate-artifacts", "--run-dir", runtime.root / "run")
    assert result.returncode == 0, result.stdout + result.stderr
    validation = json.loads(result.stdout)
    assert validation["valid"] and validation["evidence_validation"] == "strict"


def assert_written(runtime: Runtime, prediction: dict[str, Any], overlay: dict[str, Any]) -> None:
    assert overlay["writeback_allowed"]
    audit = read_jsonl(runtime.root / "run" / "writeback_audit.jsonl")[0]
    assert audit["writeback_action"] == "written" and audit["status"] == "confirmed"
    workbook = load_workbook(runtime.root / "run" / "filled_form.xlsx")
    try:
        assert workbook["实际调研"]["D2"].value == prediction["answer_value"]
        assert "Evidence" in workbook.sheetnames and "Evidence" in workbook["实际调研"]["D2"].comment.text
    finally:
        workbook.close()


def test_sufficient_primary_uses_one_acquisition_and_two_actual_qdrant_queries(runtime: Runtime) -> None:
    records = native_upload(runtime, primary_capacity=True)
    runtime.models.sufficiency = semantic_response
    configure_answer(runtime, records)
    template = make_template(runtime.root / "只需主轮.xlsx", sheets={"实际调研": [(2, "UPS容量")]})
    original = template.read_bytes()
    result = runtime.run("--template", template, "--writeback", "--sufficiency-enabled")
    runtime.assert_completed(result, count=1)
    prediction, overlay = completed(runtime)
    assert_acquisition(prediction, rounds=1, queries=2, sufficient=True)
    assert len(runtime.models.chat_requests(sufficiency=True)) == 1
    answers = runtime.models.chat_requests(sufficiency=False)
    assert len(answers) == 1 and all(hit["retrieval_round"] == 0 for hit in final_evidence(answers[0]))
    fact = next(record for record in records if record["evidence_kind"] == "structured_field")
    ref = prediction["evidence_refs"][0]
    assert ref["chunk_id"] == fact["chunk_id"]
    assert ref["file_name"] == "UPS当前配置.xlsx" and ref["sheet_name"] == "UPS现网配置"
    assert ref["row_index"] == 2 and ref["cell_range"] == "A2:B2"
    assert ref["source_text"] == fact["raw_source_text"]
    assert_written(runtime, prediction, overlay)
    assert template.read_bytes() == original
    validate_actual_cli(runtime)


@pytest.mark.parametrize("capacity_format", ["txt", "docx"])
def test_targeted_round_adds_native_capacity_deduplicates_ids_and_writes_real_addresses(runtime: Runtime, capacity_format: str) -> None:
    records = native_upload(runtime, capacity_format=capacity_format)
    runtime.models.sufficiency = semantic_response
    configure_answer(runtime, records, composite=True)
    template = make_template(runtime.root / "补容量.xlsx", sheets={"实际调研": [(2, "UPS品牌及容量")]})
    result = runtime.run("--template", template, "--writeback", "--sufficiency-enabled")
    runtime.assert_completed(result, count=1)
    prediction, overlay = completed(runtime)
    acquisition = assert_acquisition(prediction, rounds=2, queries=7, sufficient=True)
    checks = [sufficiency_payload(request) for request in runtime.models.chat_requests(sufficiency=True)]
    assert len(checks) == 2
    assert all("500kVA" not in str(hit["raw_source_text"]) for hit in checks[0]["retrieved_evidence"])
    assert acquisition["rounds"][1]["missing_facts"] == ["UPS容量"]
    assert "缺失事实=UPS容量" in acquisition["rounds"][1]["query"]
    assert acquisition["rounds"][1]["evidence_gain"] == 1
    assert acquisition["rounds"][0]["hit_count"] == 2 and acquisition["rounds"][1]["hit_count"] == len(records)
    answers = runtime.models.chat_requests(sufficiency=False)
    assert len(answers) == 1
    pack = final_evidence(answers[0])
    ids = [hit["evidence_id"] for hit in pack]
    assert len(ids) == len(set(ids)) == len(records)
    assert [hit["evidence_kind"] for hit in pack][:2] == ["structured_field", "table_row"]
    for hit in pack:
        if "500kVA" in str(hit["raw_source_text"]):
            assert hit["retrieval_layer"] == "target_text_detail"
            assert hit["retrieval_round"] == 1 and hit["triggered_by"] == ["UPS容量"]
        else:
            assert hit["retrieval_round"] == 0 and hit["triggered_by"] == []
    authority = read_jsonl(runtime.root / "run" / "retrieval_evidence.checkpoint.jsonl")[0]["top_hits"]
    assert len(authority) == len({hit["evidence_id"] for hit in authority}) == len(records)
    capacity = next(record for record in records if record["file_name"] == f"UPS容量实测.{capacity_format}")
    references = {ref["chunk_id"]: ref for ref in prediction["evidence_refs"]}
    assert len(references) == 2
    brand = next(record for record in records if record["evidence_kind"] == "structured_field")
    brand_ref = references[brand["chunk_id"]]
    assert brand_ref["file_name"] == "UPS当前配置.xlsx" and brand_ref["sheet_name"] == "UPS现网配置"
    assert brand_ref["row_index"] == 2 and brand_ref["cell_range"] == "A2:B2"
    ref = references[capacity["chunk_id"]]
    assert ref["file_name"] == capacity["file_name"] and ref["source_anchor"] == capacity["anchor"]
    assert ref["paragraph_index"] == capacity["paragraph_index"]
    assert ref["source_text"] == capacity["raw_source_text"] and ref["knowledge_base_id"] == "sufficiency-native-kb"
    assert all(ref["file_name"] != "模型伪造地址.xlsx" for ref in references.values())
    assert_written(runtime, prediction, overlay)
    validate_actual_cli(runtime)


@pytest.mark.parametrize("empty_scope", [False, True])
def test_remaining_gap_is_system_abstention_without_final_answer_or_writeback(runtime: Runtime, empty_scope: bool) -> None:
    records = native_upload(runtime)
    runtime.models.sufficiency = semantic_response
    configure_answer(runtime, records, composite=True)
    template = make_template(runtime.root / "仍缺容量.xlsx", sheets={"实际调研": [(2, "UPS品牌及容量")]})
    result = runtime.run(
        "--template", template, "--writeback", "--sufficiency-enabled",
        target="no_current_room_evidence" if empty_scope else TARGET,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    prediction, overlay = completed(runtime)
    acquisition = assert_acquisition(prediction, rounds=2, queries=7, sufficient=False)
    assert "UPS容量" in acquisition["final_sufficiency"]["missing_facts"]
    assert prediction["answer_status"] == ("not_found" if empty_scope else "partial_clue")
    assert prediction["answer_value"] == "未找到" and not prediction["source_chunk_ids"]
    assert prediction["validation"]["step15_generated"]["origin"] == "system_sufficiency_abstention"
    assert len(runtime.models.chat_requests(sufficiency=True)) == 2
    assert not runtime.models.chat_requests(sufficiency=False)
    assert not overlay["writeback_allowed"] and overlay["review_required"]
    audit = read_jsonl(runtime.root / "run" / "writeback_audit.jsonl")[0]
    assert audit["writeback_action"] != "written" and audit["status"] != "confirmed"
    workbook = load_workbook(runtime.root / "run" / "filled_form.xlsx")
    try:
        assert workbook["实际调研"]["D2"].value is None
    finally:
        workbook.close()
    validate_actual_cli(runtime)


@pytest.mark.parametrize("invalid", ["malformed", "unknown_id"])
def test_invalid_sufficiency_cannot_call_answer_or_confirm_writeback(runtime: Runtime, invalid: str) -> None:
    records = native_upload(runtime, primary_capacity=True)
    configure_answer(runtime, records)
    runtime.models.sufficiency = lambda request: (
        {"answer_status": "answered", "answer_value": "500kVA"} if invalid == "malformed" else {
            "sufficient": True, "missing_facts": [], "supporting_evidence_ids": ["outside-current-pack"], "reason": "forged support",
        }
    )
    template = make_template(runtime.root / "严格拒绝.xlsx", sheets={"实际调研": [(2, "UPS容量")]})
    result = runtime.run("--template", template, "--writeback", "--sufficiency-enabled")
    assert result.returncode == 0, result.stdout + result.stderr
    prediction, overlay = completed(runtime)
    acquisition = assert_acquisition(prediction, rounds=2, queries=7, sufficient=False)
    codes = {diagnostic["code"] for diagnostic in acquisition["final_sufficiency"]["diagnostics"]}
    assert ("SUFFICIENCY_SCHEMA_INVALID" if invalid == "malformed" else "SUFFICIENCY_SUPPORT_NOT_IN_PACK") in codes
    assert prediction["answer_status"] == "partial_clue" and prediction["answer_value"] == "未找到"
    assert len(runtime.models.chat_requests(sufficiency=True)) == 2 and not runtime.models.chat_requests(sufficiency=False)
    assert not overlay["writeback_allowed"]
    assert read_jsonl(runtime.root / "run" / "writeback_audit.jsonl")[0]["writeback_action"] != "written"
    validate_actual_cli(runtime)


@pytest.mark.parametrize("change", ["strategy", "answer_prompt", "sufficiency_prompt"])
def test_strategy_or_prompt_change_rejects_resume_before_models_or_artifact_changes(runtime: Runtime, change: str) -> None:
    records = native_upload(runtime, primary_capacity=True)
    runtime.models.sufficiency = semantic_response
    configure_answer(runtime, records)
    template = make_template(runtime.root / "冻结策略.xlsx", sheets={"实际调研": [(2, "UPS容量")]})
    initial = runtime.run("--template", template, "--sufficiency-enabled")
    runtime.assert_completed(initial, count=1)
    output = runtime.root / "run"
    contract = read_json(output / "form_input_snapshot.json")["acquisition_contract"]
    assert contract["sufficiency_enabled"] and contract["sufficiency_prompt_version"] == "evidence-sufficiency-v1"
    # A damaged checkpoint proves contract identity is checked before loading it.
    (output / "predictions.checkpoint.jsonl").write_text("{damaged checkpoint}\n", encoding="utf-8")
    before, requests = protected_files(output), runtime.models.recorded()
    if change == "strategy":
        result = runtime.run("--template", template, "--resume", "--no-sufficiency-enabled")
    elif change == "answer_prompt":
        result = runtime.run("--template", template, "--resume", "--sufficiency-enabled", "--prompt-version", "agent_v2")
    else:
        # Only the published prompt identity is changed; runtime logic stays real.
        wrapper = (
            "import sys; import nested_doc_rag.grounding.sufficiency as suff; "
            "suff.SUFFICIENCY_PROMPT_VERSION = 'evidence-sufficiency-v2-test'; "
            "from nested_doc_rag.cli import main; main(sys.argv[1:])"
        )
        result = cli(
            runtime, "run-step15-agent", "--config", runtime.config, "--qdrant-path", runtime.qdrant,
            "--qdrant-collection", COLLECTION, "--target-namespace", TARGET, "--global-namespace", GLOBAL,
            "--rows", "all", "--out-dir", output, "--no-judge", "--template", template,
            "--resume", "--sufficiency-enabled", wrapper=wrapper,
        )
    assert_rejected_unchanged(runtime, result, before, requests)
