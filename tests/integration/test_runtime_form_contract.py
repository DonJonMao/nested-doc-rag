"""Runtime form contracts through the real CLI, disk Qdrant and local HTTP.

The model substitute deliberately returns ``not_found``. These tests verify
input propagation, archival and safe checkpoint reuse, not answer quality.
No parser, runner, retriever or CLI implementation is replaced.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
import threading
from collections import Counter
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from qdrant_client import QdrantClient, models

REPO_ROOT = Path(__file__).resolve().parents[2]
TARGET = "contract_room_301"
GLOBAL = "contract_global"
COLLECTION = "runtime_form_contract"
VECTOR = [1.0, 0.0, 0.0]
DEFAULT_HEADERS = ("类别", "指标名称", "填写说明", "实际情况", "备注")


@dataclass
class LocalModels:
    """Record real curl requests while serving deterministic protocol replies."""

    requests: list[tuple[str, dict[str, Any]]] = field(default_factory=list)
    lock: threading.Lock = field(default_factory=threading.Lock)
    base_url: str = ""
    answer: dict[str, Any] | None = None
    sufficiency: Callable[[dict[str, Any]], Any] | None = None

    def chat_requests(self, *, sufficiency: bool) -> list[dict[str, Any]]:
        return [
            request for path, request in self.recorded()
            if path == "/v1/chat/completions"
            and any(
                message.get("role") == "system" and "Evidence Sufficiency Check" in str(message.get("content") or "")
                for message in request.get("messages", [])
            ) is sufficiency
        ]

    def recorded(self) -> list[tuple[str, dict[str, Any]]]:
        with self.lock:
            return list(self.requests)

    def counts(self) -> Counter[str]:
        return Counter(path for path, _ in self.recorded())

    def prompt_text(self) -> str:
        return "\n".join(
            str(message.get("content") or "")
            for path, request in self.recorded()
            if path == "/v1/chat/completions"
            for message in request.get("messages", [])
        )


@pytest.fixture
def local_models() -> Iterator[LocalModels]:
    service = LocalModels()

    class Handler(BaseHTTPRequestHandler):
        def do_POST(self) -> None:
            request = json.loads(self.rfile.read(int(self.headers.get("Content-Length", "0"))))
            with service.lock:
                service.requests.append((self.path, request))
            if self.path == "/v1/embeddings":
                response = {"data": [{"index": i, "embedding": VECTOR} for i, _ in enumerate(request["input"])]}
            elif self.path == "/rerank":
                response = {
                    "results": [
                        {"index": i, "relevance_score": 1.0}
                        for i in range(min(int(request["top_n"]), len(request["documents"])))
                    ]
                }
            elif self.path == "/v1/chat/completions":
                is_sufficiency = any(
                    message.get("role") == "system" and "Evidence Sufficiency Check" in str(message.get("content") or "")
                    for message in request.get("messages", [])
                )
                if is_sufficiency:
                    answer = service.sufficiency(request) if service.sufficiency is not None else {
                        "sufficient": False, "missing_facts": ["测试替身没有配置充分性响应"],
                        "supporting_evidence_ids": [], "reason": "Explicit localhost substitute abstention",
                    }
                else:
                    answer = service.answer or {
                        "answer_status": "not_found", "answer_value": "未找到", "confidence": 0.0,
                        "source_chunk_ids": [], "evidence_attachment_ids": [], "reference_source_documents": [],
                    }
                response = {"choices": [{"message": {"content": json.dumps(answer, ensure_ascii=False)}}]}
            else:
                self.send_error(404)
                return
            body = json.dumps(response).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, format: str, *args: Any) -> None:
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    service.base_url = f"http://127.0.0.1:{server.server_port}"
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield service
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)
        assert not thread.is_alive()


@dataclass
class Runtime:
    root: Path
    config: Path
    qdrant: Path
    models: LocalModels
    env: dict[str, str]

    @property
    def historical_items(self) -> Path:
        return self.root / "artifacts" / "12_gongkan_form_analysis" / "form_items.jsonl"

    def run(
        self,
        *extra: str | Path,
        out_dir: Path | None = None,
        rows: str = "all",
        target: str = TARGET,
        global_namespace: str = GLOBAL,
        room_context: str | None = None,
        parser_version: str | None = None,
    ) -> subprocess.CompletedProcess[str]:
        args = [
            "run-step15-agent", "--config", str(self.config),
            "--qdrant-path", str(self.qdrant), "--qdrant-collection", COLLECTION,
            "--target-namespace", target, "--global-namespace", global_namespace,
            "--rows", rows, "--out-dir", str(out_dir or self.root / "run"),
            "--no-judge", "--no-grounding-enabled", "--no-field-binding-enabled",
            "--chat-max-retries", "0", "--chat-retry-backoff-seconds", "0", "--timeout", "5",
            *map(str, extra),
        ]
        if room_context is not None:
            args.extend(["--room-context", room_context])
        if parser_version is None:
            command = [sys.executable, "-m", "nested_doc_rag.cli", *args]
        else:
            # Simulate only a parser release-version bump. All parsing and
            # runtime work remains in the actual production implementations.
            wrapper = (
                "import sys; "
                "import nested_doc_rag.form.template_parser as parser; "
                f"parser.PARSER_VERSION = {parser_version!r}; "
                "from nested_doc_rag.cli import main; main(sys.argv[1:])"
            )
            command = [sys.executable, "-c", wrapper, *args]
        return subprocess.run(
            command, cwd=self.root, env=self.env,
            capture_output=True, text=True, timeout=45, check=False,
        )

    def assert_completed(self, result: subprocess.CompletedProcess[str], *, count: int, out_dir: Path | None = None) -> None:
        assert result.returncode == 0, result.stdout + result.stderr
        output = out_dir or self.root / "run"
        manifest = read_json(output / "run_manifest.json")
        assert manifest["status"] == "completed", read_json(output / "trace_summary.json")
        assert manifest["counts"]["failed"] == 0, read_json(output / "trace_summary.json")
        assert manifest["counts"]["total_fields"] == count
        for path in ("/v1/embeddings", "/rerank", "/v1/chat/completions"):
            assert self.models.counts()[path] > 0, f"real model HTTP path was not used: {path}"
        predictions = read_jsonl(output / "predictions_raw.jsonl")
        assert len(predictions) == count
        assert len({prediction["field_id"] for prediction in predictions}) == count
        assert all(prediction["method_name"] != "step15_agent_failed" for prediction in predictions)


@pytest.fixture
def runtime(tmp_path: Path, local_models: LocalModels) -> Runtime:
    qdrant_path = tmp_path / "qdrant"
    client = QdrantClient(path=str(qdrant_path))
    try:
        client.create_collection(COLLECTION, vectors_config=models.VectorParams(size=3, distance=models.Distance.COSINE))
        client.upsert(
            COLLECTION,
            points=[models.PointStruct(id=1, vector=VECTOR, payload={
                "chunk_id": "anonymous_fact", "namespace": TARGET, "corpus_layer": "fact",
                "evidence_kind": "structured_field", "source_type": "uploaded_excel_row",
                "file_name": "匿名能力.xlsx", "relative_path": "匿名能力.xlsx",
                "sheet_name": "能力", "row_index": 2, "cell_range": "A2:C2",
                "field_name": "UPS容量", "field_value": "500kVA", "raw_text": "UPS容量：500kVA。",
            })],
            wait=True,
        )
    finally:
        # Disk Qdrant must release its lock before the separate CLI opens it.
        client.close()

    config = tmp_path / "runtime.yaml"
    config_data = {
        "paths": {
            "project_root": str(tmp_path), "data_dir": str(tmp_path / "data"),
            "artifacts_dir": str(tmp_path / "artifacts"), "qdrant_path": str(qdrant_path), "temp_dir": str(tmp_path / "tmp"),
        },
        "services": {
            "embedding_endpoint": local_models.base_url + "/v1/embeddings", "embedding_model": "local-embedding",
            "rerank_endpoint": local_models.base_url + "/rerank", "rerank_model": "local-rerank",
            "chat_endpoint": local_models.base_url + "/v1/chat/completions", "chat_model": "local-chat",
            "chat_api_key_env": "RUNTIME_CONTRACT_MODEL_KEY", "timeout_seconds": 5,
        },
        "qdrant": {"url": "", "collection_name": COLLECTION},
        "agentscope": {"enabled": False, "mode": "off"}, "agentic_mas": {"enabled": False},
        "grounding": {
            "evidence_strength_enabled": False, "field_binding_enabled": False, "field_binding_agent_enabled": False,
            "slot_decomposition_enabled": False, "pre_writeback_consistency_enabled": False,
            "relaxed_writeback_gate_enabled": False,
        },
        "retrieval": {"sufficiency_enabled": False, "expand_parent_payload": False, "vector_top_k": 3, "rerank_top_n": 2},
        "excel": {"template_path": str(tmp_path / "no-default-template.xlsx")},
    }
    config.write_text("\n".join(yaml_lines(config_data)) + "\n", encoding="utf-8")
    aliases = {
        "EMBEDDING_ENDPOINT", "EMBEDDING_MODEL", "RERANK_ENDPOINT", "RERANK_MODEL", "CHAT_ENDPOINT", "CHAT_MODEL",
        "CHAT_API_KEY_ENV", "DEEPSEEK_BASE_URL", "DEEPSEEK_MODEL", "QDRANT_COLLECTION", "QDRANT_URL",
        "QDRANT_API_KEY_ENV", "QDRANT_PATH", "TARGET_NAMESPACE", "GLOBAL_NAMESPACE", "RETRIEVAL_MODE",
        "WRITEBACK_ALLOW_UNCERTAIN", "WRITEBACK_EMBED_IMAGE", "DEEPSEEK_API_KEY", "QDRANT_API_KEY",
    }
    env = {key: value for key, value in os.environ.items() if key not in aliases and not key.startswith(("NDR_", "NESTED_DOC_RAG"))}
    env.update({
        "PYTHONPATH": str(REPO_ROOT / "src"), "RUNTIME_CONTRACT_MODEL_KEY": "local-test-only",
        "NDR_MODEL_GATEWAY_ENABLED": "false", "NESTED_DOC_RAG__QDRANT__URL": "",
        "NESTED_DOC_RAG__PATHS__PROJECT_ROOT": str(tmp_path),
    })
    return Runtime(tmp_path, config, qdrant_path, local_models, env)


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def yaml_lines(value: dict[str, Any], indent: int = 0) -> list[str]:
    """Use the production loader's supported YAML subset without PyYAML."""
    lines: list[str] = []
    for key, item in value.items():
        prefix = " " * indent + key + ":"
        if isinstance(item, dict):
            lines.extend([prefix, *yaml_lines(item, indent + 2)])
        else:
            lines.append(prefix + " " + json.dumps(item, ensure_ascii=False))
    return lines


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def write_items(path: Path, items: list[dict[str, Any]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(item, ensure_ascii=False) + "\n" for item in items), encoding="utf-8")
    return path


def explicit_item(row: int = 2, *, question: str = "显式市电路数", sheet: str = "显式调研") -> dict[str, Any]:
    return {
        "form_item_id": f"explicit_{sheet}_{row}", "file_name": "任意上传文件.xlsx", "sheet_name": sheet,
        "row_index": row, "target_cell": f"'{sheet}'!D{row}", "target_column_label": "实际情况",
        "question_text": question, "instruction_text": "请填写当前现网数据", "category_path": ["动力"],
        "needs_evidence": False, "answer_example": "",
    }


def make_template(
    path: Path,
    *,
    sheets: dict[str, list[tuple[int, str]]] | None = None,
    headers: tuple[str, ...] = DEFAULT_HEADERS,
) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, fields in (sheets or {"动力调研": [(2, "市电路数"), (3, "UPS容量")]}).items():
        worksheet = workbook.create_sheet(title)
        worksheet.append(list(headers))
        for row, question in fields:
            values = {"类别": "动力", "指标名称": question, "填写说明": f"请核对{question}", "备注": "保留人工备注"}
            for column, header in enumerate(headers, 1):
                worksheet.cell(row=row, column=column, value=values.get(header))
    workbook.save(path)
    workbook.close()
    return path


def protected_files(out_dir: Path) -> dict[str, bytes]:
    return {str(path.relative_to(out_dir)): path.read_bytes() for path in out_dir.rglob("*") if path.is_file()}


def assert_rejected_unchanged(
    runtime: Runtime,
    result: subprocess.CompletedProcess[str],
    before: dict[str, bytes],
    requests_before: list[tuple[str, dict[str, Any]]],
) -> None:
    assert result.returncode != 0, result.stdout + result.stderr
    error = (result.stderr + result.stdout).lower()
    assert "cannot resume" in error, error
    assert "changed" in error or "snapshot" in error, error
    assert "jsondecodeerror" not in error, "Input identity must be checked before reading a damaged prediction checkpoint"
    assert runtime.models.recorded() == requests_before, "Rejected resume contacted a model endpoint"
    assert protected_files(runtime.root / "run") == before, "Rejected resume overwrote an existing run artifact"


@pytest.mark.parametrize(
    ("filename", "headers", "fields"),
    [
        ("新上传动力信息.xlsx", DEFAULT_HEADERS, [(2, "市电路数"), (3, "UPS容量")]),
        ("重新排列_新版.xlsx", ("实际情况", "备注", "填写说明", "类别", "指标名称"), [(3, "UPS容量"), (7, "市电路数")]),
    ],
)
def test_uploaded_template_drives_actual_rows_columns_and_prompts(runtime: Runtime, filename, headers, fields) -> None:
    template = make_template(runtime.root / filename, sheets={"供电调研": fields}, headers=headers)
    original_template = template.read_bytes()
    result = runtime.run("--template", template)
    runtime.assert_completed(result, count=2)
    output = runtime.root / "run"
    items = read_jsonl(output / "form_items.jsonl")
    assert {(item["row_index"], item["question_text"]) for item in items} == set(fields)
    target_column = get_column_letter(headers.index("实际情况") + 1)
    assert {item["target_cell"].split("!")[-1] for item in items} == {f"{target_column}{row}" for row, _ in fields}
    assert all(item["file_name"] == filename and item["sheet_name"] == "供电调研" for item in items)
    assert all("!" in item["target_cell"] for item in items)
    predictions = read_jsonl(output / "predictions_raw.jsonl")
    assert {item["form_item_id"] for item in items} == {prediction["field_id"] for prediction in predictions}
    assert {item["target_cell"] for item in items} == {prediction["target_cell"] for prediction in predictions}
    for _, question in fields:
        assert question in runtime.models.prompt_text()
        assert f"请核对{question}" in runtime.models.prompt_text()
    snapshot = read_json(output / "form_input_snapshot.json")
    report = read_json(output / "form_parse_report.json")
    assert snapshot["input_mode"] == "template"
    assert snapshot["parser_name"] == report["parser_name"]
    assert snapshot["parser_version"] == report["parser_version"]
    assert snapshot["selected_field_count"] == 2
    assert read_json(output / "run_manifest.json")["form_input"] == snapshot
    assert not (output / "filled_form.xlsx").exists(), "Parsing must work independently of --writeback"
    assert not runtime.historical_items.exists()
    assert template.read_bytes() == original_template


def test_equal_row_numbers_on_multiple_sheets_remain_distinct(runtime: Runtime) -> None:
    template = make_template(runtime.root / "多工作表.xlsx", sheets={"动力 调研": [(2, "市电路数")], "网络调研": [(2, "网络出口数")]})
    result = runtime.run("--template", template, rows="2")
    runtime.assert_completed(result, count=2)
    items = read_jsonl(runtime.root / "run" / "form_items.jsonl")
    predictions = read_jsonl(runtime.root / "run" / "predictions_raw.jsonl")
    assert {item["sheet_name"] for item in items} == {"动力 调研", "网络调研"}
    assert {prediction["row_index"] for prediction in predictions} == {2}
    assert len({prediction["target_cell"] for prediction in predictions}) == 2
    assert all(prediction["target_cell"].split("!")[-1] == "D2" for prediction in predictions)
    assert read_json(runtime.root / "run" / "form_input_snapshot.json")["selected_field_count"] == 2


def test_all_means_all_fields_in_current_template_including_rows_after_144(runtime: Runtime) -> None:
    fields = [(2, "市电路数"), (3, "UPS容量"), (145, "网络出口数"), (146, "发电机数量")]
    template = make_template(runtime.root / "超过历史行范围.xlsx", sheets={"实际字段": fields})
    result = runtime.run("--template", template)
    runtime.assert_completed(result, count=4)
    assert {prediction["row_index"] for prediction in read_jsonl(runtime.root / "run" / "predictions_raw.jsonl")} == {2, 3, 145, 146}


def test_all_for_explicit_items_includes_row_one_and_arbitrary_filenames(runtime: Runtime) -> None:
    path = write_items(runtime.root / "explicit.jsonl", [explicit_item(row) for row in (1, 2, 145)])
    result = runtime.run("--form-items", path)
    runtime.assert_completed(result, count=3)
    assert {prediction["row_index"] for prediction in read_jsonl(runtime.root / "run" / "predictions_raw.jsonl")} == {1, 2, 145}
    assert read_json(runtime.root / "run" / "form_input_snapshot.json")["input_mode"] == "explicit_items"


def test_explicit_rows_filter_runtime_without_discarding_other_input_fields(runtime: Runtime) -> None:
    template = make_template(runtime.root / "行选择.xlsx", sheets={"动力": [(2, "市电路数"), (3, "UPS容量"), (145, "发电机数量")]})
    result = runtime.run("--template", template, rows="3,145")
    runtime.assert_completed(result, count=2)
    output = runtime.root / "run"
    items = read_jsonl(output / "form_items.jsonl")
    predictions = read_jsonl(output / "predictions_raw.jsonl")
    snapshot = read_json(output / "form_input_snapshot.json")
    assert {item["row_index"] for item in items} == {2, 3, 145}
    assert {prediction["row_index"] for prediction in predictions} == {3, 145}
    assert snapshot["selected_field_count"] == 2
    assert set(snapshot["selected_field_ids"]) == {prediction["field_id"] for prediction in predictions}
    assert read_json(output / "form_parse_report.json")["detected_fields"] == 3


def test_explicit_items_take_precedence_over_template_and_historical_artifact(runtime: Runtime) -> None:
    template = make_template(runtime.root / "不应选中的模板.xlsx")
    item = explicit_item(question="显式输入专属问题")
    path = write_items(runtime.root / "explicit.jsonl", [item])
    write_items(runtime.historical_items, [explicit_item(4, question="历史输入专属问题")])
    result = runtime.run("--form-items", path, "--template", template)
    runtime.assert_completed(result, count=1)
    assert read_jsonl(runtime.root / "run" / "form_items.jsonl") == [item]
    assert "显式输入专属问题" in runtime.models.prompt_text()
    assert "历史输入专属问题" not in runtime.models.prompt_text()
    assert "请核对市电路数" not in runtime.models.prompt_text()
    assert "请核对UPS容量" not in runtime.models.prompt_text()
    assert read_json(runtime.root / "run" / "form_input_snapshot.json")["input_mode"] == "explicit_items"
    assert "LegacyFormItemsFallbackWarning" not in result.stderr
    assert "LegacyIndexScopeWarning" in result.stderr


def test_template_takes_precedence_over_historical_artifact(runtime: Runtime) -> None:
    template = make_template(runtime.root / "新上传优先.xlsx", sheets={"当前模板": [(2, "当前模板专属问题")]})
    historical = write_items(runtime.historical_items, [explicit_item(4, question="历史输入专属问题")])
    historical_bytes = historical.read_bytes()
    result = runtime.run("--template", template)
    runtime.assert_completed(result, count=1)
    assert read_jsonl(runtime.root / "run" / "form_items.jsonl")[0]["question_text"] == "当前模板专属问题"
    assert "历史输入专属问题" not in runtime.models.prompt_text()
    assert "LegacyFormItemsFallbackWarning" not in result.stderr
    assert "LegacyIndexScopeWarning" in result.stderr
    assert historical.read_bytes() == historical_bytes


def test_missing_explicit_inputs_use_isolated_legacy_artifact_with_warning(runtime: Runtime) -> None:
    legacy = explicit_item(4, question="历史兼容字段")
    legacy["file_name"] = "基地云机房信息调研表.xlsx"
    write_items(runtime.historical_items, [legacy])
    result = runtime.run()
    runtime.assert_completed(result, count=1)
    assert "LegacyFormItemsFallbackWarning" in result.stderr, "Legacy form fallback must be externally observable"
    assert "LegacyIndexScopeWarning" in result.stderr
    assert read_jsonl(runtime.root / "run" / "form_items.jsonl") == [legacy]
    assert read_json(runtime.root / "run" / "form_input_snapshot.json")["input_mode"] == "legacy"


@pytest.mark.parametrize("invalid", ["missing", "malformed"])
def test_invalid_explicit_items_are_not_silently_replaced_by_template(runtime: Runtime, invalid: str) -> None:
    template = make_template(runtime.root / "有效模板.xlsx")
    path = runtime.root / "invalid.jsonl"
    if invalid == "malformed":
        path.write_text("{not json}\n", encoding="utf-8")
    write_items(runtime.historical_items, [explicit_item(4)])
    result = runtime.run("--form-items", path, "--template", template)
    assert result.returncode != 0
    assert not runtime.models.recorded()
    assert not (runtime.root / "run" / "predictions_raw.jsonl").exists()
    assert not (runtime.root / "run" / "form_input_snapshot.json").exists()


def test_ambiguous_template_stops_before_model_calls_and_archives_parse_diagnostics(runtime: Runtime) -> None:
    template = make_template(runtime.root / "歧义模板.xlsx", headers=("类别", "指标名称", "实际情况", "机房信息"))
    write_items(runtime.historical_items, [explicit_item(4)])
    result = runtime.run("--template", template)
    assert result.returncode != 0
    assert "ambiguous_target_columns" in result.stderr
    assert not runtime.models.recorded()
    output = runtime.root / "run"
    report = read_json(output / "form_parse_report.json")
    assert report["status"] == "error"
    assert any(error["reason"] == "ambiguous_target_columns" for error in report["errors"])
    assert not (output / "predictions_raw.jsonl").exists()
    assert not (output / "form_input_snapshot.json").exists()


def test_identical_input_resume_skips_completed_fields_without_duplicate_predictions(runtime: Runtime) -> None:
    template = make_template(runtime.root / "可恢复.xlsx", sheets={"动力": [(2, "市电路数")], "网络": [(2, "网络出口数")]})
    initial = runtime.run("--template", template)
    runtime.assert_completed(initial, count=2)
    output = runtime.root / "run"
    immutable = {name: (output / name).read_bytes() for name in ("form_items.jsonl", "form_parse_report.json", "form_input_snapshot.json", "predictions_raw.jsonl")}
    requests = runtime.models.recorded()
    result = runtime.run("--template", template, "--resume")
    runtime.assert_completed(result, count=2)
    assert runtime.models.recorded() == requests
    assert {name: (output / name).read_bytes() for name in immutable} == immutable
    state = read_json(output / "run_state.json")
    assert state["skipped_completed_count"] == 2
    assert state["fields_completed"] == 2
    assert len(read_jsonl(output / "predictions.checkpoint.jsonl")) == 2


def test_template_change_rejects_resume_before_loading_prediction_checkpoint(runtime: Runtime) -> None:
    template = make_template(runtime.root / "模板变更.xlsx", sheets={"动力": [(2, "市电路数")]})
    initial = runtime.run("--template", template)
    runtime.assert_completed(initial, count=1)
    workbook = load_workbook(template)
    workbook.active["C2"] = "修改后的填写标准：仅填写已投运线路"
    workbook.save(template)
    workbook.close()
    output = runtime.root / "run"
    (output / "predictions.checkpoint.jsonl").write_text("{invalid prediction checkpoint}\n", encoding="utf-8")
    before = protected_files(output)
    requests = runtime.models.recorded()
    result = runtime.run("--template", template, "--resume")
    assert_rejected_unchanged(runtime, result, before, requests)


def test_same_field_id_with_changed_explicit_question_rejects_resume(runtime: Runtime) -> None:
    item = explicit_item()
    path = write_items(runtime.root / "explicit.jsonl", [item])
    initial = runtime.run("--form-items", path)
    runtime.assert_completed(initial, count=1)
    item["question_text"] = "变更后的显式字段问题"
    write_items(path, [item])
    before = protected_files(runtime.root / "run")
    requests = runtime.models.recorded()
    result = runtime.run("--form-items", path, "--resume")
    assert_rejected_unchanged(runtime, result, before, requests)


def test_template_hash_is_checked_even_when_explicit_items_take_precedence(runtime: Runtime) -> None:
    template = make_template(runtime.root / "配套模板.xlsx", sheets={"动力": [(2, "市电路数")]})
    path = write_items(runtime.root / "explicit.jsonl", [explicit_item()])
    initial = runtime.run("--form-items", path, "--template", template)
    runtime.assert_completed(initial, count=1)
    workbook = load_workbook(template)
    workbook.active["E2"] = "人工备注已变更，显式字段集合保持不变"
    workbook.save(template)
    workbook.close()
    before = protected_files(runtime.root / "run")
    requests = runtime.models.recorded()
    result = runtime.run("--form-items", path, "--template", template, "--resume")
    assert_rejected_unchanged(runtime, result, before, requests)


def test_parser_release_version_change_rejects_resume(runtime: Runtime) -> None:
    template = make_template(runtime.root / "解析器升级.xlsx", sheets={"动力": [(2, "市电路数")]})
    initial = runtime.run("--template", template)
    runtime.assert_completed(initial, count=1)
    output = runtime.root / "run"
    old_version = read_json(output / "form_input_snapshot.json")["parser_version"]
    before = protected_files(output)
    requests = runtime.models.recorded()
    result = runtime.run("--template", template, "--resume", parser_version=old_version + ".contract-upgrade")
    assert_rejected_unchanged(runtime, result, before, requests)


@pytest.mark.parametrize("change", ["rows", "target", "global", "room_context"])
def test_selection_or_retrieval_scope_change_rejects_resume(runtime: Runtime, change: str) -> None:
    template = make_template(runtime.root / "范围变化.xlsx")
    initial = runtime.run("--template", template)
    runtime.assert_completed(initial, count=2)
    before = protected_files(runtime.root / "run")
    requests = runtime.models.recorded()
    options = {
        "rows": {"rows": "2"}, "target": {"target": "another_room"},
        "global": {"global_namespace": "another_global"}, "room_context": {"room_context": "机房范围已改为401"},
    }[change]
    result = runtime.run("--template", template, "--resume", **options)
    assert_rejected_unchanged(runtime, result, before, requests)


def test_historical_checkpoint_without_snapshot_cannot_be_resumed(runtime: Runtime) -> None:
    path = write_items(runtime.root / "explicit.jsonl", [explicit_item()])
    initial = runtime.run("--form-items", path)
    runtime.assert_completed(initial, count=1)
    output = runtime.root / "run"
    (output / "form_input_snapshot.json").unlink()
    before = protected_files(output)
    requests = runtime.models.recorded()
    result = runtime.run("--form-items", path, "--resume")
    assert_rejected_unchanged(runtime, result, before, requests)
