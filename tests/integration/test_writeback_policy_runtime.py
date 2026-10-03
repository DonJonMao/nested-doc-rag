"""Actual native ingest/run/writeback/validator CLIs with localhost models.

Disk Qdrant, template parsing, writer and artifact checks use production code.
Deterministic localhost embedding/rerank/chat replies verify policy contracts,
not real-model answer accuracy or a clean Docker upload workflow.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.comments import Comment
from test_runtime_form_contract import (
    Runtime,
    assert_rejected_unchanged,
    make_template,
    protected_files,
    read_json,
    read_jsonl,
)
from test_runtime_form_contract import (
    local_models as local_models,
)
from test_runtime_form_contract import (
    runtime as runtime,
)
from test_sufficiency_runtime_contract import (
    cli,
    completed,
    configure_answer,
    native_upload,
    semantic_response,
    validate_actual_cli,
)

HUMAN_VALUE = "人工确认现网值_仅本地模板_不得发送模型_780kVA"
HUMAN_COMMENT = "人工核对备注必须保留"
POLICIES = ("preserve", "overwrite_confirmed", "overwrite_all")


def setup_native_answer(runtime: Runtime) -> list[dict[str, Any]]:
    records = native_upload(runtime, primary_capacity=True)
    runtime.models.sufficiency = semantic_response
    configure_answer(runtime, records)
    return records


def template_with_value(runtime: Runtime, value: Any, *, name: str = "人工调研.xlsx") -> Path:
    template = make_template(runtime.root / name, sheets={"实际调研": [(2, "UPS容量")]})
    workbook = load_workbook(template)
    workbook["实际调研"]["D2"] = value
    workbook["实际调研"]["D2"].comment = Comment(HUMAN_COMMENT, "人工")
    workbook.save(template)
    workbook.close()
    return template


def cell_value(path: Path) -> Any:
    workbook = load_workbook(path, data_only=False)
    try:
        return workbook["实际调研"]["D2"].value
    finally:
        workbook.close()


def audit_value(record: dict[str, Any], *, old: Any, actual: Any, policy: str) -> None:
    assert record["old_value"] == old
    assert record["new_value"] == actual, "new_value must record the actual cell after the action"
    assert record["policy"] == policy
    assert record["answer_value"] == "500kVA"


@pytest.mark.parametrize("initial_value", [HUMAN_VALUE, "   ", 0])
def test_default_preserve_keeps_nonempty_human_values_without_leaking_them_to_models(runtime: Runtime, initial_value: Any) -> None:
    setup_native_answer(runtime)
    template = template_with_value(runtime, initial_value)
    original = template.read_bytes()
    result = runtime.run("--template", template, "--writeback", "--sufficiency-enabled")
    runtime.assert_completed(result, count=1)
    prediction, _ = completed(runtime)
    assert prediction["answer_status"] == "answered" and prediction["answer_value"] == "500kVA"
    audit = read_jsonl(runtime.root / "run" / "writeback_audit.jsonl")[0]
    assert audit["reason"] == "target_non_empty" and audit["error_code"] == "WB_TARGET_NON_EMPTY"
    assert audit["writeback_action"] != "written"
    audit_value(audit, old=initial_value, actual=initial_value, policy="preserve")
    filled = runtime.root / "run" / "filled_form.xlsx"
    assert cell_value(filled) == initial_value
    workbook = load_workbook(filled)
    try:
        assert workbook["实际调研"]["D2"].comment.text == HUMAN_COMMENT
    finally:
        workbook.close()
    assert any(review["reason"] == "target_non_empty" for review in read_jsonl(runtime.root / "run" / "review_items.jsonl"))
    if initial_value == HUMAN_VALUE:
        assert HUMAN_VALUE not in json.dumps(runtime.models.recorded(), ensure_ascii=False)
        assert HUMAN_VALUE not in (runtime.root / "run" / "form_items.jsonl").read_text(encoding="utf-8")
    assert template.read_bytes() == original
    manifest = read_json(runtime.root / "run" / "run_manifest.json")
    assert manifest["schema_version"] == "1.3"
    assert manifest["writeback"]["config"]["existing_value_policy"] == "preserve"
    assert manifest["writeback"]["overwrite_all_cli"] is False
    validate_actual_cli(runtime)


def test_default_preserve_writes_an_empty_confirmed_cell_and_records_actual_values(runtime: Runtime) -> None:
    setup_native_answer(runtime)
    template = template_with_value(runtime, None)
    result = runtime.run("--template", template, "--writeback", "--sufficiency-enabled")
    runtime.assert_completed(result, count=1)
    prediction, overlay = completed(runtime)
    assert prediction["answer_value"] == "500kVA" and overlay["writeback_allowed"]
    audit = read_jsonl(runtime.root / "run" / "writeback_audit.jsonl")[0]
    assert audit["writeback_action"] == "written" and audit["status"] == "confirmed"
    audit_value(audit, old=None, actual="500kVA", policy="preserve")
    assert cell_value(runtime.root / "run" / "filled_form.xlsx") == "500kVA"
    validate_actual_cli(runtime)


@pytest.mark.parametrize("policy", ["overwrite_confirmed", "overwrite_all"])
def test_explicit_overwrite_policies_replace_confirmed_ordinary_values(runtime: Runtime, policy: str) -> None:
    setup_native_answer(runtime)
    template = template_with_value(runtime, HUMAN_VALUE)
    original = template.read_bytes()
    result = runtime.run("--template", template, "--writeback", "--sufficiency-enabled", "--existing-value-policy", policy)
    runtime.assert_completed(result, count=1)
    prediction, _ = completed(runtime)
    assert prediction["answer_value"] == "500kVA"
    audit = read_jsonl(runtime.root / "run" / "writeback_audit.jsonl")[0]
    assert audit["writeback_action"] == "written" and audit["status"] == "confirmed"
    audit_value(audit, old=HUMAN_VALUE, actual="500kVA", policy=policy)
    assert cell_value(runtime.root / "run" / "filled_form.xlsx") == "500kVA"
    assert HUMAN_VALUE not in json.dumps(runtime.models.recorded(), ensure_ascii=False)
    assert template.read_bytes() == original
    contract = read_json(runtime.root / "run" / "form_input_snapshot.json")["acquisition_contract"]["writeback_policy"]
    assert contract == {"version": "writeback-policy-v1", "existing_value_policy": policy, "overwrite_all_cli": policy == "overwrite_all"}
    manifest = read_json(runtime.root / "run" / "run_manifest.json")
    assert manifest["writeback"]["config"]["existing_value_policy"] == policy
    assert manifest["writeback"]["overwrite_all_cli"] is (policy == "overwrite_all")
    validate_actual_cli(runtime)


def generated_predictions(runtime: Runtime) -> tuple[Path, Path]:
    setup_native_answer(runtime)
    template = template_with_value(runtime, None, name="仅生成答案.xlsx")
    result = runtime.run("--template", template, "--sufficiency-enabled")
    runtime.assert_completed(result, count=1)
    validate_actual_cli(runtime)
    return template, runtime.root / "run" / "predictions_raw.jsonl"


def standalone_writeback(
    runtime: Runtime, template: Path, predictions: Path, *, policy: str, suffix: str,
    authority: Path | None = None,
) -> tuple[Path, dict[str, Any]]:
    output = runtime.root / suffix
    output.mkdir()
    workbook = output / "filled.xlsx"
    args = [
        "writeback", "--template", str(template), "--pred", str(predictions), "--out", str(workbook),
        "--existing-value-policy", policy,
    ]
    if authority is not None:
        args.extend(["--retrieval-evidence", str(authority)])
    requests = runtime.models.recorded()
    result = cli(runtime, *args)
    assert result.returncode == 0, result.stdout + result.stderr
    assert runtime.models.recorded() == requests, "Standalone writer must not invoke models"
    return workbook, read_jsonl(output / "writeback_audit.jsonl")[0]


@pytest.mark.parametrize("policy", POLICIES)
def test_standalone_cli_preserves_formula_targets_under_every_policy(runtime: Runtime, policy: str) -> None:
    _, predictions = generated_predictions(runtime)
    template = template_with_value(runtime, "=1+1", name="公式保护.xlsx")
    original = template.read_bytes()
    filled, audit = standalone_writeback(runtime, template, predictions, policy=policy, suffix="formula-writer")
    assert audit["writeback_action"] == "skipped_formula" and audit["status"] != "confirmed"
    audit_value(audit, old="=1+1", actual="=1+1", policy=policy)
    assert cell_value(filled) == "=1+1" and template.read_bytes() == original


@pytest.mark.parametrize("damage", ["missing_authority", "partial_status"])
def test_cli_overwrite_all_cannot_bypass_evidence_or_status_gates(runtime: Runtime, damage: str) -> None:
    _, predictions = generated_predictions(runtime)
    template = template_with_value(runtime, HUMAN_VALUE, name="覆盖仍需证据.xlsx")
    authority = runtime.root / "run" / "retrieval_evidence.jsonl"
    if damage == "missing_authority":
        authority = runtime.root / "missing-authority.jsonl"
    else:
        rows = read_jsonl(predictions)
        rows[0].update(answer_status="partial_clue", answer_value="未找到", confidence=0.0)
        predictions = runtime.root / "partial-prediction.jsonl"
        predictions.write_text(json.dumps(rows[0], ensure_ascii=False) + "\n", encoding="utf-8")
    filled, audit = standalone_writeback(runtime, template, predictions, policy="overwrite_all", suffix="gated-writer", authority=authority)
    assert audit["writeback_action"] != "written" and audit["status"] != "confirmed"
    assert audit["old_value"] == audit["new_value"] == HUMAN_VALUE
    assert audit["policy"] == "overwrite_all"
    assert cell_value(filled) == HUMAN_VALUE
    if damage == "missing_authority":
        assert audit["error_code"] and "unresolvable_evidence" in audit["reason"]


@pytest.mark.parametrize("changed_policy", ["overwrite_confirmed", "overwrite_all"])
def test_policy_change_rejects_resume_before_models_or_archived_artifact_changes(runtime: Runtime, changed_policy: str) -> None:
    setup_native_answer(runtime)
    template = template_with_value(runtime, HUMAN_VALUE)
    initial = runtime.run("--template", template, "--writeback", "--sufficiency-enabled")
    runtime.assert_completed(initial, count=1)
    output = runtime.root / "run"
    (output / "predictions.checkpoint.jsonl").write_text("{damaged checkpoint}\n", encoding="utf-8")
    before, requests = protected_files(output), runtime.models.recorded()
    result = runtime.run("--template", template, "--writeback", "--sufficiency-enabled", "--resume", "--existing-value-policy", changed_policy)
    assert_rejected_unchanged(runtime, result, before, requests)
