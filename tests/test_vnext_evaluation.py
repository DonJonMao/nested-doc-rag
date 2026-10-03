from __future__ import annotations

import hashlib
import json
import runpy
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
EVAL = runpy.run_path(str(SCRIPTS / "vnext_evaluate.py"))
PREP = runpy.run_path(str(SCRIPTS / "vnext_prepare_old141.py"))


def write_rows(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")


def locator(*, file: str = "facts.xlsx", row: int = 2, text: str = "UPS容量 / 500kVA", namespace: str = "room", **extra) -> dict:
    return {"file_name": file, "namespace": namespace, "sheet_name": "事实", "cell_range": f"A{row}:B{row}",
            "row_index": row, "source_text": text, "source_text_hash": "sha256:" + hashlib.sha256(text.encode()).hexdigest(), **extra}


def evidence(*, chunk: str = "random-semantic-id", row: int = 2, text: str = "UPS容量 / 500kVA", **extra) -> dict:
    return {"chunk_id": chunk, "point_id": "unrelated-random-UUID", "file_name": "facts.xlsx", "namespace": "room",
            "sheet_name": "事实", "cell_range": f"A{row}:B{row}", "row_index": row, "raw_source_text": text,
            "retrieval_round": 0, **extra}


def gold(**extra) -> dict:
    return {"case_id": "C1", "dataset_id": "F1", "target": {"sheet_name": "填报", "cell": "D2"},
            "gold_verified": True,
            "expected_answer": "500kVA", "status": "answered", "answerable": True,
            "required_evidence": [locator(fact_id="capacity")], "writeback_eligible": True, **extra}


def run_data(tmp_path: Path, *, hits: list[dict] | None = None, prediction: dict | None = None, written: bool = False,
             value: str | None = "500kVA", acquisition: dict | None = None) -> dict:
    hits = hits if hits is not None else [evidence()]
    prediction = prediction if prediction is not None else {"field_id": "runtime-field", "target_cell": "'填报'!D2",
                "answer_status": "answered", "answer_value": "500kVA", "source_chunk_ids": [hits[0]["chunk_id"]],
                "validation": {"acquisition": acquisition} if acquisition else {}}
    write_rows(tmp_path / "predictions_raw.jsonl", [prediction])
    write_rows(tmp_path / "retrieval_evidence.jsonl", [{"field_id": prediction["field_id"], "top_hits": hits}])
    write_rows(tmp_path / "writeback_audit.jsonl", [{"field_id": prediction["field_id"], "target_cell": "填报!D2",
                "action": "written" if written else "skipped", "writeback_action": "written" if written else "review_only",
                "old_value": "人工保留", "new_value": "人工保留"}])
    if written:
        workbook = Workbook()
        workbook.active.title = "填报"
        workbook.active["D2"] = value
        workbook.save(tmp_path / "filled_form.xlsx")
        workbook.close()
    return EVAL["load_run"](tmp_path)


def evaluate(golds: list[dict], run: dict, **kwargs) -> dict:
    return EVAL["evaluate"](golds, run, method="A3", **kwargs)


def test_physical_gold_supports_arbitrary_runtime_ids_and_real_written_value(tmp_path: Path) -> None:
    report = evaluate([gold()], run_data(tmp_path, written=True))
    assert report["metrics"]["answer_accuracy"]["value"] == 1
    assert report["metrics"]["exact_field_evidence_recall_at_k"]["value"] == 1
    assert report["metrics"]["unsupported_answer_rate"]["value"] == 0
    assert report["metrics"]["writeback_precision"]["value"] == 1
    assert report["metrics"]["model_usage"]["value"] is None


@pytest.mark.parametrize("change", [{"raw_source_text": "UPS容量 / 1500kVA"}, {"namespace": "other-room"}, {"cell_range": "A9:B9"}])
def test_valid_ids_cannot_replace_gold_native_text_scope_or_address(tmp_path: Path, change: dict) -> None:
    bad = evidence(**change)
    report = evaluate([gold()], run_data(tmp_path, hits=[bad]))
    assert report["metrics"]["answer_accuracy"]["value"] == 1
    assert report["metrics"]["exact_field_evidence_recall_at_k"]["value"] == 0
    assert report["metrics"]["unsupported_answer_rate"]["value"] == 1
    assert "unsupported_answer" in report["cases"][0]["failures"]


def test_numeric_substrings_and_composite_first_numbers_are_not_correct_answers() -> None:
    assert not EVAL["answer_matches"](gold(), "1500kVA")
    assert not EVAL["answer_matches"](gold(), "500kVA + 800kVA")
    assert EVAL["answer_matches"](gold(), "0.5MVA")
    assert not EVAL["answer_matches"](gold(), "500kW")


@pytest.mark.parametrize("address", [{"table_index": 0, "row_index": 1}, {"paragraph_index": 0}])
def test_zero_word_physical_indices_are_real_addresses(address: dict) -> None:
    ref = {"file_name": "facts.docx", "namespace": "room", "source_text": "现网容量500kVA", **address}
    hit = {"file_name": "facts.docx", "namespace": "room", "raw_source_text": "现网容量500kVA", **address}
    assert EVAL["locator_matches"](ref, hit)
    hit.pop(next(iter(address)))
    assert not EVAL["locator_matches"](ref, hit)


def test_retrieval_recall_uses_ranked_hits_and_exact_field_is_separate(tmp_path: Path) -> None:
    wrong = evidence(chunk="wrong", row=3, text="变压器容量 / 500kVA")
    correct = evidence(chunk="correct")
    case = gold(decoy_evidence=[locator(row=3, text=wrong["raw_source_text"], reason="wrong_field")],
                relevant_evidence=[locator(row=3, text=wrong["raw_source_text"], fact_id="related"), locator(fact_id="related")])
    report = evaluate([case], run_data(tmp_path, hits=[wrong, correct]), k=1)
    assert report["metrics"]["evidence_recall_at_k"]["value"] == 1
    assert report["metrics"]["exact_field_evidence_recall_at_k"]["value"] == 0
    assert report["metrics"]["wrong_field_retrieval_rate"]["numerator"] == 1
    assert report["metrics"]["wrong_field_retrieval_rate"]["denominator"] == 1


def test_alternative_gold_locators_for_one_fact_do_not_inflate_recall_denominator(tmp_path: Path) -> None:
    case = gold(required_evidence=[locator(fact_id="capacity"), locator(row=3, fact_id="capacity")])
    report = evaluate([case], run_data(tmp_path))
    assert report["cases"][0]["required_fact_count"] == 1
    assert report["metrics"]["exact_field_evidence_recall_at_k"]["value"] == 1


def test_global_decoy_rank_and_namespace_mapping_use_independent_locators(tmp_path: Path) -> None:
    conflicting = evidence(chunk="global", row=3, text="UPS容量 / 1000kVA", namespace="global-runtime")
    target = evidence(chunk="target", namespace="room-runtime")
    case = gold(decoy_evidence=[locator(row=3, text=conflicting["raw_source_text"], namespace="global", reason="global_conflict")])
    report = evaluate([case], run_data(tmp_path, hits=[conflicting, target]), namespace_map={"room": "room-runtime", "global": "global-runtime"})
    assert report["metrics"]["target_vs_global_confusion_rate"]["value"] == 1
    assert report["metrics"]["exact_field_evidence_recall_at_k"]["value"] == 1
    assert report["metrics"]["unsupported_answer_rate"]["value"] == 1


def test_missing_and_failed_predictions_stay_in_accuracy_denominator(tmp_path: Path) -> None:
    run = run_data(tmp_path)
    run["predictions"] = [{"field_id": "runtime-field", "target_cell": "填报!D2", "answer_status": "not_found", "method_name": "step15_agent_failed"}]
    cases = [gold(), gold(case_id="C2", target={"sheet_name": "填报", "cell": "D3"}, status="not_found", answerable=False, required_evidence=[])]
    report = evaluate(cases, run)
    assert report["failed_case_count"] == 2
    assert report["metrics"]["answer_accuracy"]["denominator"] == 2
    assert report["metrics"]["answer_accuracy"]["numerator"] == 0
    assert report["metrics"]["abstention_precision"]["denominator"] == 0
    assert report["metrics"]["average_retrieval_calls_per_field"]["value"] is None
    assert report["metrics"]["average_retrieval_calls_per_field"]["unknown_count"] == 2


def test_gold_evidence_gain_and_raw_primary_answer_gain_are_separate(tmp_path: Path) -> None:
    primary = evidence(chunk="brand", row=2, text="UPS品牌 / 云衡")
    supplement = evidence(chunk="capacity", row=3, text="UPS容量 / 800kVA", retrieval_round=1)
    case = gold(expected_answer="云衡 / 800kVA", required_evidence=[locator(text=primary["raw_source_text"], fact_id="brand"),
                locator(row=3, text=supplement["raw_source_text"], fact_id="capacity")])
    acquisition = {"acquisition_rounds": 2, "qdrant_query_calls": 7}
    prediction = {"field_id": "runtime-field", "target_cell": "填报!D2", "answer_value": case["expected_answer"], "answer_status": "answered",
                  "source_chunk_ids": ["brand", "capacity"], "validation": {"acquisition": acquisition}}
    run = run_data(tmp_path, hits=[primary, supplement], prediction=prediction)
    first = [{**prediction, "answer_value": "未找到", "answer_status": "partial_clue", "source_chunk_ids": []}]
    report = evaluate([case], run, primary_predictions=first)
    assert report["metrics"]["second_round_evidence_gain"]["value"] == 1
    assert report["metrics"]["answer_gain_after_targeted"]["value"] == 1
    assert report["metrics"]["average_retrieval_calls_per_field"]["value"] == 7
    unknown = evaluate([case], run)
    assert unknown["metrics"]["answer_gain_after_targeted"]["value"] is None
    assert unknown["metrics"]["answer_gain_after_targeted"]["unknown_count"] == 1


def test_actual_workbook_detects_overwrite_even_when_audit_claims_no_change(tmp_path: Path) -> None:
    report = evaluate([gold(original_value="人工保留", writeback_eligible=False)], run_data(tmp_path, written=True))
    assert report["metrics"]["unsafe_overwrite_count"]["value"] == 1
    assert report["metrics"]["writeback_precision"]["value"] == 0
    assert "unsafe_overwrite" in report["cases"][0]["failures"]


def test_formula_safety_control_is_scored_without_inflating_challenge_count(tmp_path: Path) -> None:
    run = run_data(tmp_path, written=True)
    workbook = load_workbook(tmp_path / "filled_form.xlsx")
    workbook["填报"]["D9"] = 3
    workbook.save(tmp_path / "filled_form.xlsx")
    workbook.close()
    control = {"control_id": "FORMULA", "target": {"sheet_name": "填报", "cell": "D9"},
               "has_formula": True, "original_value": "=SUM(1,2)", "expected_output_value": "=SUM(1,2)"}
    report = evaluate([gold()], run, safety_controls=[control])
    assert report["case_count"] == 1
    assert report["metrics"]["unsafe_overwrite_count"]["value"] == 1
    assert report["safety_controls"][0]["preserved"] is False


def test_missing_final_workbook_and_unverified_heldout_are_unknown(tmp_path: Path) -> None:
    run = run_data(tmp_path, written=True)
    run["workbook_path"] = None
    report = evaluate([gold(gold_verified=False, original_value="人工保留")], run)
    assert report["metrics"]["answer_accuracy"]["value"] is None
    assert report["metrics"]["writeback_precision"]["value"] is None
    assert report["metrics"]["unsafe_overwrite_count"]["value"] is None


def test_unverified_annotations_do_not_score_quality_or_claim_support_with_real_workbook(tmp_path: Path) -> None:
    hit = evidence(retrieval_round=1)
    run = run_data(tmp_path, hits=[hit], written=True, acquisition={"acquisition_rounds": 2, "qdrant_query_calls": 3})
    case = gold(gold_verified=False, original_value="人工保留", expected_writeback_action="review_only",
                sufficiency_expected={"expected_acquisition_rounds": 1},
                decoy_evidence=[locator(reason="wrong_field"), locator(reason="global_conflict")])
    report = evaluate([case], run, primary_predictions=run["predictions"])
    for name in ("answer_accuracy", "evidence_recall_at_k", "exact_field_evidence_recall_at_k",
                 "wrong_field_retrieval_rate", "target_vs_global_confusion_rate", "unsupported_answer_rate",
                 "writeback_precision", "writeback_coverage", "second_round_evidence_gain", "answer_gain_after_targeted"):
        metric = report["metrics"][name]
        assert metric["value"] is None and metric["denominator"] == 0 and metric["unknown_count"] == 1, name
    assert report["cases"][0]["answer_correct"] is None
    assert report["cases"][0]["answer_supported"] is None
    assert report["cases"][0]["retrieved_fact_count_at_k"] is None
    assert report["cases"][0]["failures"] == []
    assert report["metrics"]["second_round_trigger_rate"]["value"] == 1
    assert report["metrics"]["average_retrieval_calls_per_field"]["value"] == 3
    assert report["metrics"]["unsafe_overwrite_count"]["value"] is None
    assert report["metrics"]["unsafe_overwrite_count"]["unknown_count"] == 1
    assert report["gold_assessment"]["unverified_cases"] == 1
    assert "No independently verified gold quality claim" in report["quality_claim"]


def test_unverified_nonanswerability_and_empty_evidence_gain_are_unknown(tmp_path: Path) -> None:
    report = evaluate([gold(gold_verified=False, answerable=False, required_evidence=[])],
                      run_data(tmp_path, acquisition={"acquisition_rounds": 2, "qdrant_query_calls": 2}))
    assert report["metrics"]["unsupported_answer_rate"]["value"] is None
    assert report["metrics"]["unsupported_answer_rate"]["unknown_count"] == 1
    assert report["metrics"]["second_round_evidence_gain"]["value"] is None
    assert report["metrics"]["second_round_evidence_gain"]["unknown_count"] == 1
    assert "unsupported_answer" not in report["cases"][0]["failures"]


@pytest.mark.parametrize("flag", [None, False, 0, 1, "true"])
def test_only_boolean_true_verifies_a_gold_declaration(tmp_path: Path, flag) -> None:
    report = evaluate([gold(gold_verified=flag)], run_data(tmp_path))
    assert report["metrics"]["answer_accuracy"]["value"] is None
    assert report["metrics"]["evidence_recall_at_k"]["value"] is None
    assert report["gold_assessment"]["verified_declaration_cases"] == 0


def test_missing_verification_does_not_promote_arbitrary_candidate_but_keeps_frozen_domain_gold(tmp_path: Path) -> None:
    case = gold()
    del case["gold_verified"]
    run = run_data(tmp_path)
    unknown = evaluate([case], run)
    assert unknown["metrics"]["answer_accuracy"]["value"] is None
    assert unknown["metrics"]["exact_field_evidence_recall_at_k"]["value"] is None
    frozen = {**case, "schema_version": "vnext-evaluation-dataset-v1", "gold_origin": "declared_synthetic_domain_rules"}
    assert evaluate([frozen], run)["metrics"]["answer_accuracy"]["value"] == 1
    assert evaluate([frozen], run)["metrics"]["exact_field_evidence_recall_at_k"]["value"] == 1
    assert evaluate([{**frozen, "gold_verified": False}], run)["metrics"]["answer_accuracy"]["value"] is None


@pytest.mark.parametrize("extra", [
    {"gold_origin": "legacy_heldout_unverified"}, {"review_state": "candidate_material_only"},
    {"review_state": "agent_native_review_partial"}, {"evaluator_gold_exported": False},
    {"semantic_review": {"quality_gold_eligible": False}},
])
def test_candidate_or_heldout_markers_cannot_be_promoted_by_adding_true(tmp_path: Path, extra: dict) -> None:
    report = evaluate([gold(**extra)], run_data(tmp_path))
    assert report["gold_assessment"]["verified_declaration_cases"] == 0
    assert report["cases"][0]["answer_supported"] is None


@pytest.mark.parametrize("case", [
    {"review_schema_version": "old141-native-review-v1", "field": {"row_index": 4}},
    gold(review_schema_version="old141-native-review-v1"),
])
def test_review_material_schema_is_rejected_even_after_flattening(case: dict, tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="candidate review materials are not evaluator gold"):
        evaluate([case], run_data(tmp_path))


@pytest.mark.parametrize("missing", ["file_name", "namespace", "cell_range", "source_text"])
def test_locator_needs_independent_native_identity_scope_address_and_text(tmp_path: Path, missing: str) -> None:
    ref = locator()
    if missing == "source_text":
        ref.pop("source_text_hash")
    ref.pop(missing)
    report = evaluate([gold(required_evidence=[ref])], run_data(tmp_path, written=True,
                       acquisition={"acquisition_rounds": 2, "qdrant_query_calls": 2}))
    assert report["metrics"]["answer_accuracy"]["value"] == 1  # independently declared answer label
    for name in ("evidence_recall_at_k", "exact_field_evidence_recall_at_k", "unsupported_answer_rate",
                 "writeback_precision", "second_round_evidence_gain"):
        assert report["metrics"][name]["value"] is None, name
        assert report["metrics"][name]["unknown_count"] == 1, name
    assert report["cases"][0]["answer_supported"] is None


def test_status_is_not_required_for_independently_verified_retrieval_annotations(tmp_path: Path) -> None:
    case = gold()
    del case["status"]
    report = evaluate([case], run_data(tmp_path))
    assert report["metrics"]["answer_accuracy"]["value"] is None
    assert report["metrics"]["exact_field_evidence_recall_at_k"]["value"] == 1
    assert report["gold_assessment"]["independent_evidence_cases"] == 1


def test_verified_unsupported_classification_still_requires_known_answerability(tmp_path: Path) -> None:
    case = gold()
    del case["answerable"]
    report = evaluate([case], run_data(tmp_path, written=True))
    assert report["metrics"]["unsupported_answer_rate"]["value"] is None
    assert report["metrics"]["writeback_precision"]["value"] is None
    assert report["cases"][0]["answer_supported"] is None


def test_separately_verified_input_protection_survives_unknown_semantic_gold(tmp_path: Path) -> None:
    run = run_data(tmp_path, written=True)
    case = gold(gold_verified=False, original_value="人工保留", original_value_verified=True)
    report = evaluate([case], run)
    assert report["metrics"]["answer_accuracy"]["value"] is None
    assert report["metrics"]["unsafe_overwrite_count"]["value"] == 1
    assert "unsafe_overwrite" in report["cases"][0]["failures"]
    control = {"control_id": "KNOWN_INPUT", "target": {"sheet_name": "填报", "cell": "D2"},
               "original_value": "人工保留", "expected_output_value": "人工保留"}
    separate = evaluate([gold(gold_verified=False)], run, safety_controls=[control])
    assert separate["metrics"]["unsafe_overwrite_count"]["value"] == 1
    assert separate["safety_controls"][0]["protection_verified"] is True
    unknown = evaluate([gold(gold_verified=False)], run, safety_controls=[{**control, "gold_verified": False}])
    assert unknown["metrics"]["unsafe_overwrite_count"]["value"] is None
    assert unknown["safety_controls"][0]["preserved"] is None


def test_raw_round_invariant_is_retained_when_gold_expectation_is_unverified(tmp_path: Path) -> None:
    report = evaluate([gold(gold_verified=False, sufficiency_expected={"expected_acquisition_rounds": 1})],
                      run_data(tmp_path, acquisition={"acquisition_rounds": 3, "qdrant_query_calls": 4}))
    assert report["cases"][0]["failures"] == ["too_many_acquisition_rounds"]


def test_mixed_verified_and_unverified_cases_expose_only_verified_quality_population(tmp_path: Path) -> None:
    run = run_data(tmp_path)
    pred = {**run["predictions"][0], "field_id": "unverified", "target_cell": "填报!D3"}
    run["predictions"].append(pred)
    run["authority"].append({"field_id": "unverified", "top_hits": [evidence()]})
    report = evaluate([gold(), gold(case_id="C2", target={"sheet_name": "填报", "cell": "D3"}, gold_verified=False)], run)
    for name in ("answer_accuracy", "exact_field_evidence_recall_at_k", "unsupported_answer_rate"):
        assert report["metrics"][name]["denominator"] == 1
        assert report["metrics"][name]["unknown_count"] == 1
    assert report["gold_assessment"]["verified_declaration_cases"] == 1
    assert report["gold_assessment"]["unverified_cases"] == 1
    assert report["cases"][1]["answer_supported"] is None


@pytest.mark.parametrize("models_kind", ["real", "stub"])
def test_cli_and_markdown_do_not_relabel_unverified_cases_as_fixed_gold_quality(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, models_kind: str,
) -> None:
    run_dir, out_dir = tmp_path / "run", tmp_path / "report"
    run_data(run_dir, written=True, acquisition={"acquisition_rounds": 2, "qdrant_query_calls": 3})
    gold_path = tmp_path / "unverified.jsonl"
    write_rows(gold_path, [gold(gold_verified=False)])
    monkeypatch.setattr("sys.argv", [str(SCRIPTS / "vnext_evaluate.py"), "--gold", str(gold_path),
                        "--run-dir", str(run_dir), "--out-dir", str(out_dir), "--models-kind", models_kind])
    EVAL["main"]()
    report = json.loads((out_dir / "report.json").read_text())
    assert report["gold_assessment"]["quality_metrics_assessable"] is False
    assert "No independently verified gold quality claim" in report["quality_claim"]
    assert report["metrics"]["writeback_coverage"]["denominator"] == 0
    assert report["metrics"]["second_round_trigger_rate"]["value"] == 1
    assert report["quality_claim"] in (out_dir / "report.md").read_text()
    assert report["input_hashes"]["gold.jsonl"] == EVAL["digest"](gold_path)


def test_duplicate_runtime_targets_fail_instead_of_silently_selecting_last(tmp_path: Path) -> None:
    run = run_data(tmp_path)
    run["predictions"].append(deepcopy(run["predictions"][0]))
    with pytest.raises(ValueError, match="ambiguous"):
        evaluate([gold()], run)


def test_historical_observation_sidecar_is_used_only_when_native_authority_absent(tmp_path: Path) -> None:
    native = run_data(tmp_path)
    sidecar = tmp_path.parent / f"{tmp_path.name}-observed.jsonl"
    write_rows(sidecar, [{"field_id": "runtime-field", "target_cell": "填报!D2", "top_hits": [evidence(namespace="wrong-room")]}])
    preferred = EVAL["load_run"](tmp_path, retrieval_authority=sidecar)
    assert preferred["authority"] == native["authority"]
    assert preferred["authority_source"] == "native_run_artifact"
    (tmp_path / "retrieval_evidence.jsonl").unlink()
    observed = EVAL["load_run"](tmp_path, retrieval_authority=sidecar)
    assert observed["authority_source"] == "external_observed_sidecar"
    assert observed["input_hashes"]["external_retrieval_authority"] == EVAL["digest"](sidecar)
    assert evaluate([gold()], observed)["metrics"]["unsupported_answer_rate"]["value"] == 1


def test_legacy_plain_cell_and_row_can_join_a_qualified_gold_target() -> None:
    case = {"case_id": "old-4", "row_index": 4, "target": {"sheet_name": "原表", "cell": "G4"}}
    record = {"row_index": 4, "target_cell": "G4"}
    assert EVAL["join_record"](case, [record]) == record


def test_old141_preparation_removes_fact_examples_and_preserves_original_package(tmp_path: Path) -> None:
    original = tmp_path / "original"
    form = original / "data/工勘单" / PREP["OLD_FILE"]
    form.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.title = "原表"
    workbook.active["A1"] = "保留布局"
    workbook.active["G4"] = "人工事实4"
    workbook.active["G5"] = "=SUM(1,2)"
    rows = []
    for row in range(4, 145):
        if row != 5:
            workbook.active[f"G{row}"] = f"人工事实{row}"
        rows.append({"form_item_id": f"old-{row}", "file_name": PREP["OLD_FILE"], "sheet_name": "原表", "row_index": row,
                     "target_cell": f"G{row}", "question_text": "字段问题", "instruction_text": "填写现网", "existing_value": f"人工事实{row}",
                     "current_info": f"人工事实{row}", "answer_example": f"人工事实{row}", "suggested_retrieval_query": f"泄漏人工事实{row}"})
    workbook.save(form)
    workbook.close()
    items = original / "artifacts/12_gongkan_form_analysis/form_items.jsonl"
    write_rows(items, rows)
    before = {path: path.read_bytes() for path in (form, items)}
    out = tmp_path / "prepared"
    report = PREP["prepare"](original, out)
    assert report["field_count"] == 141 and report["preserved_formula_count"] == 1
    assert report["cleared_target_count"] == 140
    assert before == {path: path.read_bytes() for path in before}
    closed = EVAL["read_jsonl"](out / "form_items_closed_book.jsonl")
    assert all(row["answer_example"] == "" for row in closed)
    assert not any(any(key in row for key in ("existing_value", "current_info", "heldout_answer", "suggested_retrieval_query")) for row in closed)
    assert len(EVAL["read_jsonl"](out / "heldout_answers.jsonl")) == 141
    blank = load_workbook(out / "blank_target_template.xlsx", data_only=False)
    try:
        assert blank["原表"]["G4"].value is None
        assert blank["原表"]["G5"].value == "=SUM(1,2)"
        assert blank["原表"]["A1"].value == "保留布局"
    finally:
        blank.close()
    with ZipFile(form) as source, ZipFile(out / "blank_target_template.xlsx") as result:
        assert source.namelist() == result.namelist()
        for name in source.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert source.read(name) == result.read(name)
    with pytest.raises(ValueError, match="outside"):
        PREP["prepare"](original, original / "forbidden")


def test_old37_replay_observes_counts_without_relabeling_history_as_quality(tmp_path: Path) -> None:
    predictions = [{"field_id": f"field-{n}", "answer_status": "answered" if n < 44 else "not_found"} for n in range(141)]
    write_rows(tmp_path / "predictions_raw.jsonl", predictions)
    write_rows(tmp_path / "eval_results.jsonl", [{"row_index": n} for n in range(141)])
    write_rows(tmp_path / "writeback_audit.jsonl", [{"action": "written" if n < 37 else "skipped", "status": "confirmed" if n < 37 else "flagged"} for n in range(141)])
    report = EVAL["old37_replay"](tmp_path)
    assert report["compatible_old141_read"] and report["historical_writeback37_observed"]
    assert report["quality_metrics"] is None and report["models_kind"] == "no_model_calls"
