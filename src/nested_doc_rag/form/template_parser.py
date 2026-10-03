"""Deterministic parsing of the workbook submitted for a single form run.

Only explicit header roles establish question and writable target locations.
Existing target values are held out of every input used to generate answers.
"""
from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import quote_sheetname

from nested_doc_rag.io import write_json, write_jsonl

PARSER_NAME = "native_excel_form_template"
PARSER_VERSION = "1"
TITLE_HEADER_LOOKAHEAD_ROWS = 5
PROOF_TERMS = ("证明", "材料", "截图", "照片", "图纸", "报告", "备案", "证书", "资质", "验收", "记录", "CAD", "PDF")
EMPTY_QUESTION_MARKERS = {"", "/", "\\", "-", "—", "／"}


@dataclass(frozen=True)
class TemplateParseResult:
    items: list[dict[str, Any]]
    report: dict[str, Any]


class TemplateParseError(ValueError):
    def __init__(self, message: str, *, report: dict[str, Any]) -> None:
        super().__init__(message)
        self.report = report


def _text(value: Any) -> str:
    return "" if value is None else re.sub(r"\s+", " ", str(value)).strip()


def _header_role(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    key = re.sub(r"\s+", "", value).replace("（", "(").replace("）", ")").casefold().strip("：:")
    if any(token in key for token in ("示例", "范例", "样例", "example", "sample")):
        return "example"
    key = re.sub(r"\([^)]*\)", "", key)
    aliases = {
        "category": {"类别", "分类", "一级分类", "专业", "系统", "工勘项", "category"},
        "subcategory": {"子类", "子分类", "二级分类", "分组", "group", "subcategory"},
        "question": {"指标名称", "指标", "字段", "字段名", "名称", "项目", "问题", "问题描述", "题目", "子项", "条目", "确认项", "评估内容", "检查项", "参数名称", "field", "fieldname", "field_name", "question", "questiontext", "item"},
        "target": {"答案", "应答", "答复", "实际情况", "实际值", "当前实际值", "当前值", "值", "填写值", "填写结果", "填写内容", "待填写", "用户填写", "机房信息", "实际信息", "机房现状", "当前信息", "现状", "满足情况", "answer", "value", "response", "actualvalue", "target", "targetvalue"},
        "instruction": {"填写说明", "填写说明及标准", "填写要求", "说明", "要求", "标准", "参考内容", "instruction", "instructions", "requirement"},
    }
    for role, names in aliases.items():
        if key in names:
            return role
    if key.startswith(("填写说明", "填写要求")):
        return "instruction"
    return None


class _SheetView:
    def __init__(self, sheet: Any) -> None:
        self.sheet = sheet
        self.merges: dict[str, Any] = {}
        for merged in sheet.merged_cells.ranges:
            for row in range(merged.min_row, merged.max_row + 1):
                for column in range(merged.min_col, merged.max_col + 1):
                    self.merges[f"{get_column_letter(column)}{row}"] = merged

    def cell(self, row: int, column: int) -> Any:
        cell = self.sheet.cell(row, column)
        merged = self.merges.get(cell.coordinate)
        return self.sheet.cell(merged.min_row, merged.min_col) if merged else cell

    def role(self, row: int, column: int) -> str | None:
        cell = self.cell(row, column)
        return None if cell.data_type == "f" else _header_role(cell.value)


def _header_band(view: _SheetView) -> tuple[int, int] | None:
    sheet = view.sheet
    for row in range(1, sheet.max_row + 1):
        if not any(view.role(row, column) in {"question", "target"} for column in range(1, sheet.max_column + 1)):
            continue
        occupied = [cell for cell in sheet[row] if cell.value is not None]
        if len(occupied) <= 1 and any(
            {"question", "target"}.issubset({view.role(later, column) for column in range(1, sheet.max_column + 1)})
            for later in range(row + 1, min(sheet.max_row, row + TITLE_HEADER_LOOKAHEAD_ROWS) + 1)
        ):
            # A single title such as "机房信息" must not become a target
            # header spanning the actual field header below it.
            continue
        end = row
        for merged in sheet.merged_cells.ranges:
            if merged.min_row <= row <= merged.max_row and view.role(row, merged.min_col):
                end = max(end, merged.max_row)
        # A horizontal parent header may have explicit child labels on the next
        # row (for example example/actual value). Those child roles take priority.
        horizontal_parent = any(
            merged.min_row == row and merged.max_col > merged.min_col and view.role(row, merged.min_col)
            for merged in sheet.merged_cells.ranges
        )
        next_row = end + 1
        if horizontal_parent and next_row <= sheet.max_row and any(
            _header_role(cell.value) in {"question", "target", "example", "instruction"}
            for cell in sheet[next_row] if cell.data_type != "f"
        ):
            end = next_row
        return row, end
    return None


def _header_columns(view: _SheetView, start: int, end: int) -> dict[str, list[dict[str, Any]]]:
    columns: dict[int, tuple[str, str]] = {}
    for row in range(start, end + 1):
        for column in range(1, view.sheet.max_column + 1):
            role = view.role(row, column)
            if role:
                columns[column] = (role, view.cell(row, column).coordinate)
    grouped: dict[tuple[str, str], list[int]] = defaultdict(list)
    for column, identity in columns.items():
        grouped[identity].append(column)
    output: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for (role, origin), indices in grouped.items():
        output[role].append({"header_cell": origin, "columns": indices, "letters": [get_column_letter(index) for index in indices]})
    return dict(output)


def _diagnostic(sheet: Any, row: int | None, reason: str, **details: Any) -> dict[str, Any]:
    return {"sheet_name": sheet.title, "row_index": row, "reason": reason, **details}


def _column_signature(layout: dict[str, list[dict[str, Any]]]) -> dict[str, list[list[int]]]:
    return {role: [definition["columns"] for definition in definitions] for role, definitions in layout.items()}


def _read_role(view: _SheetView, row: int, definition: dict[str, Any]) -> tuple[Any | None, str | None]:
    columns = definition["columns"]
    cells = {view.cell(row, column).coordinate: view.cell(row, column) for column in columns}
    if len(cells) != 1:
        return None, "unmerged_header_range"
    cell = next(iter(cells.values()))
    if cell.column not in columns:
        return None, "merged_role_overlap"
    return cell, None


def _update_categories(
    view: _SheetView, row: int, definitions: list[dict[str, Any]], categories: dict[str, str], report: dict[str, Any],
) -> list[Any]:
    cells = []
    for index, definition in enumerate(definitions):
        if not any(_text(view.cell(row, column).value) for column in definition["columns"]):
            continue
        cell, reason = _read_role(view, row, definition)
        if reason:
            report["errors"].append(_diagnostic(view.sheet, row, reason, role="category"))
            continue
        assert cell is not None
        cells.append(cell)
        value = "" if cell.data_type == "f" else _text(cell.value)
        key = definition["header_cell"]
        if value and value != categories.get(key):
            for child in definitions[index + 1:]:
                categories.pop(child["header_cell"], None)
            categories[key] = value
    return cells


def _parse_sheet(
    view: _SheetView, *, template: Path, sheet_index: int, report: dict[str, Any], include_heldout_answers: bool,
) -> list[dict[str, Any]]:
    sheet = view.sheet
    band = _header_band(view)
    sheet_report: dict[str, Any] = {"sheet_name": sheet.title, "sheet_index": sheet_index, "detected_fields": 0}
    report["sheets"].append(sheet_report)
    if band is None:
        sheet_report["status"] = "skipped"
        sheet_report["reason"] = "no_explicit_form_header"
        return []
    start, end = band
    columns = _header_columns(view, start, end)
    sheet_report.update({"header_rows": list(range(start, end + 1)), "column_roles": columns})
    for role in ("question", "target"):
        candidates = columns.get(role) or []
        if len(candidates) != 1:
            reason = f"ambiguous_{role}_columns" if candidates else f"missing_{role}_column"
            error = _diagnostic(sheet, start, reason, candidates=candidates)
            report["errors"].append(error)
            report["ambiguous_rows"].append(error)
    if any(error["sheet_name"] == sheet.title for error in report["errors"]):
        sheet_report["status"] = "error"
        return []
    question_definition, target_definition = columns["question"][0], columns["target"][0]
    target_columns = target_definition["letters"]
    report["target_column_by_sheet"][sheet.title] = target_columns[0] if len(target_columns) == 1 else ":".join((target_columns[0], target_columns[-1]))
    category_definitions = sorted(
        columns.get("category", []) + columns.get("subcategory", []), key=lambda definition: definition["columns"][0],
    )
    categories: dict[str, str] = {}
    used_targets: dict[str, int] = {}
    items: list[dict[str, Any]] = []
    for row in range(end + 1, sheet.max_row + 1):
        repeated_columns = _header_columns(view, row, row)
        if repeated_columns.get("question") and repeated_columns.get("target"):
            if _column_signature(repeated_columns) != _column_signature(columns):
                error = _diagnostic(sheet, row, "inconsistent_header_layout", column_roles=repeated_columns)
                report["errors"].append(error)
                report["ambiguous_rows"].append(error)
            else:
                report["skipped_rows"].append(_diagnostic(sheet, row, "repeated_header"))
            categories.clear()
            continue
        category_cells = _update_categories(view, row, category_definitions, categories, report)
        if all(_text(view.cell(row, column).value) in EMPTY_QUESTION_MARKERS for column in question_definition["columns"]):
            report["skipped_rows"].append(_diagnostic(sheet, row, "empty_question"))
            continue
        question_cell, reason = _read_role(view, row, question_definition)
        if reason:
            report["errors"].append(_diagnostic(sheet, row, reason, role="question"))
            continue
        assert question_cell is not None
        question = _text(question_cell.value)
        if question in EMPTY_QUESTION_MARKERS:
            report["skipped_rows"].append(_diagnostic(sheet, row, "empty_question"))
            continue
        if question_cell.data_type == "f":
            report["skipped_rows"].append(_diagnostic(sheet, row, "formula_question", cell=question_cell.coordinate))
            continue
        target_cell, reason = _read_role(view, row, target_definition)
        if reason:
            error = _diagnostic(sheet, row, reason, role="target", columns=target_definition["letters"])
            report["errors"].append(error)
            report["ambiguous_rows"].append(error)
            continue
        assert target_cell is not None
        if target_cell.data_type == "f":
            report["skipped_rows"].append(_diagnostic(sheet, row, "formula_target", target_cell=target_cell.coordinate))
            continue
        if target_cell.coordinate in used_targets:
            error = _diagnostic(sheet, row, "shared_merged_target", target_cell=target_cell.coordinate, other_row=used_targets[target_cell.coordinate])
            report["errors"].append(error)
            report["ambiguous_rows"].append(error)
            continue
        prompt_cells = [question_cell, *category_cells]
        instructions: list[str] = []
        examples: list[str] = []
        for role, values in (("instruction", instructions), ("example", examples)):
            for definition in columns.get(role) or []:
                cell, reason = _read_role(view, row, definition)
                if reason:
                    report["errors"].append(_diagnostic(sheet, row, reason, role=role))
                    continue
                assert cell is not None
                prompt_cells.append(cell)
                if cell.data_type == "f":
                    report["warnings"].append(_diagnostic(sheet, row, "formula_prompt_cell_skipped", role=role, cell=cell.coordinate))
                elif _text(cell.value):
                    values.append(_text(cell.value))
        if target_cell.coordinate in {cell.coordinate for cell in prompt_cells}:
            report["errors"].append(_diagnostic(sheet, row, "target_overlaps_prompt", target_cell=target_cell.coordinate))
            continue
        category_path = list(dict.fromkeys(categories[definition["header_cell"]] for definition in category_definitions if categories.get(definition["header_cell"])))
        instruction = " / ".join(instructions)
        example = " / ".join(examples) or None
        item: dict[str, Any] = {
            "file_name": template.name, "relative_path": template.name,
            "sheet_index": sheet_index, "sheet_name": sheet.title, "row_index": row,
            "target_cell": f"{quote_sheetname(sheet.title)}!{target_cell.coordinate}", "category_path": category_path,
            "target_column_label": _text(sheet[target_definition["header_cell"]].value),
            "question_text": question, "instruction_text": instruction, "answer_example": example,
            "needs_evidence": any(term.casefold() in " ".join((question, instruction, example or "")).casefold() for term in PROOF_TERMS),
            "parser_name": PARSER_NAME, "parser_version": PARSER_VERSION,
        }
        identity = {key: item[key] for key in ("sheet_name", "target_cell", "category_path", "question_text", "instruction_text", "answer_example", "parser_version")}
        item["form_item_id"] = "form_" + hashlib.sha256(json.dumps(identity, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
        item["suggested_retrieval_query"] = " ".join(value for value in [*category_path, question, instruction, example or ""] if value)
        if include_heldout_answers:
            item["heldout_answer"] = _text(target_cell.value)
        used_targets[target_cell.coordinate] = row
        items.append(item)
    sheet_report.update({"detected_fields": len(items), "status": "error" if any(error["sheet_name"] == sheet.title for error in report["errors"]) else "parsed"})
    return items


def parse_form_template(template: Path, *, include_heldout_answers: bool = False) -> TemplateParseResult:
    template = Path(template)
    report: dict[str, Any] = {
        "parser_name": PARSER_NAME, "parser_version": PARSER_VERSION, "template_file_name": template.name,
        "include_heldout_answers": include_heldout_answers, "sheet_count": 0, "detected_fields": 0,
        "sheets": [], "ambiguous_rows": [], "skipped_rows": [], "target_column_by_sheet": {}, "errors": [], "warnings": [],
    }
    if template.suffix.casefold() not in {".xlsx", ".xlsm"}:
        report["errors"].append({"reason": "unsupported_template_format", "suffix": template.suffix})
        report["status"] = "error"
        raise TemplateParseError("form template must be an .xlsx or .xlsm workbook", report=report)
    try:
        template_bytes = template.read_bytes()
        report["template_sha256"] = "sha256:" + hashlib.sha256(template_bytes).hexdigest()
        workbook = load_workbook(BytesIO(template_bytes), read_only=False, data_only=False)
    except Exception as exc:
        report["errors"].append({"reason": "workbook_read_failed", "detail": str(exc)})
        report["status"] = "error"
        raise TemplateParseError(f"cannot read form template: {exc}", report=report) from exc
    try:
        report["sheet_count"] = len(workbook.worksheets)
        items = [
            item for sheet_index, sheet in enumerate(workbook.worksheets, 1)
            for item in _parse_sheet(_SheetView(sheet), template=template, sheet_index=sheet_index, report=report, include_heldout_answers=include_heldout_answers)
        ]
    finally:
        workbook.close()
    report["detected_fields"] = len(items)
    if not items and not report["errors"]:
        report["errors"].append({"reason": "no_form_fields", "detail": "no rows have an explicit question and writable target"})
    report["status"] = "error" if report["errors"] else "parsed"
    if report["errors"]:
        reasons = ", ".join(dict.fromkeys(error["reason"] for error in report["errors"]))
        raise TemplateParseError(f"form template has unresolved structure: {reasons}", report=report)
    return TemplateParseResult(items=items, report=report)


def parse_template_to_run_dir(template: Path, out_dir: Path, *, include_heldout_answers: bool = False) -> Path:
    """Standalone convenience writer; resumed runs validate their snapshot first."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        result = parse_form_template(template, include_heldout_answers=include_heldout_answers)
    except TemplateParseError as exc:
        write_json(out_dir / "form_parse_report.json", exc.report)
        raise
    write_json(out_dir / "form_parse_report.json", result.report)
    path = out_dir / "form_items.jsonl"
    write_jsonl(path, result.items)
    return path
