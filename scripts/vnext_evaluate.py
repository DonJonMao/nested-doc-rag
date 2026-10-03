#!/usr/bin/env python3
"""Offline, locator-grounded Phase 5 evaluation. Never contacts a model service."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

ABSTENTIONS = {"not_found", "partial_clue", "conflict_unresolved"}
WRITTEN = {"written", "written_red_comment"}


def gold_declaration_verified(gold: dict[str, Any]) -> bool:
    """Trust an explicit gold declaration, never a candidate or heldout label.

    The frozen synthetic dataset predates gold_verified, but already declares
    its versioned domain-rule origin. No human signoff is required for it.
    """
    if gold.get("gold_origin") == "legacy_heldout_unverified" or gold.get("review_state") in {
        "candidate_material_only", "agent_native_review_partial",
    } or gold.get("evaluator_gold_exported") is False or (
        isinstance(gold.get("semantic_review"), dict)
        and gold["semantic_review"].get("quality_gold_eligible") is False
    ):
        return False
    if "gold_verified" in gold:
        return gold["gold_verified"] is True
    return gold.get("schema_version") == "vnext-evaluation-dataset-v1" and gold.get("gold_origin") == "declared_synthetic_domain_rules"


def independent_locator(locator: dict[str, Any]) -> bool:
    """A selected ID alone cannot supply source identity, scope or native text."""
    native_hash = str(locator.get("source_text_hash") or "").removeprefix("sha256:")
    has_hash = re.fullmatch(r"[0-9a-fA-F]{64}", native_hash) is not None and native_hash != hashlib.sha256(b"").hexdigest()
    has_text = isinstance(locator.get("source_text"), str) and bool(locator["source_text"].strip())
    document_hash = str(locator.get("document_sha256") or "").removeprefix("sha256:")
    has_identity = any(isinstance(locator.get(key), str) and locator[key].strip() for key in ("file_name", "relative_path")) or re.fullmatch(r"[0-9a-fA-F]{64}", document_hash) is not None
    address = {**locator, **(locator.get("address") or {})}
    has_address = (
        bool(address.get("sheet_name") and address.get("cell_range"))
        or (type(address.get("table_index")) is int and address["table_index"] >= 0
            and type(address.get("row_index")) is int and address["row_index"] >= 0)
        or (type(address.get("paragraph_index")) is int and address["paragraph_index"] >= 0)
        or bool(address.get("source_anchor"))
    )
    return bool(has_identity and isinstance(locator.get("namespace"), str) and locator["namespace"].strip()
                and has_address and (has_text or has_hash))


def independent_groups(groups: dict[str, list[dict[str, Any]]]) -> bool:
    return bool(groups) and all(any(independent_locator(locator) for locator in locators) for locators in groups.values())


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    records = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    if any(not isinstance(record, dict) for record in records):
        raise ValueError(f"JSONL requires object records: {path}")
    return records


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalized(value: Any) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(value if value is not None else "")).casefold())


def answer_matches(gold: dict[str, Any], actual: Any) -> bool:
    expected = gold.get("expected_answer", gold.get("expected_value"))
    accepted = [expected, *(gold.get("accepted_answers") or gold.get("accepted_aliases") or [])]
    if any(normalized(value) == normalized(actual) for value in accepted):
        return True
    # Numeric equality must consume the whole value. A 500 substring in 1500,
    # or the first number in a composite answer, is never semantic support.
    number = re.compile(r"([-+]?\d+(?:\.\d+)?)([a-z%℃°]*)")
    units = {"mw": ("kw", 1000), "kw": ("kw", 1), "w": ("kw", .001),
             "mva": ("kva", 1000), "kva": ("kva", 1)}
    left, right = number.fullmatch(normalized(expected)), number.fullmatch(normalized(actual))
    if not left or not right:
        return False
    lu, lm = units.get(left[2], (left[2], 1))
    ru, rm = units.get(right[2], (right[2], 1))
    return lu == ru and math.isclose(float(left[1]) * lm, float(right[1]) * rm, rel_tol=1e-6, abs_tol=1e-6)


def target_key(record: dict[str, Any]) -> tuple[str, str] | None:
    target = record.get("target") or {}
    if target.get("sheet_name") and target.get("cell"):
        return str(target["sheet_name"]), str(target["cell"]).replace("$", "").upper()
    value = record.get("target_cell")
    if isinstance(value, str) and "!" in value:
        sheet, cell = value.rsplit("!", 1)
        return sheet.strip("'").replace("''", "'"), cell.replace("$", "").upper()
    if isinstance(value, str) and record.get("sheet_name"):
        return str(record["sheet_name"]), value.replace("$", "").upper()
    if record.get("sheet_name") and record.get("cell"):
        return str(record["sheet_name"]), str(record["cell"]).replace("$", "").upper()
    return None


def join_record(gold: dict[str, Any], records: list[dict[str, Any]]) -> dict[str, Any] | None:
    key = target_key(gold)
    matches = [row for row in records if key is not None and target_key(row) == key]
    if not matches and gold.get("field_id"):
        matches = [row for row in records if row.get("field_id") == gold["field_id"]]
    if not matches and gold.get("row_index") is not None:
        matches = [row for row in records if row.get("row_index") == gold["row_index"] and (
            key is None or (target_key(row) is None and row.get("target_cell") == key[1]))]
    if len(matches) > 1:
        raise ValueError(f"ambiguous run identity for gold case {gold['case_id']}")
    return matches[0] if matches else None


def hit_values(hit: dict[str, Any]) -> dict[str, Any]:
    """Expose physical aliases without promoting scores or embedding text."""
    source = hit.get("source") if isinstance(hit.get("source"), dict) else {}
    address = hit.get("address") if isinstance(hit.get("address"), dict) else {}
    local = source.get("local_anchor") if isinstance(source.get("local_anchor"), dict) else {}
    value = {**local, **source, **hit, **{k: v for k, v in address.items() if v not in (None, "")}}
    value["cell_range"] = value.get("cell_range") or value.get("cell") or value.get("source_cell")
    value["source_text"] = hit.get("raw_source_text", source.get("raw_source_text", hit.get("source_text", hit.get("raw_text"))))
    value["document_sha256"] = value.get("source_document_hash") or value.get("document_hash")
    value["source_anchor"] = value.get("source_anchor") or value.get("anchor")
    return value


def locator_matches(locator: dict[str, Any], hit: dict[str, Any], namespace_map: dict[str, str] | None = None) -> bool:
    if not independent_locator(locator):
        return False
    value = hit_values(hit)
    expected = {**locator, **(locator.get("address") or {})}
    filename = expected.get("file_name")
    relative = expected.get("relative_path")
    if not filename and not relative and not expected.get("document_sha256"):
        return False
    if filename and str(value.get("file_name") or value.get("source_document") or "") != str(filename):
        return False
    if relative and str(value.get("relative_path") or "") != str(relative):
        return False
    physical = ("sheet_name", "cell_range", "table_index", "row_index", "paragraph_index", "source_anchor")
    if not any(expected.get(key) is not None for key in physical[1:]):
        return False
    for key in physical:
        actual = value.get(key)
        if expected.get(key) is not None and (actual is None or str(actual) != str(expected[key])):
            return False
    if expected.get("namespace") is not None:
        namespace = (namespace_map or {}).get(expected["namespace"], expected["namespace"])
        if value.get("namespace") != namespace:
            return False
    if expected.get("document_sha256") and str(value.get("document_sha256") or "").removeprefix("sha256:") != str(expected["document_sha256"]).removeprefix("sha256:"):
        return False
    raw = value.get("source_text")
    if expected.get("source_text") is not None:
        if not isinstance(raw, str):
            return False
        if expected.get("source_text_contains", False):
            if str(expected["source_text"]) not in raw:
                return False
        elif raw != expected["source_text"]:
            return False
    if expected.get("source_text_hash"):
        if not isinstance(raw, str) or hashlib.sha256(raw.encode("utf-8")).hexdigest() != str(expected["source_text_hash"]).removeprefix("sha256:"):
            return False
    return True


def evidence_groups(gold: dict[str, Any], *, key: str = "required_evidence") -> dict[str, list[dict[str, Any]]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for index, locator in enumerate(gold.get(key) or []):
        groups[str(locator.get("fact_id") or index)].append(locator)
    return groups


def matched_facts(gold: dict[str, Any], hits: list[dict[str, Any]], namespace_map: dict[str, str], *, key: str = "required_evidence") -> set[str]:
    return {fact for fact, locators in evidence_groups(gold, key=key).items()
            if any(locator_matches(locator, hit, namespace_map) for locator in locators for hit in hits)}


def metric(numerator: float | None, denominator: int, *, population: str, reason: str | None = None, unknown: int = 0) -> dict[str, Any]:
    return {"value": numerator / denominator if numerator is not None and denominator else None,
            "numerator": numerator, "denominator": denominator, "population": population,
            "unknown_count": unknown, "reason": reason or ("empty denominator" if not denominator else None)}


def load_run(run_dir: Path, *, retrieval_authority: Path | None = None) -> dict[str, Any]:
    raw = run_dir / "predictions_raw.jsonl"
    predictions = read_jsonl(raw if raw.is_file() else run_dir / "predictions.jsonl")
    evaluations = read_jsonl(run_dir / "eval_results.jsonl")
    native_authority = run_dir / "retrieval_evidence.jsonl"
    authority_path = native_authority if native_authority.is_file() else retrieval_authority
    authority = read_jsonl(authority_path) if authority_path is not None else []
    trace = read_jsonl(run_dir / "trace.jsonl")
    audit_path = run_dir / "writeback_audit.jsonl"
    workbook_path = run_dir / "filled_form.xlsx"
    usage_path = run_dir / "model_usage.json"
    artifacts = [raw, run_dir / "predictions.jsonl", run_dir / "eval_results.jsonl", run_dir / "retrieval_evidence.jsonl",
                 run_dir / "trace.jsonl", audit_path, run_dir / "run_manifest.json", workbook_path, usage_path]
    return {"predictions": predictions, "evaluations": evaluations, "authority": authority, "trace": trace,
            "audit": read_jsonl(audit_path), "audit_available": audit_path.is_file(),
            "workbook_path": workbook_path if workbook_path.is_file() else None,
            "usage": json.loads(usage_path.read_text()) if usage_path.is_file() else None,
            "authority_source": "native_run_artifact" if native_authority.is_file() else ("external_observed_sidecar" if retrieval_authority else "historical_eval_or_missing"),
            "input_hashes": {**{path.name: digest(path) for path in artifacts if path.is_file()},
                             **({"external_retrieval_authority": digest(authority_path)} if authority_path is not None and authority_path != native_authority else {})}}


def retrieval_for(gold: dict[str, Any], prediction: dict[str, Any] | None, run: dict[str, Any]) -> tuple[list[dict[str, Any]], bool]:
    rows = [row for row in run["authority"] if prediction and row.get("field_id") == prediction.get("field_id")]
    if len(rows) > 1:
        raise ValueError(f"duplicate retrieval authority for {gold['case_id']}")
    if rows:
        if not isinstance(rows[0].get("top_hits"), list):
            raise ValueError("retrieval authority top_hits must be a list")
        return rows[0]["top_hits"], True
    historical = join_record(gold, run["evaluations"])
    if historical is not None and isinstance(historical.get("top_hits"), list):
        return historical["top_hits"], True
    return [], False


def selected_hits(prediction: dict[str, Any], hits: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], bool]:
    ids = {str(value) for value in prediction.get("source_chunk_ids") or []}
    selected = [hit for hit in hits if str(hit.get("chunk_id") or hit.get("evidence_id") or "") in ids]
    # IDs only identify the chosen pack entries. Support is checked against
    # independent gold locators and native source text below.
    present = {str(hit.get("chunk_id") or hit.get("evidence_id") or "") for hit in selected}
    return selected, bool(ids) and present == ids


def acquisition_for(prediction: dict[str, Any], run: dict[str, Any]) -> dict[str, Any] | None:
    acquisition = (prediction.get("validation") or {}).get("acquisition")
    if isinstance(acquisition, dict):
        return acquisition
    candidates = [(row.get("payload") or {}).get("acquisition") for row in run["trace"]
                  if row.get("field_id") == prediction.get("field_id") and row.get("step") == "layered_retrieval_finished"]
    return next((value for value in candidates if isinstance(value, dict)), None)


def evaluate(golds: list[dict[str, Any]], run: dict[str, Any], *, method: str, k: int = 5,
             namespace_map: dict[str, str] | None = None, primary_predictions: list[dict[str, Any]] | None = None,
             safety_controls: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    if any("review_schema_version" in gold for gold in golds):
        raise ValueError("candidate review materials are not evaluator gold; export independently approved annotations separately")
    if k < 1 or not golds or any(not gold.get("case_id") for gold in golds) or len({gold["case_id"] for gold in golds}) != len(golds):
        raise ValueError("evaluation requires positive K and unique, nonempty gold cases")
    namespace_map = namespace_map or {}
    workbook_values: dict[tuple[str, str], Any] = {}
    if run.get("workbook_path") is not None:
        from openpyxl import load_workbook
        workbook = load_workbook(run["workbook_path"], read_only=True, data_only=False)
        try:
            for gold in [*golds, *(safety_controls or [])]:
                key = target_key(gold)
                if key is not None and key[0] in workbook.sheetnames:
                    workbook_values[key] = workbook[key[0]][key[1]].value
        finally:
            workbook.close()
    rows: list[dict[str, Any]] = []
    sums: Counter[str] = Counter()
    denominators: Counter[str] = Counter()
    unknowns: Counter[str] = Counter()
    for gold in golds:
        prediction = join_record(gold, run["predictions"])
        hits, retrieval_available = retrieval_for(gold, prediction, run)
        prediction = prediction or {}
        missing = not prediction
        failed = missing or prediction.get("method_name") == "step15_agent_failed"
        status = prediction.get("answer_status")
        groups = evidence_groups(gold)
        expected_status = gold.get("status", gold.get("expected_status"))
        declaration_verified = gold_declaration_verified(gold)
        verified = declaration_verified and expected_status in {"answered", *ABSTENTIONS} and (
            expected_status != "answered" or gold.get("expected_answer", gold.get("expected_value")) is not None)
        evidence_verified = declaration_verified and independent_groups(groups)
        answerable = gold.get("answerable")
        correct_status = status in (gold.get("accepted_statuses") or [expected_status])
        correct = not failed and correct_status and (status != "answered" or answer_matches(gold, prediction.get("answer_value")))
        selected, ids_valid = selected_hits(prediction, hits)
        facts = matched_facts(gold, selected, namespace_map)
        supported = verified and evidence_verified and ids_valid and set(groups) <= facts and answerable is True and answer_matches(gold, prediction.get("answer_value"))
        support_assessable = declaration_verified and (answerable is False or (
            answerable is True and evidence_verified and retrieval_available and verified))
        failures = []
        if missing:
            failures.append("missing_prediction")
        elif failed:
            failures.append("field_failed")
        if verified:
            denominators["answer_accuracy"] += 1
            sums["answer_accuracy"] += int(correct)
            if not correct:
                failures.append("answer_or_status_mismatch")
        else:
            unknowns["answer_accuracy"] += 1
        if not failed and status in ABSTENTIONS:
            if answerable is not None and verified:
                denominators["abstention_precision"] += 1
                sums["abstention_precision"] += int(not answerable)
            else:
                unknowns["abstention_precision"] += 1
        if status == "answered":
            if support_assessable:
                denominators["unsupported_answer_rate"] += 1
                sums["unsupported_answer_rate"] += int(not supported)
                if not supported:
                    failures.append("unsupported_answer")
            else:
                unknowns["unsupported_answer_rate"] += 1
        top = hits[:k]
        recalled = matched_facts(gold, top, namespace_map)
        relevant_key = "relevant_evidence" if gold.get("relevant_evidence") else "required_evidence"
        relevant_groups = evidence_groups(gold, key=relevant_key)
        relevant_recalled = matched_facts(gold, top, namespace_map, key=relevant_key)
        relevant_verified = declaration_verified and independent_groups(relevant_groups)
        if relevant_groups:
            if relevant_verified and (retrieval_available or failed):
                denominators["evidence_recall_at_k"] += 1
                sums["evidence_recall_at_k"] += len(relevant_recalled) / len(relevant_groups)
            else:
                unknowns["evidence_recall_at_k"] += 1
        if groups:
            if evidence_verified and (retrieval_available or failed):
                denominators["exact_field_evidence_recall_at_k"] += 1
                # Gold support locators explicitly bind field, room, status and
                # address; same-valued decoys cannot earn exact-field credit.
                sums["exact_field_evidence_recall_at_k"] += len(recalled) / len(groups)
            else:
                unknowns["exact_field_evidence_recall_at_k"] += 1
        decoys = gold.get("decoy_evidence") or []
        if any(locator.get("reason") == "wrong_field" for locator in decoys):
            field_labels_verified = declaration_verified and any(
                locator.get("reason") == "wrong_field" and independent_locator(locator) for locator in decoys)
            for hit in top:
                support_hit = evidence_verified and any(locator_matches(locator, hit, namespace_map) for locators in groups.values() for locator in locators)
                reasons = {locator.get("reason") for locator in decoys if locator_matches(locator, hit, namespace_map)}
                if field_labels_verified and (support_hit or reasons):
                    denominators["wrong_field_retrieval_rate"] += 1
                    sums["wrong_field_retrieval_rate"] += int("wrong_field" in reasons)
                else:
                    unknowns["wrong_field_retrieval_rate"] += 1
        if any(locator.get("reason") == "global_conflict" for locator in decoys):
            conflict_verified = evidence_verified and any(
                locator.get("reason") == "global_conflict" and independent_locator(locator) for locator in decoys)
            if conflict_verified and (retrieval_available or failed):
                denominators["target_vs_global_confusion_rate"] += 1
                # Retrieval confusion: a contradictory global fact precedes
                # every target support in the returned ranked pack.
                labels = []
                for hit in top:
                    if any(locator_matches(locator, hit, namespace_map) for locators in groups.values() for locator in locators):
                        labels.append("target")
                    elif any(locator.get("reason") == "global_conflict" and locator_matches(locator, hit, namespace_map) for locator in decoys):
                        labels.append("global")
                sums["target_vs_global_confusion_rate"] += int(bool(labels) and labels[0] == "global")
            else:
                unknowns["target_vs_global_confusion_rate"] += 1
        audit = join_record(gold, run["audit"])
        wrote = audit is not None and (audit.get("action") == "written" or audit.get("writeback_action") in WRITTEN)
        eligible = gold.get("writeback_eligible")
        if eligible is not None:
            if not declaration_verified or type(eligible) is not bool:
                unknowns["writeback_coverage"] += 1
            elif run["audit_available"]:
                if eligible:
                    if wrote and target_key(gold) not in workbook_values:
                        unknowns["writeback_coverage"] += 1
                    else:
                        denominators["writeback_coverage"] += 1
                        sums["writeback_coverage"] += int(wrote)
            else:
                unknowns["writeback_coverage"] += int(bool(eligible))
        if wrote:
            actual_available = target_key(gold) in workbook_values
            actual_correct = actual_available and answer_matches(gold, workbook_values[target_key(gold)])
            if verified and actual_available and support_assessable:
                denominators["writeback_precision"] += 1
                sums["writeback_precision"] += int(correct and supported and actual_correct and eligible is not False)
                if not (correct and supported and actual_correct and eligible is not False):
                    failures.append("unsafe_or_incorrect_writeback")
            else:
                unknowns["writeback_precision"] += 1
        protected = gold.get("has_formula") is True or gold.get("original_value") not in (None, "")
        protection_verified = gold.get("original_value_verified", declaration_verified) is True
        if protected:
            key = target_key(gold)
            if not protection_verified or key not in workbook_values:
                unknowns["unsafe_overwrite_count"] += 1
            else:
                changed = gold.get("original_value") != workbook_values[key]
                permitted = gold.get("has_formula") is not True and gold.get("overwrite_authorized") is True
                denominators["unsafe_overwrite_count"] += 1
                sums["unsafe_overwrite_count"] += int(changed and not permitted)
                if changed and not permitted:
                    failures.append("unsafe_overwrite")
        expected_action = gold.get("expected_writeback_action", gold.get("writeback_action"))
        if declaration_verified and expected_action is not None and audit is not None and audit.get("writeback_action") != expected_action:
            failures.append("writeback_action_mismatch")
        acquisition = acquisition_for(prediction, run)
        rounds = acquisition.get("acquisition_rounds") if acquisition else None
        if type(rounds) is int:
            denominators["second_round_trigger_rate"] += 1
            sums["second_round_trigger_rate"] += int(rounds == 2)
            if rounds > 2:
                failures.append("too_many_acquisition_rounds")
            expected_rounds = (gold.get("sufficiency_expected") or {}).get("expected_acquisition_rounds")
            if declaration_verified and expected_rounds is not None and method in {"A3", "A4"} and rounds != expected_rounds:
                failures.append("acquisition_round_expectation_mismatch")
            if rounds == 2:
                if not evidence_verified or not retrieval_available or any("retrieval_round" not in hit for hit in hits):
                    unknowns["second_round_evidence_gain"] += 1
                else:
                    primary = [hit for hit in hits if hit["retrieval_round"] == 0]
                    gain = len(matched_facts(gold, hits, namespace_map) - matched_facts(gold, primary, namespace_map))
                    denominators["second_round_evidence_gain"] += 1
                    sums["second_round_evidence_gain"] += gain
                previous = join_record(gold, primary_predictions or [])
                if previous is not None and verified and retrieval_available:
                    prior_correct = previous.get("answer_status") in (gold.get("accepted_statuses") or [expected_status]) and (
                        previous.get("answer_status") != "answered" or answer_matches(gold, previous.get("answer_value")))
                    denominators["answer_gain_after_targeted"] += 1
                    sums["answer_gain_after_targeted"] += int(correct) - int(prior_correct)
                else:
                    unknowns["answer_gain_after_targeted"] += 1
        else:
            unknowns["second_round_trigger_rate"] += 1
        calls = acquisition.get("qdrant_query_calls") if acquisition else None
        if type(calls) is int and calls >= 0:
            denominators["average_retrieval_calls_per_field"] += 1
            sums["average_retrieval_calls_per_field"] += calls
        else:
            unknowns["average_retrieval_calls_per_field"] += 1
        rows.append({"case_id": gold["case_id"], "dataset_id": gold.get("dataset_id"), "category": gold.get("category"),
                     "gold_declaration_verified": declaration_verified, "answer_label_verified": verified,
                     "evidence_annotations_verified": evidence_verified, "protection_verified": protection_verified if protected else None,
                     "prediction_present": not missing, "field_failed": failed, "answer_correct": correct if verified else None,
                     "answer_supported": supported if status == "answered" and support_assessable else None,
                     "required_fact_count": len(groups), "retrieved_fact_count_at_k": len(recalled) if evidence_verified and (retrieval_available or failed) else None,
                     "retrieval_available": retrieval_available, "written": wrote if run["audit_available"] else None,
                     "acquisition_rounds": rounds, "qdrant_query_calls": calls, "failures": failures})
    control_rows = []
    for control in safety_controls or []:
        key = target_key(control)
        observed = workbook_values.get(key)
        available = key in workbook_values
        protection_verified = control.get("original_value_verified", control.get("gold_verified", True)) is True
        safe = observed == control.get("expected_output_value", control.get("original_value")) if available and protection_verified else None
        if control.get("has_formula") or control.get("original_value") not in (None, ""):
            if available and protection_verified:
                denominators["unsafe_overwrite_count"] += 1
                sums["unsafe_overwrite_count"] += int(not safe)
            else:
                unknowns["unsafe_overwrite_count"] += 1
        control_rows.append({"control_id": control["control_id"], "target": control["target"], "preserved": safe,
                             "protection_verified": protection_verified})
    populations = {
        "evidence_recall_at_k": "macro required-fact recall in ranked retrieval top K; cases with gold evidence",
        "exact_field_evidence_recall_at_k": "macro exact field/scope/status/native locator recall in ranked retrieval top K",
        "wrong_field_retrieval_rate": "gold-classified top K hits in cases with wrong-field decoys; unclassified hits reported unknown",
        "target_vs_global_confusion_rate": "target/global-conflict cases; contradictory global evidence ranks ahead of target",
        "answer_accuracy": "all verified gold cases, including failed and missing predictions",
        "abstention_precision": "actual nonfailed abstentions with gold answerability; includes conflict_unresolved",
        "unsupported_answer_rate": "answered fields with independently assessable support",
        "writeback_precision": "written fields with independently assessable correctness and support",
        "writeback_coverage": "gold fields eligible for writeback under the fixed policy",
        "second_round_trigger_rate": "fields with raw acquisition-round measurements",
        "second_round_evidence_gain": "supplemented fields; mean new independently gold-supported facts",
        "answer_gain_after_targeted": "supplemented fields with raw primary-only predictions; signed correctness gain",
        "average_retrieval_calls_per_field": "fields with actual instrumented Qdrant calls, including retry attempts",
    }
    metrics = {name: metric(sums[name], denominators[name], population=population, unknown=unknowns[name])
               for name, population in populations.items()}
    metrics["unsafe_overwrite_count"] = {"value": sums["unsafe_overwrite_count"] if not unknowns["unsafe_overwrite_count"] else None,
                                         "observed_count": sums["unsafe_overwrite_count"], "denominator": denominators["unsafe_overwrite_count"],
                                         "unknown_count": unknowns["unsafe_overwrite_count"], "population": "protected gold targets compared with the actual filled workbook"}
    usage_path = run.get("usage")
    metrics["model_usage"] = {"value": usage_path, "reason": None if usage_path is not None else "raw provider usage not archived; unknown, never zero"}
    if not primary_predictions:
        metrics["answer_gain_after_targeted"]["reason"] = "raw primary-only predictions not supplied"
    quality_assessable = any(value["denominator"] for name, value in metrics.items()
                             if name not in {"unsafe_overwrite_count", "second_round_trigger_rate", "average_retrieval_calls_per_field", "model_usage"})
    return {"schema_version": "vnext-evaluation-v1", "method": method, "k": k, "case_count": len(golds),
            "failed_case_count": sum(row["field_failed"] for row in rows), "metrics": metrics, "cases": rows,
            "safety_controls": control_rows,
            "input_hashes": run["input_hashes"], "namespace_map": namespace_map,
            "authority_source": run.get("authority_source"),
            "gold_assessment": {"verified_declaration_cases": sum(row["gold_declaration_verified"] for row in rows),
                                "unverified_cases": sum(not row["gold_declaration_verified"] for row in rows),
                                "answer_label_verified_cases": sum(row["answer_label_verified"] for row in rows),
                                "independent_evidence_cases": sum(row["evidence_annotations_verified"] for row in rows),
                                "quality_metrics_assessable": quality_assessable,
                                "verification_basis": "explicit gold_verified=true, or frozen versioned synthetic domain-rule declaration; per-metric annotation completeness is separate"},
            "models_kind": "offline_artifact_evaluation",
            "quality_claim": "limited to independently declared, assessable gold annotations; unverified cases remain unknown" if quality_assessable else
                             "No independently verified gold quality claim; measured run telemetry and separately verified input protection only"}


def old37_replay(run_dir: Path) -> dict[str, Any]:
    run = load_run(run_dir)
    summary_path = run_dir / "summary.json"
    summary = json.loads(summary_path.read_text()) if summary_path.is_file() else {}
    validation = {"valid": None, "reason": "run_manifest.json absent; no artifact-validator claim"}
    if (run_dir / "run_manifest.json").is_file():
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
        from nested_doc_rag.artifacts import ArtifactValidationError, validate_step15_artifacts
        try:
            result = validate_step15_artifacts(run_dir)
            validation = {"valid": result["valid"], "evidence_validation": result["evidence_validation"]}
        except (ArtifactValidationError, ValueError, OSError) as exc:
            validation = {"valid": False, "reason": str(exc)}
    return {"schema_version": "old37-artifact-replay-v1", "runtime_kind": "historical_artifact_replay", "models_kind": "no_model_calls",
            "runner_kind": "offline_reader", "data_origin": "original_readonly_artifacts", "storage_initial_state": "historical",
            "counts": {"predictions": len(run["predictions"]), "evaluations": len(run["evaluations"]), "audit": len(run["audit"]),
                       "status": dict(Counter(row.get("answer_status") for row in run["predictions"])),
                       "action": dict(Counter(row.get("action") for row in run["audit"])),
                       "confirmed": sum(row.get("status") == "confirmed" for row in run["audit"])},
            "compatible_old141_read": len(run["predictions"]) == 141 and len(run["evaluations"]) == 141 and len(run["audit"]) == 141,
            "legacy_artifact_validation": validation,
            "historical_writeback37_observed": sum(row.get("action") == "written" for row in run["audit"]) == 37,
            "recorded_summary_written": (summary.get("writeback_summary") or {}).get("written_count"),
            "input_hashes": {**run["input_hashes"], **({"summary.json": digest(summary_path)} if summary_path.is_file() else {})},
            "quality_metrics": None, "quality_reason": "historical outputs/judge are compatibility evidence, not independent gold or a current live rerun"}


def write_report(report: dict[str, Any], out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if "cases" in report:
        (out_dir / "cases.jsonl").write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in report["cases"]), encoding="utf-8")
        lines = [f"# Phase 5 offline evaluation: {report['method']}", "", f"Cases: {report['case_count']}; failed: {report['failed_case_count']}; K={report['k']}.", "",
                 report["quality_claim"], "",
                 "| Metric | Value | Numerator | Denominator | Unknown |", "| --- | --- | --- | --- | --- |"]
        for name, value in report["metrics"].items():
            lines.append(f"| {name} | {value.get('value')} | {value.get('numerator', value.get('observed_count'))} | {value.get('denominator')} | {value.get('unknown_count')} |")
        lines.extend(["", "Exact populations and missing-data reasons are recorded in report.json. Failed/missing predictions remain in the accuracy denominator."])
        (out_dir / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--gold", type=Path)
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, required=True)
    parser.add_argument("--method", choices=["A0", "A1", "A2", "A3", "A4"], default="A3")
    parser.add_argument("--models-kind", choices=["real", "stub", "replay", "unknown"], default="unknown", help="Origin of the run's model outputs; offline scoring never implies real models.")
    parser.add_argument("--k", type=int, default=5)
    parser.add_argument("--namespace-map", type=Path)
    parser.add_argument("--dataset-id", help="Select one pair from a shared multi-dataset gold JSONL.")
    parser.add_argument("--dataset-manifest", type=Path, help="Include the selected pair's independently declared formula/protection controls.")
    parser.add_argument("--primary-predictions", type=Path)
    parser.add_argument("--retrieval-authority", type=Path, help="Actual observed retrieval sidecar for historical stages lacking a native authority artifact.")
    parser.add_argument("--replay-old37", action="store_true")
    args = parser.parse_args()
    if args.out_dir.resolve().is_relative_to(args.run_dir.resolve()):
        parser.error("report output must be outside the immutable run artifact directory")
    if args.replay_old37:
        report = old37_replay(args.run_dir)
    else:
        if args.gold is None:
            parser.error("--gold is required for quality evaluation")
        namespace_map = json.loads(args.namespace_map.read_text()) if args.namespace_map else {}
        golds = read_jsonl(args.gold)
        if args.dataset_id:
            golds = [gold for gold in golds if gold.get("dataset_id") == args.dataset_id]
        elif len({gold.get("dataset_id") for gold in golds}) > 1:
            parser.error("--dataset-id is required for a shared multi-dataset gold file")
        controls = []
        if args.dataset_manifest:
            manifest = json.loads(args.dataset_manifest.read_text())
            pairs = [pair for pair in manifest["pairs"] if pair["id"] == args.dataset_id]
            if len(pairs) != 1:
                parser.error("--dataset-manifest requires one valid --dataset-id")
            controls = pairs[0]["template"].get("safety_controls", [])
        if args.retrieval_authority is not None and not args.retrieval_authority.is_file():
            parser.error("--retrieval-authority must identify an existing observed sidecar")
        report = evaluate(golds, load_run(args.run_dir, retrieval_authority=args.retrieval_authority), method=args.method, k=args.k, namespace_map=namespace_map,
                          primary_predictions=read_jsonl(args.primary_predictions) if args.primary_predictions else None, safety_controls=controls)
        report["input_hashes"]["gold.jsonl"] = digest(args.gold)
        if args.primary_predictions:
            report["input_hashes"]["primary_predictions.jsonl"] = digest(args.primary_predictions)
        if args.dataset_manifest:
            report["input_hashes"]["dataset_manifest.json"] = digest(args.dataset_manifest)
        if args.namespace_map:
            report["input_hashes"]["namespace_map.json"] = digest(args.namespace_map)
        report["evaluation_kind"] = "offline_artifact_evaluation"
        report["models_kind"] = args.models_kind
        if args.models_kind != "real":
            report["quality_claim"] = "No real-model accuracy claim; " + report["quality_claim"]
    write_report(report, args.out_dir)
    print(json.dumps({"report": str(args.out_dir / "report.json"), "model_requests": 0, "schema_version": report["schema_version"]}))


if __name__ == "__main__":
    main()
