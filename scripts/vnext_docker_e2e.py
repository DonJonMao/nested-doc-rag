#!/usr/bin/env python3
"""Durable fresh-input acceptance driver; never builds, starts, or removes Docker.

prepare -> verify-clean -> externally run the printed Compose up command -> run.
Real and explicitly requested stub runs have separate identities. Neither mode
silently falls back. A successful stub run proves the engineering contract only.
All inputs, private runtime env, IDs, API responses, SSE and downloads are retained.
This driver covers A3 fresh Docker E2E; old141 and A0-A2 need separate ledgers.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import mimetypes
import os
import re
import secrets
import shlex
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import UTC, datetime
from pathlib import Path, PurePosixPath
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
COMPOSE = ROOT / "go-server/deployments/docker-compose.vnext-acceptance.yaml"
sys.path.insert(0, str(ROOT / "src"))
TERMINAL = {"succeeded", "completed_with_failures", "failed", "canceled", "cancelled"}
SECRET_KEY = re.compile(r"password|secret|(^|_)token($|_)|api_key$|access_key$|database_dsn", re.I)
REQUIRED_TYPES = {
    "run_manifest", "predictions_raw", "predictions", "agent_overlays",
    "predictions_agent_view", "review_items", "trace", "trace_summary",
    "run_summary", "summary", "filled_form", "writeback_audit", "evidence_map",
    "evidence_provenance", "retrieval_evidence", "form_input_snapshot", "index_scopes",
}


class AcceptanceError(RuntimeError):
    pass


def now() -> str:
    return datetime.now(UTC).isoformat()


def sha256(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    with temporary.open("w", encoding="utf-8") as stream:
        os.chmod(temporary, 0o600)
        json.dump(value, stream, ensure_ascii=False, indent=2)
        stream.write("\n")
    temporary.replace(path)


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def read_env(path: Path) -> dict[str, str]:
    result: dict[str, str] = {}
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        if line.startswith("export "):
            line = line[7:]
        if "=" not in line:
            raise AcceptanceError(f"invalid env line in {path.name}")
        key, value = line.split("=", 1)
        key, value = key.strip(), value.strip()
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", key):
            raise AcceptanceError(f"invalid env key in {path.name}")
        if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'":
            quote = value[0]
            value = value[1:-1]
            if quote == "'":
                value = value.replace("\\'", "'")
            else:
                value = value.replace('\\"', '"').replace("\\\\", "\\")
        result[key] = value
    return result


def write_env(path: Path, values: dict[str, str]) -> None:
    # Compose single quotes preserve literal $ in credentials. Never print this.
    with path.open("x", encoding="utf-8") as stream:
        os.chmod(path, 0o600)
        for key, value in sorted(values.items()):
            if "\n" in value or "\r" in value:
                raise AcceptanceError(f"multiline env value is unsupported: {key}")
            escaped = value.replace("'", "\\'")
            stream.write(f"{key}='{escaped}'\n")


def source_path(parent: Path, relative: str) -> Path:
    path = (parent / relative).resolve()
    if not path.is_relative_to(ROOT) or not path.is_file():
        raise AcceptanceError(f"dataset input must be a regular file inside vNext: {relative}")
    return path


def checked_source(parent: Path, record: dict[str, Any]) -> dict[str, Any]:
    path = source_path(parent, str(record["path"]))
    actual = sha256(path)
    expected = str(record.get("sha256", "")).removeprefix("sha256:")
    if not expected or actual != expected:
        raise AcceptanceError(f"input hash mismatch: {record['path']}")
    return {**record, "absolute_path": str(path), "sha256": actual, "file_size": path.stat().st_size}


def load_dataset(path: Path, pair_ids: list[str] | None) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    manifest = read_json(path)
    pairs = manifest.get("pairs", [])
    if not isinstance(pairs, list) or not pairs:
        raise AcceptanceError("manifest requires a nonempty pairs list")
    found: set[str] = set()
    selected = []
    from nested_doc_rag.form.template_parser import parse_form_template

    for original in pairs:
        pair = dict(original)
        pair_id = str(pair["id"])
        if not re.fullmatch(r"[A-Za-z0-9_-]+", pair_id) or pair_id in found:
            raise AcceptanceError("pair IDs must be unique safe directory names")
        found.add(pair_id)
        if pair_ids and pair_id not in pair_ids:
            continue
        namespaces = pair.get("namespaces") or {}
        if not namespaces.get("target") or not namespaces.get("global") or namespaces["target"] == namespaces["global"] or namespaces["target"] == "global":
            raise AcceptanceError(f"invalid target/global namespaces for {pair_id}")
        if not pair.get("collection"):
            raise AcceptanceError(f"collection is required for {pair_id}")
        for role in ("target", "global"):
            sources = pair.get(f"{role}_sources")
            if not isinstance(sources, list) or not sources:
                raise AcceptanceError(f"{pair_id} requires fresh {role} source files")
            pair[f"{role}_sources"] = [checked_source(path.parent, item) for item in sources]
            for item in pair[f"{role}_sources"]:
                if item.get("document_role", "knowledge_base") not in {"knowledge_base", "intro_doc", "proof_attachment", "misc"}:
                    raise AcceptanceError(f"invalid document role for {pair_id}")
                if item.get("file_name", Path(item["absolute_path"]).name) != Path(item["absolute_path"]).name:
                    raise AcceptanceError("source filename must match the uploaded file basename")
        pair["template"] = checked_source(path.parent, pair["template"])
        parsed = parse_form_template(Path(pair["template"]["absolute_path"]))
        expected = pair["template"].get("expected_fields")
        if isinstance(expected, list):
            expected = len(expected)
        if not isinstance(expected, int) or isinstance(expected, bool) or expected <= 0 or expected != len(parsed.items):
            raise AcceptanceError(f"{pair_id} template.expected_fields must match parsed count {len(parsed.items)}")
        pair["template"]["expected_fields"] = expected
        pair["template"]["parse_report"] = parsed.report
        pair["template"]["parsed_targets"] = [
            {key: item.get(key) for key in ("form_item_id", "sheet_name", "row_index", "target_cell")}
            for item in parsed.items
        ]
        declared_targets = pair["template"].get("expected_targets")
        if declared_targets is not None:
            actual_targets = {(item["sheet_name"], item["target_cell"].rpartition("!")[2]) for item in parsed.items}
            if actual_targets != {(item["sheet_name"], item["cell"]) for item in declared_targets}:
                raise AcceptanceError(f"{pair_id} parsed physical targets differ from the dataset")
        gold = dict(pair.get("gold") or {})
        gold_path = source_path(path.parent, str(gold["path"]))
        if gold.get("sha256") and sha256(gold_path) != str(gold["sha256"]).removeprefix("sha256:"):
            raise AcceptanceError(f"{pair_id} gold hash mismatch")
        gold.update(absolute_path=str(gold_path), sha256=sha256(gold_path))
        rows = jsonl(gold_path)
        dataset_id = gold.get("dataset_id", pair_id)
        matching = [row for row in rows if row.get("dataset_id", dataset_id) == dataset_id]
        if not matching:
            raise AcceptanceError(f"{pair_id} has no gold records")
        if {item["form_item_id"] for item in parsed.items} != {item.get("field_id") for item in matching}:
            raise AcceptanceError(f"{pair_id} parsed field identities differ from gold")
        gold["gold_count"] = len(matching)
        pair["gold"] = gold
        selected.append(pair)
    if pair_ids and set(pair_ids) - found:
        raise AcceptanceError(f"unknown pair IDs: {', '.join(sorted(set(pair_ids) - found))}")
    if not selected:
        raise AcceptanceError("no dataset pairs selected")
    return manifest, selected


def compose_args(state: dict[str, Any]) -> list[str]:
    return ["docker", "compose", "--env-file", state["runtime_env"], "-p", state["compose_project"], "-f", state["compose_file"]]


def prepare(args: argparse.Namespace) -> int:
    from nested_doc_rag.config import load_app_config

    manifest_path = args.manifest.resolve()
    raw_manifest, pairs = load_dataset(manifest_path, args.pair)
    model_config = args.model_config.resolve()
    if not model_config.is_file():
        raise AcceptanceError("model config does not exist")
    # Only explicitly supplied model variables and the host model env are used.
    # load_app_config(env=...) does NOT discover/load the workspace .env.
    model_env = dict(os.environ)
    explicit_env = read_env(args.models_env.resolve()) if args.models_env else {}
    model_env.update(explicit_env)
    cfg = load_app_config(model_config, env=model_env, project_root=ROOT)
    services = cfg.services
    if not cfg.retrieval.sufficiency_enabled or cfg.retrieval.schema_first_enabled or cfg.writeback.existing_value_policy != "preserve" or cfg.agentscope.mode != "off":
        raise AcceptanceError("fresh A3 acceptance requires sufficiency enabled, schema-first disabled, preserve writeback and agentscope off")
    if args.models_kind == "stub" and (not args.models_env or not args.model_note):
        raise AcceptanceError("stub mode requires --models-env and --model-note; it never proves real model quality")
    endpoints = {"chat": services.chat_endpoint, "embedding": services.embedding_endpoint, "rerank": services.rerank_endpoint}
    for kind, endpoint in endpoints.items():
        parsed_url = urllib.parse.urlsplit(endpoint)
        if parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc or parsed_url.username or parsed_url.password or parsed_url.query:
            raise AcceptanceError(f"{kind} endpoint must be a credential-free HTTP URL")
    suffix = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ") + "-" + secrets.token_hex(3)
    project = args.project or ("vnext-" + args.models_kind + "-" + suffix.lower())
    if not re.fullmatch(r"[a-z0-9][a-z0-9_-]+", project):
        raise AcceptanceError("Compose project must use lowercase letters, digits, - or _")
    state_dir = (args.state_dir or (ROOT / "artifacts/vnext/phase5/docker" / project)).resolve()
    if not state_dir.is_relative_to(ROOT / "artifacts"):
        raise AcceptanceError("private acceptance evidence must be under ignored vNext artifacts/")
    if state_dir.exists() and any(state_dir.iterdir()):
        raise AcceptanceError("state directory is not empty; create a new acceptance identity")
    state_dir.mkdir(parents=True, exist_ok=True)
    os.chmod(state_dir, 0o700)
    runtime_env = state_dir / "runtime.env"
    pg_password, minio_password = secrets.token_urlsafe(24), secrets.token_urlsafe(24)
    values = {
        "VNEXT_RUNTIME_ENV": str(runtime_env), "VNEXT_MODEL_CONFIG": str(model_config),
        "VNEXT_API_PORT": str(args.api_port), "VNEXT_API_IMAGE": args.api_image,
        "VNEXT_WORKER_IMAGE": args.worker_image,
        "POSTGRES_PASSWORD": pg_password, "MINIO_ROOT_USER": "vnext" + secrets.token_hex(8),
        "MINIO_ROOT_PASSWORD": minio_password,
        "GONGKAN_DATABASE_DSN": f"postgres://gongkan:{pg_password}@postgres:5432/gongkan?sslmode=disable",
        "GONGKAN_REDIS_ADDR": "redis:6379", "GONGKAN_STORAGE_TYPE": "minio",
        "GONGKAN_MINIO_ENDPOINT": "minio:9000", "GONGKAN_MINIO_BUCKET": "vnext-acceptance",
        "GONGKAN_MINIO_ACCESS_KEY": "", "GONGKAN_MINIO_SECRET_KEY": minio_password,
        "GONGKAN_MINIO_USE_SSL": "false", "GONGKAN_MODEL_GATEWAY_ENABLED": "false",
        "GONGKAN_JWT_SECRET": secrets.token_urlsafe(48),
        "GONGKAN_BOOTSTRAP_ADMIN_PASSWORD": secrets.token_urlsafe(24),
        "GONGKAN_PYTHON_CONFIG_PATH": "config/vnext_acceptance.yaml",
        "GONGKAN_PYTHON_ARTIFACT_VALIDATION_ENABLED": "true",
        "GONGKAN_JOBS_MAX_ATTEMPTS": "1", "GONGKAN_JOBS_REDIS_NAMESPACE": project,
        "GONGKAN_JOBS_EVENT_CHANNEL": project + ":run_events",
        "NDR_CHAT_ENDPOINT": services.chat_endpoint, "NDR_CHAT_MODEL": services.chat_model,
        "NDR_CHAT_API_KEY_ENV": services.chat_api_key_env,
        "NDR_EMBEDDING_ENDPOINT": services.embedding_endpoint, "NDR_EMBEDDING_MODEL": services.embedding_model,
        "NDR_RERANK_ENDPOINT": services.rerank_endpoint, "NDR_RERANK_MODEL": services.rerank_model,
        "NDR_QDRANT_URL": "http://qdrant:6333", "NDR_QDRANT_API_KEY_ENV": "QDRANT_API_KEY",
        "QDRANT_API_KEY": "", "VNEXT_MODELS_KIND": args.models_kind,
    }
    values["GONGKAN_MINIO_ACCESS_KEY"] = values["MINIO_ROOT_USER"]
    # Freeze other effective Python overrides rather than inheriting arbitrary
    # host environment into the container. Never carry host storage namespaces.
    excluded = ("QDRANT", "TARGET_NAMESPACE", "GLOBAL_NAMESPACE", "PATHS__", "PATH", "COLLECTION")
    for key, value in model_env.items():
        if key.startswith(("NESTED_DOC_RAG__", "NDR_")) and not any(part in key for part in excluded):
            values[key] = value
    api_key_name = services.chat_api_key_env
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", api_key_name):
        raise AcceptanceError("invalid chat API key env name")
    values[api_key_name] = model_env.get(api_key_name, "")
    # The isolated storage override always wins over copied model env.
    values.update({"NDR_QDRANT_URL": "http://qdrant:6333", "QDRANT_API_KEY": ""})
    write_env(runtime_env, values)
    commit = subprocess.run(["git", "rev-parse", "HEAD"], cwd=ROOT, capture_output=True, text=True, check=True).stdout.strip()
    dirty = subprocess.run(["git", "status", "--short"], cwd=ROOT, capture_output=True, text=True, check=True).stdout.splitlines()
    state = {
        "schema_version": "vnext-docker-e2e-v1", "created_at": now(), "status": "prepared",
        "state_dir": str(state_dir), "root": str(ROOT), "git_commit": commit, "git_status": dirty,
        "driver_sha256": sha256(Path(__file__)), "compose_file": str(COMPOSE), "compose_sha256": sha256(COMPOSE),
        "compose_project": project, "runtime_env": str(runtime_env),
        "base_url": f"http://127.0.0.1:{args.api_port}", "manifest_path": str(manifest_path),
        "manifest_sha256": sha256(manifest_path), "pairs": pairs,
        "identity": {
            "runtime_kind": "docker", "models_kind": args.models_kind, "runner_kind": "production-api-worker-python",
            "data_kind": "fresh-upload", "storage_kind": "isolated-new-compose-volumes", "method": "A3",
            "api_image": args.api_image, "worker_image": args.worker_image,
            "model_config": str(model_config), "model_config_sha256": sha256(model_config),
            "model_note": args.model_note or "existing model configuration; upstream fingerprints unknown",
            "effective_services": {**endpoints, "chat_model": services.chat_model, "embedding_model": services.embedding_model, "rerank_model": services.rerank_model},
            "effective_contract": {"sufficiency_enabled": True, "schema_first_enabled": False, "existing_value_policy": "preserve", "agentscope_mode": "off"},
            "credentials_available": {api_key_name: bool(values[api_key_name])},
            "usage": "unknown", "cost": "unknown", "upstream_weight_fingerprints": "unknown",
        },
        "commands": [], "http": [], "pair_results": [], "failures": [],
        "acceptance_scope": "fresh A3 engineering E2E; old141 and A0-A2 are separate required runs",
        "final_quality_gates_claimed": False,
    }
    write_json(state_dir / "dataset_manifest.json", raw_manifest)
    write_json(state_dir / "ledger.json", state)
    command = compose_args(state)
    print(json.dumps({
        "state_dir": str(state_dir), "models_kind": args.models_kind, "pair_ids": [p["id"] for p in pairs],
        "private_env": str(runtime_env),
        "verify_clean": shlex.join([sys.executable, str(Path(__file__)), "verify-clean", "--state-dir", str(state_dir)]),
        "compose_build": shlex.join(command + ["build"]),
        "compose_up": shlex.join(command + ["up", "-d", "--no-build"]),
        "run": shlex.join([sys.executable, str(Path(__file__)), "run", "--state-dir", str(state_dir)]),
    }, ensure_ascii=False, indent=2))
    return 0


class Ledger:
    def __init__(self, state_dir: Path):
        self.path = state_dir.resolve() / "ledger.json"
        self.state = read_json(self.path)
        self.directory = self.path.parent
        self.env = read_env(Path(self.state["runtime_env"]))
        self.secret_values = [value for key, value in self.env.items() if SECRET_KEY.search(key) and len(value) >= 6]
        for key in (self.env.get("NDR_CHAT_API_KEY_ENV", ""), "MINIO_ROOT_USER"):
            if len(self.env.get(key, "")) >= 6:
                self.secret_values.append(self.env[key])
        self.token = ""
        self.auth_at = 0.0

    def redact(self, value: Any) -> Any:
        if isinstance(value, dict):
            return {key: "[REDACTED]" if SECRET_KEY.search(key) else self.redact(item) for key, item in value.items()}
        if isinstance(value, list):
            return [self.redact(item) for item in value]
        if isinstance(value, str):
            for secret in self.secret_values:
                value = value.replace(secret, "[REDACTED]")
            if self.token:
                value = value.replace(self.token, "[REDACTED]")
        return value

    def save(self) -> None:
        write_json(self.path, self.redact(self.state))

    def evidence(self, name: str, value: Any) -> None:
        write_json(self.directory / name, self.redact(value))

    def command(self, argv: list[str], name: str, *, check: bool = True, timeout: int = 120) -> subprocess.CompletedProcess[str]:
        started = time.monotonic()
        env = dict(os.environ)
        env["PYTHONPATH"] = str(ROOT / "src") + (os.pathsep + env["PYTHONPATH"] if env.get("PYTHONPATH") else "")
        try:
            result = subprocess.run(argv, cwd=ROOT, env=env, capture_output=True, text=True, timeout=timeout)
        except subprocess.TimeoutExpired as error:
            self.state["commands"].append({"name": name, "argv": argv, "exit_code": None, "error": "TimeoutExpired", "elapsed_seconds": round(time.monotonic() - started, 3)})
            self.save()
            raise AcceptanceError(f"{name} timed out") from error
        record = {"name": name, "argv": argv, "exit_code": result.returncode, "elapsed_seconds": round(time.monotonic() - started, 3), "stdout": result.stdout, "stderr": result.stderr}
        self.evidence(f"commands/{name}.json", record)
        self.state["commands"].append({key: value for key, value in record.items() if key not in {"stdout", "stderr"}})
        self.save()
        if check and result.returncode:
            raise AcceptanceError(f"{name} exited {result.returncode}; see commands/{name}.json")
        return result

    def request(self, method: str, route: str, *, payload: dict[str, Any] | None = None, file: Path | None = None, fields: dict[str, str] | None = None, name: str, raw: bool = False, timeout: int = 60) -> Any:
        if self.token and route != "/auth/login" and time.monotonic() - self.auth_at >= 900:
            refreshed = self.request("POST", "/auth/login", payload={"username": "admin", "password": self.env["GONGKAN_BOOTSTRAP_ADMIN_PASSWORD"]}, name=name + "-reauth")
            self.token = refreshed["access_token"]
        headers = {}
        if self.token:
            headers["Authorization"] = "Bearer " + self.token
        body = None
        if file:
            boundary = "vnext-" + secrets.token_hex(20)
            pieces = []
            for key, value in (fields or {}).items():
                pieces.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{key}"\r\n\r\n{value}\r\n'.encode())
            filename = file.name.replace('"', "_").replace("\r", "_").replace("\n", "_")
            mime = mimetypes.guess_type(file.name)[0] or "application/octet-stream"
            pieces.append(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{filename}"\r\nContent-Type: {mime}\r\n\r\n'.encode())
            pieces.extend([file.read_bytes(), f"\r\n--{boundary}--\r\n".encode()])
            body = b"".join(pieces)
            headers["Content-Type"] = "multipart/form-data; boundary=" + boundary
        elif payload is not None:
            body = json.dumps(payload, ensure_ascii=False).encode()
            headers["Content-Type"] = "application/json"
        url = self.state["base_url"] + "/api/v1" + route
        started = time.monotonic()
        record: dict[str, Any] = {"name": name, "method": method, "route": route, "at": now()}
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        request = urllib.request.Request(url, data=body, headers=headers, method=method)
        try:
            with opener.open(request, timeout=timeout) as response:
                data = response.read()
                record.update(http_status=response.status, response_sha256=hashlib.sha256(data).hexdigest(), response_bytes=len(data))
            parsed = data if raw else json.loads(data)
            if not raw:
                self.evidence(f"api/{name}.json", parsed)
        except urllib.error.HTTPError as error:
            record.update(http_status=error.code, error_type="HTTPError")
            self.evidence(f"api/{name}.error.json", {"status": error.code, "body": error.read().decode(errors="replace")})
            raise AcceptanceError(f"{method} {route} returned HTTP {error.code}") from error
        except (OSError, ValueError) as error:
            record.update(error_type=type(error).__name__)
            raise AcceptanceError(f"{method} {route} failed: {type(error).__name__}") from error
        finally:
            record["elapsed_seconds"] = round(time.monotonic() - started, 3)
            self.state["http"].append(record)
            self.save()
        if raw:
            return parsed
        if not isinstance(parsed, dict) or "data" not in parsed:
            raise AcceptanceError(f"{name} lacks API data envelope")
        if route == "/auth/login":
            self.auth_at = time.monotonic()
        return parsed["data"]


def verify_clean(args: argparse.Namespace) -> int:
    ledger = Ledger(args.state_dir)
    if ledger.state["status"] != "prepared":
        raise AcceptanceError("verify-clean must precede Compose up and API work")
    project = ledger.state["compose_project"]
    volumes = ledger.command(["docker", "volume", "ls", "--filter", f"label=com.docker.compose.project={project}", "--format", "{{.Name}}"], "before-volume-list").stdout.splitlines()
    # Include manually created volumes with the same names as well as labels.
    all_names = ledger.command(["docker", "volume", "ls", "--format", "{{.Name}}"], "before-all-volume-names").stdout.splitlines()
    expected = [f"{project}_{name}" for name in ("postgres_data", "redis_data", "minio_data", "qdrant_data", "api_runtime", "worker_runtime", "worker_python_artifacts", "worker_python_tmp")]
    collisions = sorted(set(volumes) | (set(expected) & set(all_names)))
    containers = ledger.command(["docker", "ps", "-a", "--filter", f"label=com.docker.compose.project={project}", "--format", "{{.ID}}"], "before-container-list").stdout.splitlines()
    report = {"at": now(), "compose_project": project, "expected_volumes": expected, "existing_volumes": collisions, "existing_containers": containers, "passed": not collisions and not containers}
    ledger.evidence("storage-before.json", report)
    ledger.state["storage_before"] = report
    ledger.save()
    if not report["passed"]:
        raise AcceptanceError("acceptance project already has storage or containers; choose a fresh project")
    print(json.dumps({"passed": True, "compose_project": project, "evidence": str(ledger.directory / "storage-before.json")}))
    return 0


def capture_docker(ledger: Ledger, stage: str) -> None:
    command = compose_args(ledger.state)
    ledger.command(command + ["ps", "-a", "--format", "json"], stage + "-compose-ps", check=False)
    ids = ledger.command(command + ["ps", "-a", "-q"], stage + "-container-ids", check=False).stdout.splitlines()
    if ids:
        # Deliberately exclude Config.Env: docker inspect otherwise leaks keys.
        template = '{"id":{{json .Id}},"image_id":{{json .Image}},"image":{{json .Config.Image}},"created":{{json .Created}},"status":{{json .State.Status}},"mounts":{{json .Mounts}},"labels":{{json .Config.Labels}}}'
        result = ledger.command(["docker", "inspect", "--format", template, *ids], stage + "-container-inspect", check=False)
        inspected = [json.loads(line) for line in result.stdout.splitlines() if line.strip()] if result.returncode == 0 else []
        ledger.evidence(f"storage-{stage}.json", {"at": now(), "containers": inspected})
        if stage == "start":
            services = {item["labels"].get("com.docker.compose.service"): item for item in inspected}
            if set(services) != {"api", "worker", "postgres", "redis", "minio", "qdrant"}:
                raise AcceptanceError("acceptance project must contain exactly API, Worker, PostgreSQL, Redis, MinIO and Qdrant")
            for item in inspected:
                for mount in item["mounts"]:
                    if mount["Type"] == "bind" and not (item["labels"].get("com.docker.compose.service") == "worker" and mount["Destination"] == "/app/python-core/config/vnext_acceptance.yaml" and mount.get("RW") is False):
                        raise AcceptanceError("unexpected bind mount in acceptance project")
                    if mount["Type"] == "volume" and mount.get("Name") not in ledger.state["storage_before"]["expected_volumes"]:
                        raise AcceptanceError("unexpected storage volume in acceptance project")
            ledger.state["docker_identity"] = inspected
            ledger.save()
    elif stage == "start":
        raise AcceptanceError("no acceptance containers found; externally start the prepared Compose project first")
    project = ledger.state["compose_project"]
    names = ledger.command(["docker", "volume", "ls", "--filter", f"label=com.docker.compose.project={project}", "--format", "{{.Name}}"], stage + "-volume-list", check=False).stdout.splitlines()
    if names:
        ledger.command(["docker", "volume", "inspect", *names], stage + "-volume-inspect", check=False)


def wait_success(ledger: Ledger, route: str, name: str, timeout: int, interval: float) -> dict[str, Any]:
    deadline = time.monotonic() + timeout
    poll = 0
    while True:
        poll += 1
        result = ledger.request("GET", route, name=f"{name}-poll-{poll:04d}")
        # Fill detail exposes a normalized UI status ("completed") alongside
        # the domain status ("succeeded"). Jobs/ingestion use status directly.
        status = result.get("raw_status") or result.get("status")
        if status in TERMINAL:
            if status != "succeeded":
                raise AcceptanceError(f"{name} terminal status is {status}")
            return result
        if time.monotonic() >= deadline:
            raise AcceptanceError(f"{name} timed out with status {status}")
        time.sleep(min(interval, max(0, deadline - time.monotonic())))


def capture_sse(ledger: Ledger, route: str, name: str, *, required: set[str] | None = None, duration: int = 6, last_event_id: str | None = None) -> dict[str, Any]:
    headers = {"Authorization": "Bearer " + ledger.token, "Accept": "text/event-stream"}
    if last_event_id:
        headers["Last-Event-ID"] = last_event_id
    request = urllib.request.Request(ledger.state["base_url"] + "/api/v1" + route, headers=headers)
    chunks: list[bytes] = []
    timeout_expected = False
    deadline = time.monotonic() + duration
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    try:
        with opener.open(request, timeout=duration) as response:
            while time.monotonic() < deadline:
                line = response.readline()
                if not line:
                    break
                chunks.append(line)
    except TimeoutError:
        timeout_expected = True
    finally:
        raw = ledger.redact(b"".join(chunks).decode(errors="replace"))
        path = ledger.directory / "sse" / (name + ".sse")
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(raw, encoding="utf-8")
        os.chmod(path, 0o600)
    events = []
    for block in raw.replace("\r\n", "\n").split("\n\n"):
        parsed: dict[str, Any] = {}
        data_lines = []
        for line in block.splitlines():
            if line.startswith("id:"):
                parsed["id"] = line[3:].strip()
            elif line.startswith("event:"):
                parsed["event"] = line[6:].strip()
            elif line.startswith("data:"):
                data_lines.append(line[5:].strip())
        if parsed.get("event"):
            if data_lines:
                parsed["data"] = json.loads("\n".join(data_lines))
            events.append(parsed)
    sequences = [event["data"]["sequence"] for event in events if isinstance(event.get("data"), dict) and "sequence" in event["data"]]
    event_types = {event["event"] for event in events}
    report = {"route": route, "events": events, "timeout_expected": timeout_expected, "sequences": sequences, "last_event_id": events[-1].get("id") if events else None, "sha256": sha256(path)}
    ledger.evidence("sse/" + name + ".json", report)
    if required and required - event_types:
        raise AcceptanceError(f"{name} SSE lacks {', '.join(sorted(required - event_types))}")
    if required and (not sequences or sequences != sorted(set(sequences)) or any(not event.get("id") for event in events)):
        raise AcceptanceError(f"{name} SSE sequence/id contract is invalid")
    return report


def safe_relative(value: str) -> PurePosixPath:
    if "\\" in value or "\x00" in value:
        raise AcceptanceError("unsafe archive relative path")
    path = PurePosixPath(value)
    if path.is_absolute() or not path.parts or any(part in {".", ".."} for part in path.parts) or ":" in path.parts[0]:
        raise AcceptanceError("unsafe archive relative path")
    return path


def download_artifacts(ledger: Ledger, run_id: str, workspace_id: str, destination: Path, name: str, *, fill: bool) -> dict[str, Any]:
    query = urllib.parse.urlencode({"workspace_id": workspace_id})
    records = ledger.request("GET", f"/runs/{run_id}/artifacts?{query}", name=name + "-artifact-list")
    if not isinstance(records, list) or not records:
        raise AcceptanceError(f"{name} has no archived artifacts")
    by_type: dict[str, dict[str, Any]] = {}
    for item in records:
        if item["artifact_type"] in by_type:
            raise AcceptanceError("duplicate archive artifact type")
        if item.get("workspace_id") != workspace_id or item.get("run_id") != run_id:
            raise AcceptanceError("archive artifact ownership mismatch")
        by_type[item["artifact_type"]] = item
    if "run_manifest" not in by_type:
        raise AcceptanceError("archive has no run manifest")
    destination.mkdir(parents=True, exist_ok=True)
    downloaded = []

    def download(item: dict[str, Any], relative: str) -> bytes:
        path = destination / safe_relative(relative)
        if not path.resolve().is_relative_to(destination.resolve()):
            raise AcceptanceError("download escapes archive directory")
        data = ledger.request("GET", f"/artifacts/{item['id']}/download", name=f"{name}-download-{item['artifact_type']}", raw=True)
        actual_hash = hashlib.sha256(data).hexdigest()
        if actual_hash != item.get("sha256") or len(data) != item.get("file_size"):
            raise AcceptanceError(f"download hash/size mismatch: {item['artifact_type']}")
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.exists() and path.read_bytes() != data:
            raise AcceptanceError("archive relative path collision")
        path.write_bytes(data)
        os.chmod(path, 0o600)
        downloaded.append({"artifact_id": item["id"], "artifact_type": item["artifact_type"], "relative_path": relative, "sha256": actual_hash, "file_size": len(data)})
        return data

    manifest = json.loads(download(by_type["run_manifest"], "run_manifest.json"))
    artifact_paths = manifest.get("artifacts") or {}
    if not isinstance(artifact_paths, dict):
        raise AcceptanceError("run manifest artifacts must be a map")
    for artifact_type, item in by_type.items():
        if artifact_type == "run_manifest":
            continue
        relative = artifact_paths.get(artifact_type)
        if not isinstance(relative, str) or not relative:
            raise AcceptanceError(f"archive artifact absent from manifest: {artifact_type}")
        download(item, relative)
    declared = {key for key, path in artifact_paths.items() if path}
    if declared - set(by_type):
        raise AcceptanceError(f"manifest artifacts missing from archive: {', '.join(sorted(declared - set(by_type)))}")
    if fill and REQUIRED_TYPES - set(by_type):
        raise AcceptanceError(f"fill archive lacks {', '.join(sorted(REQUIRED_TYPES - set(by_type)))}")
    report = {"run_id": run_id, "workspace_id": workspace_id, "run_dir": str(destination), "artifacts": downloaded, "manifest_schema_version": manifest.get("schema_version")}
    ledger.evidence(str(destination.relative_to(ledger.directory)) + "/download_receipt.json", report)
    return report


def inspect_fill(ledger: Ledger, pair: dict[str, Any], created: dict[str, Any], result: dict[str, Any], run_dir: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    snapshot = read_json(run_dir / "form_input_snapshot.json")
    expected_fields = pair["template"]["expected_fields"]
    if snapshot.get("selected_field_count") != expected_fields or result.get("progress_total") != expected_fields or result.get("progress_done") != expected_fields:
        raise AcceptanceError("actual parsed/processed field count differs from dataset")
    if snapshot.get("template_sha256") != pair["template"]["sha256"] or snapshot.get("input_mode") != "template":
        raise AcceptanceError("run did not use the exact uploaded fresh template")
    if snapshot.get("target_namespace") != pair["namespaces"]["target"] or snapshot.get("global_namespace") != pair["namespaces"]["global"]:
        raise AcceptanceError("frozen template namespaces differ from dataset")
    scopes = read_json(run_dir / "index_scopes.json")
    actual = {(scope["knowledge_base_id"], scope["index_version_id"], scope["namespace"], scope["collection"], scope["storage_contract"]) for scope in scopes}
    expected = {(created[role + "_scope"]["knowledge_base_id"], created[role + "_scope"]["index_version_id"], pair["namespaces"][role], pair["collection"], "versioned_v1") for role in ("target", "global")}
    if actual != expected:
        raise AcceptanceError("downloaded frozen index scopes differ from published KB versions")
    workbook = load_workbook(run_dir / "filled_form.xlsx", data_only=False)
    original = load_workbook(Path(pair["template"]["absolute_path"]), data_only=False)
    protected = []
    overwritten = []
    formulas = []
    try:
        for sheet in original.worksheets:
            if sheet.title not in workbook.sheetnames:
                raise AcceptanceError("filled workbook lost an original sheet")
            output = workbook[sheet.title]
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formulas.append({"sheet": sheet.title, "cell": cell.coordinate})
                    if cell.value is not None and cell.value != "":
                        protected.append((sheet.title, cell.coordinate))
                        if output[cell.coordinate].value != cell.value:
                            overwritten.append({"sheet": sheet.title, "cell": cell.coordinate})
    finally:
        workbook.close()
        original.close()
    if overwritten:
        ledger.evidence(f"pairs/{pair['id']}/unsafe-overwrites.json", overwritten)
        raise AcceptanceError("preserve policy changed nonempty/formula input cells")
    predictions = jsonl(run_dir / "predictions_raw.jsonl")
    if len(predictions) != expected_fields:
        raise AcceptanceError("raw prediction count differs from parsed template fields")
    return {"selected_field_count": expected_fields, "prediction_count": len(predictions), "protected_nonempty_cells": len(protected), "formula_cells": formulas, "unsafe_overwrite_count": len(overwritten), "scope_count": len(scopes), "template_sha256": snapshot["template_sha256"]}


def inspect_ingestion(info: dict[str, Any], run_dir: Path, workspace_id: str) -> dict[str, Any]:
    snapshot_path = run_dir / "input_snapshot.json"
    snapshot = read_json(snapshot_path)
    receipt = read_json(run_dir / "validation_receipt.json")
    expected = {"knowledge_base_id": info["knowledge_base_id"], "index_version_id": info["index_version_id"], "namespace": info["namespace"], "collection": info["collection"]}
    if snapshot.get("workspace_id") != workspace_id or any(snapshot.get(key) != value or receipt.get(key) != value for key, value in expected.items()):
        raise AcceptanceError("ingestion archive scope differs from created KB/version")
    original_sources = {(item["document_id"], item["file_id"], item["file_name"], item["sha256"]) for item in info["uploads"]}
    frozen_sources = {(item["document_id"], item["file_id"], item["filename"], item["sha256"]) for item in snapshot["documents"]}
    if original_sources != frozen_sources:
        raise AcceptanceError("ingestion frozen sources differ from uploaded input files")
    if receipt.get("input_snapshot_hash") != sha256(snapshot_path) or receipt.get("input_snapshot_hash") != info["publication"].get("input_snapshot_hash") or receipt.get("document_count") != len(original_sources) or not receipt.get("source_hashes_verified") or not receipt.get("smoke_passed") or receipt.get("expected_evidence_count", 0) <= 0 or receipt.get("expected_evidence_count") != receipt.get("actual_evidence_count") or receipt.get("expected_schema_count") != receipt.get("actual_schema_count"):
        raise AcceptanceError("ingestion receipt hash/count/smoke validation failed")
    return receipt


def run_pair(ledger: Ledger, pair: dict[str, Any], args: argparse.Namespace) -> None:
    pair_id = pair["id"]
    result: dict[str, Any] = {"pair_id": pair_id, "status": "running", "started_at": now(), "input_hashes": {"template": pair["template"]["sha256"], "gold": pair["gold"]["sha256"]}, "knowledge_bases": {}}
    ledger.state["pair_results"].append(result)
    ledger.save()
    workspace = ledger.request("POST", "/workspaces", payload={"name": f"vNext {pair_id} {ledger.state['compose_project']}", "description": "Fresh isolated acceptance inputs"}, name=pair_id + "-workspace-create")
    workspace_id = workspace["id"]
    result["workspace_id"] = workspace_id
    ledger.save()
    existing = ledger.request("GET", "/knowledge-bases?" + urllib.parse.urlencode({"workspace_id": workspace_id}), name=pair_id + "-initial-kbs")
    if existing.get("knowledge_bases"):
        raise AcceptanceError("new acceptance workspace already contains knowledge bases")
    for role in ("target", "global"):
        prefix = f"{pair_id}-{role}"
        kb = ledger.request("POST", "/knowledge-bases", payload={"workspace_id": workspace_id, "name": f"{pair_id} {role}", "namespace": pair["namespaces"][role], "qdrant_collection": pair["collection"]}, name=prefix + "-kb-create")
        kb_id = kb["id"]
        info = {"knowledge_base_id": kb_id, "namespace": pair["namespaces"][role], "collection": pair["collection"], "uploads": []}
        result["knowledge_bases"][role] = info
        ledger.save()
        for index, source in enumerate(pair[role + "_sources"], 1):
            uploaded = ledger.request("POST", f"/knowledge-bases/{kb_id}/documents", file=Path(source["absolute_path"]), fields={"document_role": source.get("document_role", "knowledge_base"), "namespace": pair["namespaces"][role]}, name=f"{prefix}-upload-{index:02d}")
            info["uploads"].append({"document_key": source.get("document_key"), "document_id": uploaded["id"], "file_id": uploaded["file_id"], "file_name": uploaded["filename"], "sha256": source["sha256"]})
            ledger.save()
        ingestion = ledger.request("POST", f"/knowledge-bases/{kb_id}/ingestion-runs", payload={}, name=prefix + "-ingest-create")
        info.update(ingestion_run_id=ingestion["id"], index_version_id=ingestion["index_version_id"], job_id=ingestion["job_id"])
        ledger.save()
        completed = wait_success(ledger, f"/ingestion-runs/{ingestion['id']}", prefix + "-ingest", args.wait_timeout, args.poll_interval)
        wait_success(ledger, f"/jobs/{ingestion['job_id']}", prefix + "-job", args.wait_timeout, args.poll_interval)
        info["completed_ingestion"] = completed
        info["sse"] = capture_sse(ledger, f"/ingestion-runs/{ingestion['id']}/events?after_sequence=0", prefix + "-ingest", required={"succeeded", "index_version_ready", "artifacts_registered"})
        kb_detail = ledger.request("GET", f"/knowledge-bases/{kb_id}", name=prefix + "-published-kb")
        versions = ledger.request("GET", f"/knowledge-bases/{kb_id}/index-versions", name=prefix + "-index-versions")["index_versions"]
        version = next((item for item in versions if item["id"] == ingestion["index_version_id"]), None)
        if kb_detail.get("status") != "ready" or kb_detail.get("current_index_version_id") != ingestion["index_version_id"] or not version or any(version.get(key) != value for key, value in {"storage_contract": "versioned_v1", "validation_state": "validated", "publication_state": "activated", "status": "ready"}.items()) or version.get("chunk_count", 0) <= 0 or version.get("document_count") != len(info["uploads"]):
            raise AcceptanceError(f"{prefix} index was not validated and atomically published")
        info["publication"] = version
        ingestion_dir = ledger.directory / "pairs" / pair_id / "ingestions" / role
        info["download"] = download_artifacts(ledger, ingestion["id"], workspace_id, ingestion_dir, prefix + "-ingest", fill=False)
        info["receipt_checks"] = inspect_ingestion(info, ingestion_dir, workspace_id)
        ledger.save()
    form = ledger.request("POST", "/forms", file=Path(pair["template"]["absolute_path"]), fields={"workspace_id": workspace_id}, name=pair_id + "-form-upload")
    result["form_file_id"] = form["id"]
    payload = {"workspace_id": workspace_id, "knowledge_base_id": result["knowledge_bases"]["target"]["knowledge_base_id"], "global_knowledge_base_id": result["knowledge_bases"]["global"]["knowledge_base_id"], "form_file_id": form["id"], "room_context": pair.get("room_context", ""), "rows": "all", "retrieval_mode": "layered", "prompt_version": "step15_compat", "judge": False, "use_judge_cache": False, "writeback": True, "name": f"vNext {pair_id} fresh A3"}
    fill = ledger.request("POST", "/fill-runs", payload=payload, name=pair_id + "-fill-create")
    result.update(fill_run_id=fill["id"], fill_job_id=fill["job_id"], fill_request=payload)
    for role in ("target", "global"):
        scope = fill.get(role + "_scope") or {}
        info = result["knowledge_bases"][role]
        if scope.get("index_version_id") != info["index_version_id"] or scope.get("knowledge_base_id") != info["knowledge_base_id"]:
            raise AcceptanceError("fill did not freeze the exact published KB versions")
    ledger.save()
    completed = wait_success(ledger, f"/fill-runs/{fill['id']}", pair_id + "-fill", args.wait_timeout, args.poll_interval)
    job = wait_success(ledger, f"/jobs/{fill['job_id']}", pair_id + "-fill-job", args.wait_timeout, args.poll_interval)
    result["completed_fill"] = completed
    result["completed_job"] = job
    query = urllib.parse.urlencode({"workspace_id": workspace_id, "after_sequence": 0})
    result["sse"] = capture_sse(ledger, f"/runs/{fill['id']}/events?{query}", pair_id + "-fill", required={"succeeded", "artifact_validation_succeeded", "artifacts_registered", "python_started", "python_finished"})
    if result["sse"]["last_event_id"]:
        result["sse_replay"] = capture_sse(ledger, f"/runs/{fill['id']}/events?workspace_id={workspace_id}", pair_id + "-fill-last-event-replay", last_event_id=result["sse"]["last_event_id"], duration=2)
        if any(sequence <= max(result["sse"]["sequences"]) for sequence in result["sse_replay"]["sequences"]):
            raise AcceptanceError("Last-Event-ID replay re-emitted an earlier event")
    run_dir = ledger.directory / "pairs" / pair_id / "run"
    result["download"] = download_artifacts(ledger, fill["id"], workspace_id, run_dir, pair_id + "-fill", fill=True)
    workbook = ledger.request("GET", f"/fill-runs/{fill['id']}/downloads/filled-form", name=pair_id + "-filled-form-convenience", raw=True)
    if workbook != (run_dir / "filled_form.xlsx").read_bytes():
        raise AcceptanceError("filled-form convenience endpoint differs from archived workbook")
    result["workbook_convenience_sha256"] = hashlib.sha256(workbook).hexdigest()
    # The creation response owns immutable pins and the worker output path;
    # the GET detail response owns completion/progress and public UI data.
    out_dir = fill.get("out_dir")
    if not out_dir or not PurePosixPath(out_dir).is_relative_to(PurePosixPath("/app/python-core/artifacts")):
        raise AcceptanceError("worker run output directory is outside its isolated artifacts volume")
    ledger.command(compose_args(ledger.state) + ["exec", "-T", "worker", "python", "-m", "nested_doc_rag.cli", "validate-artifacts", "--run-dir", out_dir], pair_id + "-worker-validator")
    ledger.command([sys.executable, "-m", "nested_doc_rag.cli", "validate-artifacts", "--run-dir", str(run_dir)], pair_id + "-download-validator")
    result["input_and_writeback_checks"] = inspect_fill(ledger, pair, fill, completed, run_dir)
    if args.evaluate:
        evaluator = ROOT / "scripts/vnext_evaluate.py"
        if not evaluator.is_file():
            raise AcceptanceError("--evaluate requested but evaluator is absent")
        evaluation_dir = ledger.directory / "pairs" / pair_id / "evaluation"
        namespace_map_path = ledger.directory / "pairs" / pair_id / "namespace_map.json"
        ledger.evidence(str(namespace_map_path.relative_to(ledger.directory)), {pair["gold"].get("dataset_id", pair_id): pair["namespaces"]["target"], "global": pair["namespaces"]["global"]})
        ledger.command([sys.executable, str(evaluator), "--gold", pair["gold"]["absolute_path"], "--dataset-id", pair["gold"].get("dataset_id", pair_id), "--dataset-manifest", str(ledger.directory / "dataset_manifest.json"), "--namespace-map", str(namespace_map_path), "--run-dir", str(run_dir), "--out-dir", str(evaluation_dir), "--method", "A3", "--models-kind", ledger.state["identity"]["models_kind"], "--k", str(args.k)], pair_id + "-evaluation", timeout=300)
        result["evaluation_dir"] = str(evaluation_dir)
    result.update(status="succeeded", finished_at=now())
    ledger.save()
    print(json.dumps({"pair_id": pair_id, "status": "succeeded", "models_kind": ledger.state["identity"]["models_kind"], "run_id": fill["id"], "selected_fields": pair["template"]["expected_fields"], "run_dir": str(run_dir)}, ensure_ascii=False), flush=True)


def run(args: argparse.Namespace) -> int:
    ledger = Ledger(args.state_dir)
    if ledger.state["status"] != "prepared":
        raise AcceptanceError("each prepared acceptance identity may run only once; keep failures and prepare a new state")
    if not ledger.state.get("storage_before", {}).get("passed"):
        raise AcceptanceError("verify-clean evidence is required before Compose up")
    if sha256(COMPOSE) != ledger.state["compose_sha256"] or sha256(Path(__file__)) != ledger.state["driver_sha256"] or sha256(Path(ledger.state["identity"]["model_config"])) != ledger.state["identity"]["model_config_sha256"]:
        raise AcceptanceError("driver/Compose/model config changed since prepare; create a new identity")
    # Freeze content hashes again at execution, even when input files are ignored.
    for pair in ledger.state["pairs"]:
        for item in [pair["template"], pair["gold"], *pair["target_sources"], *pair["global_sources"]]:
            if sha256(Path(item["absolute_path"])) != item["sha256"]:
                raise AcceptanceError("dataset input changed since prepare")
    ledger.state.update(status="running", started_at=now())
    ledger.save()
    try:
        capture_docker(ledger, "start")
        login = ledger.request("POST", "/auth/login", payload={"username": "admin", "password": ledger.env["GONGKAN_BOOTSTRAP_ADMIN_PASSWORD"]}, name="login")
        ledger.token = login["access_token"]
        ledger.state["actor"] = ledger.redact({key: value for key, value in login.items() if "token" not in key.lower()})
        ledger.save()
        for pair in ledger.state["pairs"]:
            # Login again for each pair; long model runs may exceed token TTL.
            login = ledger.request("POST", "/auth/login", payload={"username": "admin", "password": ledger.env["GONGKAN_BOOTSTRAP_ADMIN_PASSWORD"]}, name=pair["id"] + "-login")
            ledger.token = login["access_token"]
            run_pair(ledger, pair, args)
        ledger.state.update(status="succeeded", finished_at=now(), completed_fresh_pairs=len(ledger.state["pair_results"]))
        ledger.state["engineering_contract_passed"] = True
        ledger.state["real_docker_chain_passed"] = ledger.state["identity"]["models_kind"] == "real"
    except Exception as error:
        ledger.state.update(status="failed", finished_at=now(), engineering_contract_passed=False, real_docker_chain_passed=False)
        ledger.state["failures"].append({"at": now(), "type": type(error).__name__, "message": ledger.redact(str(error))})
        if ledger.state["pair_results"]:
            current = ledger.state["pair_results"][-1]
            if current["status"] == "running":
                current.update(status="failed", finished_at=now())
            # Retain failure SSE without changing the original failure reason.
            if ledger.token:
                candidates = [(info.get("ingestion_run_id"), f"/ingestion-runs/{info.get('ingestion_run_id')}/events?after_sequence=0", role) for role, info in current["knowledge_bases"].items() if info.get("ingestion_run_id")]
                if current.get("fill_run_id"):
                    candidates.append((current["fill_run_id"], f"/runs/{current['fill_run_id']}/events?workspace_id={current['workspace_id']}&after_sequence=0", "fill"))
                for _, route, role in candidates:
                    try:
                        capture_sse(ledger, route, f"{current['pair_id']}-{role}-failure", duration=2)
                    except Exception as capture_error:
                        ledger.state["failures"].append({"type": "FailureCapture", "message": type(capture_error).__name__})
                for role, info in current["knowledge_bases"].items():
                    if info.get("ingestion_run_id") and not info.get("download"):
                        try:
                            info["failure_download"] = download_artifacts(ledger, info["ingestion_run_id"], current["workspace_id"], ledger.directory / "pairs" / current["pair_id"] / "failure" / role, current["pair_id"] + "-" + role + "-failure", fill=False)
                        except Exception as capture_error:
                            ledger.state["failures"].append({"type": "IngestionFailureArtifacts", "message": type(capture_error).__name__})
                if current.get("fill_run_id") and not current.get("download"):
                    try:
                        current["failure_download"] = download_artifacts(ledger, current["fill_run_id"], current["workspace_id"], ledger.directory / "pairs" / current["pair_id"] / "failure" / "fill", current["pair_id"] + "-fill-failure", fill=False)
                    except Exception as capture_error:
                        ledger.state["failures"].append({"type": "FillFailureArtifacts", "message": type(capture_error).__name__})
        print(json.dumps({"status": "failed", "models_kind": ledger.state["identity"]["models_kind"], "error": ledger.redact(str(error)), "ledger": str(ledger.path)}, ensure_ascii=False), file=sys.stderr)
    finally:
        ledger.save()
        try:
            capture_docker(ledger, "finish")
            ledger.command(compose_args(ledger.state) + ["logs", "--no-color", "--timestamps", "api", "worker"], "finish-api-worker-logs", check=False)
        except Exception as error:
            ledger.state["failures"].append({"type": "DockerEvidenceCapture", "message": type(error).__name__})
        ledger.save()
    print(json.dumps({"status": ledger.state["status"], "models_kind": ledger.state["identity"]["models_kind"], "completed_fresh_pairs": sum(item["status"] == "succeeded" for item in ledger.state["pair_results"]), "final_quality_gates_claimed": False, "ledger": str(ledger.path)}, ensure_ascii=False))
    return 0 if ledger.state["status"] == "succeeded" else 1


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    subparsers = parser.add_subparsers(dest="command", required=True)
    preparation = subparsers.add_parser("prepare", help="Validate local inputs and generate a private env; no Docker or model calls")
    preparation.add_argument("--manifest", type=Path, default=ROOT / "docs/vnext/datasets/manifest.json")
    preparation.add_argument("--model-config", type=Path, default=ROOT / "config/docker.yaml")
    preparation.add_argument("--models-env", type=Path)
    preparation.add_argument("--models-kind", choices=("real", "stub"), default="real")
    preparation.add_argument("--model-note")
    preparation.add_argument("--pair", action="append", help="Select pair ID; repeat to select multiple pairs")
    preparation.add_argument("--state-dir", type=Path)
    preparation.add_argument("--project")
    preparation.add_argument("--api-port", type=int, default=18080)
    preparation.add_argument("--api-image", default="datacenter-vnext-api:phase5")
    preparation.add_argument("--worker-image", default="datacenter-vnext-worker:phase5")
    preparation.set_defaults(func=prepare)
    clean = subparsers.add_parser("verify-clean", help="Read-only Docker evidence of absent project containers/volumes BEFORE up")
    clean.add_argument("--state-dir", type=Path, required=True)
    clean.set_defaults(func=verify_clean)
    execution = subparsers.add_parser("run", help="Fresh API upload/ingest/fill/download chain; never builds/starts/stops Docker")
    execution.add_argument("--state-dir", type=Path, required=True)
    execution.add_argument("--wait-timeout", type=int, default=3600)
    execution.add_argument("--poll-interval", type=float, default=3)
    execution.add_argument("--evaluate", action="store_true", help="Run vnext_evaluate.py on each downloaded archive")
    execution.add_argument("--k", type=int, default=5)
    execution.set_defaults(func=run)
    args = parser.parse_args()
    if args.command == "prepare" and not 1 <= args.api_port <= 65535:
        parser.error("api port must be 1..65535")
    if args.command == "run" and (args.wait_timeout <= 0 or args.poll_interval <= 0):
        parser.error("timeouts and intervals must be positive")
    try:
        return args.func(args)
    except (AcceptanceError, OSError, KeyError, ValueError) as error:
        # Never emit raw exception values here: malformed inputs may contain keys.
        print(json.dumps({"status": "failed", "error_type": type(error).__name__, "message": str(error) if isinstance(error, AcceptanceError) else "local input/command error; inspect the retained evidence"}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
