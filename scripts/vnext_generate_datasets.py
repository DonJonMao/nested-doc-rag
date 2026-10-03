#!/usr/bin/env python3
"""Generate anonymous native Phase 5 fixtures and independent physical gold.

Uses only dependencies already required by the production project. The recipes
below are the gold authority; native parsers validate them, never supply answers.
No embeddings, Qdrant, retrieval, model service or previous artifact is accessed.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import re
import sys
import zipfile
from collections import Counter
from datetime import datetime
from importlib.metadata import version
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT / "src"))

COLLECTION = "vnext_phase5_native"
FIXED_DATE = datetime(2026, 1, 1)
SCHEMA = "vnext-evaluation-dataset-v1"
INSTRUCTION = "仅填写所问机房与阶段的事实，保留数值和单位。"
CATEGORIES = {
    "same_value_wrong_field": 5,
    "same_field_wrong_room": 5,
    "current_vs_planned": 4,
    "numeric_substring": 4,
    "target_global_conflict": 4,
    "table_detail": 4,
    "second_round": 4,
}


def digest(data: bytes) -> str:
    return "sha256:" + hashlib.sha256(data).hexdigest()


def json_bytes(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode()


def fixed_zip(path: Path) -> None:
    """OOXML zip metadata must not make otherwise identical fixtures drift."""
    original = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            data = source.read(name)
            if name == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                    rb"\g<1>2026-01-01T00:00:00Z\g<2>",
                    data,
                )
            target.writestr(info, data)


def workbook(sheets: dict[str, list[list[Any]]], path: Path) -> None:
    book = Workbook()
    book.remove(book.active)
    book.properties.creator = "Anonymous synthetic evaluation"
    book.properties.created = book.properties.modified = FIXED_DATE
    edge = Side(style="thin", color="D9D9D9")
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                cell.font = Font(name="Arial", size=11, color="000000")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
        for cell in sheet[1]:
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="294766")
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 29 if col == 1 else 24
        for row in range(1, sheet.max_row + 1):
            sheet.row_dimensions[row].height = 30
        sheet.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    book.close()
    fixed_zip(path)


def word(paragraphs: list[str], tables: list[list[list[str]]], path: Path) -> None:
    doc = Document()
    doc.core_properties.author = "Anonymous synthetic evaluation"
    doc.core_properties.created = doc.core_properties.modified = FIXED_DATE
    doc.sections[0].top_margin = doc.sections[0].bottom_margin = Inches(0.7)
    style = doc.styles["Normal"]
    style.font.name, style.font.size, style.font.color.rgb = "Arial", Pt(11), RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(6)
    for index, text in enumerate(paragraphs):
        doc.add_paragraph(text, style="Title" if index == 0 else "Normal")
    for rows in tables:
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.style = "Table Grid"
        for values in rows:
            cells = table.add_row().cells
            for cell, text in zip(cells, values, strict=True):
                cell.text = text
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    fixed_zip(path)


def native_ref(document: dict[str, Any], values: list[str] | str, *, fact_id: str = "", room: str = "", **address: Any) -> dict[str, Any]:
    text = " / ".join(values) if isinstance(values, list) else values
    return {
        "fact_id": fact_id,
        "document_key": document["document_key"],
        "file_name": document["file_name"],
        "namespace": document["namespace"],
        "namespace_role": "global" if document["namespace"] == "global" else "target",
        "room_scope": room,
        **address,
        "source_text": text,
        "source_text_contains": False,
        "source_text_hash": digest(text.encode()),
        "source_text_hash_space": "native_full_record",
    }


def decoy(ref: dict[str, Any], reason: str) -> dict[str, Any]:
    return {**ref, "reason": reason}


def case(
    number: int,
    dataset: str,
    category: str,
    question: str,
    answer: str,
    refs: list[dict[str, Any]],
    decoys: list[dict[str, Any]] | None = None,
    *,
    rounds: int | None = None,
    missing: list[str] | None = None,
    answerable: bool = True,
    original: Any = None,
) -> dict[str, Any]:
    return {
        "schema_version": SCHEMA,
        "case_id": f"C{number:02d}",
        "dataset_id": dataset,
        "category": category,
        "question": question,
        "expected_answer": answer,
        "accepted_answers": [answer],
        "expected_status": "answered" if answerable else "partial_clue",
        "expected_answer_status": "answered" if answerable else "partial_clue",
        "expected_writeback_status": "confirmed" if answerable else "flagged",
        "answerable": answerable,
        "required_evidence": refs,
        "decoy_evidence": decoys or [],
        "sufficiency_expected": {
            "eventually_sufficient": answerable,
            "expected_acquisition_rounds": rounds,
            "missing_facts_after_primary": missing or [],
            "max_acquisition_rounds": 2,
        },
        "writeback_eligible": answerable and original is None,
        "expected_writeback_action": "skipped_non_empty_cell" if original is not None else ("written" if answerable else "review_only"),
        "original_value": original,
        "has_formula": False,
        "gold_origin": "declared_synthetic_domain_rules",
    }


def document(root: Path, dataset: str, role: str, filename: str) -> dict[str, Any]:
    return {
        "document_key": f"{dataset}.{role}.{filename}",
        "file_name": filename,
        "artifact_relative_path": f"{dataset}/{role}/{filename}",
        "namespace": "global" if role == "global" else dataset,
        "document_role": "knowledge_base",
        "local_path": root / dataset / role / filename,
    }


def xref(doc: dict[str, Any], sheet: str, row: int, values: list[str], fact: str, room: str = "") -> dict[str, Any]:
    return native_ref(
        doc, values, fact_id=fact, room=room, sheet_name=sheet, row_index=row, cell_range=f"A{row}:{get_column_letter(len(values))}{row}"
    )


def pref(doc: dict[str, Any], paragraph: int, text: str, fact: str, room: str = "") -> dict[str, Any]:
    return native_ref(doc, text, fact_id=fact, room=room, paragraph_index=paragraph, source_anchor=f"paragraph {paragraph}")


def tref(doc: dict[str, Any], table: int, row: int, values: list[str], fact: str, room: str = "") -> dict[str, Any]:
    return native_ref(doc, values, fact_id=fact, room=room, table_index=table, row_index=row, source_anchor=f"table {table} row {row}")


def f1(root: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    main = document(root, "f1", "target", "北辰设施.xlsx")
    note = document(root, "f1", "target", "北辰电池实测.docx")
    global_doc = document(root, "f1", "global", "北辰企业标准.xlsx")
    rows = [
        ["字段", "当前实际值"],
        ["市电路数", "2路"],
        ["UPS路数", "2路"],
        ["已投运UPS容量", "500kVA"],
        ["规划UPS容量", "800kVA"],
        ["已投运机架数", "24架"],
        ["规划机架数", "48架"],
        ["机房净面积", "150㎡"],
        ["所在生产楼建筑面积", "1500㎡"],
        ["单机架额定功率", "6kW"],
        ["UPS后备时间", "2小时"],
        ["单台柴油发电机容量", "600kW"],
        ["柴油发电机总容量", "1200kW"],
        ["发电机维护编号", "6001"],
    ]
    other = [["字段", "当前实际值"], ["市电路数", "1路"], ["UPS路数", "1路"], ["已投运UPS容量", "750kVA"], ["机房净面积", "180㎡"]]
    details = [["机房", "设备", "型号", "阶段"], ["北辰301", "柴油发电机", "DG-600A", "已投运"]]
    workbook({"北辰301": rows, "北辰302": other, "设备明细": details}, main["local_path"])
    paragraphs = ["北辰301机房电池实测", "本资料仅记录北辰301机房已投运电池配置。", "北辰301机房当前UPS电池配置为4组。"]
    word(paragraphs, [], note["local_path"])
    global_rows = [["字段", "标准值"], ["企业默认UPS容量", "800kVA"], ["企业默认市电路数", "4路"]]
    workbook({"企业标准": global_rows}, global_doc["local_path"])
    r = {n: xref(main, "北辰301", n, rows[n - 1], f"f1.row{n}", "北辰301") for n in range(2, len(rows) + 1)}
    o = {n: xref(main, "北辰302", n, other[n - 1], f"f1.other{n}", "北辰302") for n in range(2, len(other) + 1)}
    g = xref(global_doc, "企业标准", 2, global_rows[1], "f1.global_ups", "企业默认")
    b = pref(note, 3, paragraphs[2], "f1.battery_groups", "北辰301")
    cases = [
        case(1, "f1", "same_value_wrong_field", "北辰301机房市电路数", "2路", [r[2]], [decoy(r[3], "wrong_field")], rounds=1),
        case(2, "f1", "same_value_wrong_field", "北辰301机房UPS路数", "2路", [r[3]], [decoy(r[2], "wrong_field")], rounds=1),
        case(3, "f1", "same_field_wrong_room", "请填写北辰301机房已投运UPS容量", "500kVA", [r[4]], [decoy(o[4], "wrong_scope")], rounds=1),
        case(4, "f1", "same_field_wrong_room", "北辰301机房净面积是多少", "150㎡", [r[8]], [decoy(o[5], "wrong_scope")], rounds=1),
        case(5, "f1", "current_vs_planned", "北辰301机房已投运机架数", "24架", [r[6]], [decoy(r[7], "planning")], rounds=1),
        case(
            6,
            "f1",
            "numeric_substring",
            "北辰301机房单台柴油发电机容量",
            "600kW",
            [r[12]],
            [decoy(r[13], "wrong_field"), decoy(r[14], "substring")],
            rounds=1,
        ),
        case(7, "f1", "numeric_substring", "北辰301机房净面积而非生产楼总建筑面积", "150㎡", [r[8]], [decoy(r[9], "substring")], rounds=1),
        case(
            8,
            "f1",
            "target_global_conflict",
            "北辰301机房已投运UPS容量以现场记录为准",
            "500kVA",
            [r[4]],
            [decoy(g, "global_conflict")],
            rounds=1,
        ),
        case(
            9,
            "f1",
            "table_detail",
            "北辰301机房已投运柴油发电机型号",
            "DG-600A",
            [xref(main, "设备明细", 2, details[1], "f1.generator_model", "北辰301")],
        ),
        case(10, "f1", "second_round", "北辰301机房UPS后备时间及电池组数", "2小时；4组", [r[11], b], rounds=2, missing=["电池组数"]),
    ]
    return [main, note], [global_doc], cases


def f2(root: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    main = document(root, "f2", "target", "杉岚运行资料.docx")
    facts = document(root, "f2", "target", "杉岚确认清单.xlsx")
    global_doc = document(root, "f2", "global", "杉岚企业标准.xlsx")
    paragraphs = [
        "杉岚201机房运行资料",
        "本资料区分201和202机房，规划记录不能替代已投运记录。",
        "杉岚201机房当前消防分区数为2个。",
        "杉岚201机房当前消防门数为2樘。",
        "杉岚201机房当前UPS套数为3套。",
        "杉岚201机房当前精密空调台数为3台。",
        "杉岚201机房空调配置记录采用当前已投运台数口径。",
        "杉岚201机房规划精密空调台数为5台。",
        "杉岚201机房当前实测PUE为1.42。",
        "杉岚201机房节能改造目标PUE为1.20。",
        "杉岚201机房当前电池后备时间为60分钟。",
        "杉岚202机房当前电池后备时间为90分钟。",
        "杉岚201机房已投运气体灭火介质为七氟丙烷。",
        "杉岚201机房两路市电分别来自杉岚东变电站和杉岚西变电站。",
        "杉岚201机房本月冷量测试记录号为1800。",
    ]
    tables = [
        [
            ["机房", "设备", "单台容量", "数量", "总容量", "品牌"],
            ["杉岚201", "精密空调", "60kW", "3台", "180kW", "岚峰"],
            ["杉岚202", "精密空调", "60kW", "4台", "240kW", "岚峰"],
        ],
        [["机房", "参数", "记录"], ["杉岚201", "市电路数", "2路"]],
    ]
    word(paragraphs, tables, main["local_path"])
    fact_rows = [["字段", "当前实际值"], ["UPS套数", "3套"]]
    workbook({"杉岚201确认": fact_rows}, facts["local_path"])
    global_rows = [["字段", "标准值"], ["企业默认灭火介质", "IG541"], ["企业默认UPS套数", "4套"]]
    workbook({"企业标准": global_rows}, global_doc["local_path"])
    p = {
        n: pref(main, n, paragraphs[n - 1], f"f2.paragraph{n}", "杉岚202" if n == 12 else "杉岚201") for n in range(3, len(paragraphs) + 1)
    }
    t = tref(main, 1, 2, tables[0][1], "f2.cooling_total", "杉岚201")
    alternative = {**p[5], "fact_id": "f2.ups_count"}
    ups = xref(facts, "杉岚201确认", 2, fact_rows[1], "f2.ups_count", "杉岚201")
    g = xref(global_doc, "企业标准", 2, global_rows[1], "f2.global_fire", "企业默认")
    cases = [
        case(11, "f2", "same_value_wrong_field", "杉岚201机房消防分区数", "2个", [p[3]], [decoy(p[4], "wrong_field")], rounds=2),
        case(12, "f2", "same_value_wrong_field", "杉岚201机房UPS套数", "3套", [ups, alternative], [decoy(p[6], "wrong_field")], rounds=1),
        case(
            13,
            "f2",
            "same_field_wrong_room",
            "杉岚201机房精密空调总容量",
            "180kW",
            [t],
            [decoy(tref(main, 1, 3, tables[0][2], "f2.other_cooling", "杉岚202"), "wrong_scope")],
        ),
        case(14, "f2", "same_field_wrong_room", "杉岚201机房电池后备时间", "60分钟", [p[11]], [decoy(p[12], "wrong_scope")], rounds=2),
        case(
            15,
            "f2",
            "current_vs_planned",
            "杉岚201机房已投运精密空调台数",
            "3台",
            [{**p[6], "fact_id": "f2.cooling_count"}, {**t, "fact_id": "f2.cooling_count"}],
            [decoy(p[8], "planning")],
        ),
        case(16, "f2", "current_vs_planned", "杉岚201机房当前实测PUE", "1.42", [p[9]], [decoy(p[10], "planning")], rounds=2),
        case(17, "f2", "numeric_substring", "杉岚201机房精密空调总容量而非测试记录号", "180kW", [t], [decoy(p[15], "substring")]),
        case(
            18,
            "f2",
            "target_global_conflict",
            "杉岚201机房已投运气体灭火介质",
            "七氟丙烷",
            [p[13]],
            [decoy(g, "global_conflict")],
            rounds=2,
        ),
        case(19, "f2", "table_detail", "杉岚201机房精密空调品牌", "岚峰", [{**t, "fact_id": "f2.cooling_brand"}]),
        case(
            20,
            "f2",
            "second_round",
            "杉岚201机房市电路数及两路来源变电站",
            "2路；杉岚东变电站和杉岚西变电站",
            [tref(main, 2, 2, tables[1][1], "f2.mains_count", "杉岚201"), p[14]],
            rounds=2,
            missing=["两路来源变电站"],
        ),
    ]
    return [main, facts], [global_doc], cases


def f3(root: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    main = document(root, "f3", "target", "澄海现网.xlsx")
    note = document(root, "f3", "target", "澄海UPS巡检.docx")
    global_doc = document(root, "f3", "global", "澄海企业标准.xlsx")
    rows = [
        ["字段", "当前实际值"],
        ["门禁控制器数量", "4台"],
        ["电池组数", "4组"],
        ["单机架额定功率", "8kW"],
        ["市电路数", "3路"],
        ["电池后备时间", "30分钟"],
        ["当前外网上联带宽", "10Gbps"],
        ["规划外网上联带宽", "40Gbps"],
        ["当前UPS额定容量", "500kVA"],
        ["UPS最大扩容容量", "1500kVA"],
        ["UPS冗余模式", "N+1"],
        ["柴油发电机额定容量", "600kW"],
    ]
    other = [["字段", "当前实际值"], ["单机架额定功率", "6kW"], ["当前UPS额定容量", "750kVA"]]
    details = [
        ["机房", "设备", "型号", "参数说明"],
        ["澄海南501", "配电断路器", "QF-160", "已投运"],
        ["澄海南501", "机架", "R-42U", "额定承重1200kg"],
    ]
    workbook({"南501": rows, "北502": other, "设备台账": details}, main["local_path"])
    paragraphs = [
        "澄海南501机房UPS巡检",
        "本记录仅列出现场已投运UPS模块配置。",
        "澄海南501机房当前UPS可用模块为3个。",
        "澄海北502机房当前UPS可用模块为4个。",
    ]
    word(paragraphs, [], note["local_path"])
    global_rows = [["字段", "标准值"], ["企业默认市电路数", "2路"], ["企业默认电池后备时间", "90分钟"]]
    workbook({"企业标准": global_rows}, global_doc["local_path"])
    r = {n: xref(main, "南501", n, rows[n - 1], f"f3.row{n}", "澄海南501") for n in range(2, len(rows) + 1)}
    g = {n: xref(global_doc, "企业标准", n, global_rows[n - 1], f"f3.global{n}", "企业默认") for n in (2, 3)}
    cases = [
        case(21, "f3", "same_value_wrong_field", "澄海南501机房门禁控制器数量", "4台", [r[2]], [decoy(r[3], "wrong_field")], rounds=1),
        case(
            22,
            "f3",
            "same_field_wrong_room",
            "澄海南501机房单机架额定功率",
            "8kW",
            [r[4]],
            [decoy(xref(main, "北502", 2, other[1], "f3.other_power", "澄海北502"), "wrong_scope")],
            rounds=1,
        ),
        case(23, "f3", "current_vs_planned", "澄海南501机房当前外网上联带宽", "10Gbps", [r[7]], [decoy(r[8], "planning")], rounds=1),
        case(24, "f3", "numeric_substring", "澄海南501机房当前UPS额定容量", "500kVA", [r[9]], [decoy(r[10], "substring")], rounds=1),
        case(25, "f3", "target_global_conflict", "澄海南501机房市电路数", "3路", [r[5]], [decoy(g[2], "global_conflict")], rounds=1),
        case(26, "f3", "target_global_conflict", "澄海南501机房电池后备时间", "30分钟", [r[6]], [decoy(g[3], "global_conflict")], rounds=1),
        case(
            27,
            "f3",
            "table_detail",
            "澄海南501机房配电断路器型号",
            "QF-160",
            [xref(main, "设备台账", 2, details[1], "f3.breaker_model", "澄海南501")],
        ),
        case(
            28,
            "f3",
            "table_detail",
            "澄海南501机房机架额定承重",
            "1200kg",
            [xref(main, "设备台账", 3, details[2], "f3.rack_load", "澄海南501")],
            original="人工确认1200kg",
        ),
        case(
            29,
            "f3",
            "second_round",
            "澄海南501机房UPS冗余模式及可用模块数",
            "N+1；3个",
            [r[11], pref(note, 3, paragraphs[2], "f3.usable_modules", "澄海南501")],
            rounds=2,
            missing=["可用模块数"],
        ),
        case(
            30,
            "f3",
            "second_round",
            "澄海南501机房柴油发电机额定容量及柴油储量",
            "未找到",
            [r[12]],
            rounds=2,
            missing=["柴油储量"],
            answerable=False,
        ),
    ]
    return [main, note], [global_doc], cases


def form(root: Path, dataset: str, cases: list[dict[str, Any]]) -> dict[str, Any]:
    path = root / dataset / "form" / {"f1": "北辰踏勘表.xlsx", "f2": "杉岚专项调研.xlsx", "f3": "澄海改造核查.xlsx"}[dataset]
    book = Workbook()
    book.remove(book.active)
    book.properties.creator = "Anonymous synthetic evaluation"
    book.properties.created = book.properties.modified = FIXED_DATE
    layouts = {
        "f1": [("采集", 6, 2, 5, list(range(10)))],
        "f2": [("动力消防", 2, 2, 3, list(range(5))), ("运行专项", 8, 4, 6, list(range(5, 10)))],
        "f3": [("现场记录", 3, 4, 2, list(range(5))), ("改造核查", 145, 1, 4, list(range(5, 10)))],
    }
    for name, header, question_col, target_col, selected in layouts[dataset]:
        sheet = book.create_sheet(name)
        category_col = next(i for i in range(1, 7) if i not in (question_col, target_col))
        instruction_col = next(i for i in range(1, 7) if i not in (question_col, target_col, category_col))
        sheet.cell(header, question_col, "字段")
        sheet.cell(header, target_col, "填写结果")
        sheet.cell(header, category_col, "类别")
        sheet.cell(header, instruction_col, "填写说明")
        for offset, selected_index in enumerate(selected, 1):
            item = cases[selected_index]
            row = header + offset
            sheet.cell(row, question_col, item["question"])
            sheet.cell(row, target_col, item["original_value"])
            sheet.cell(row, category_col, "现场参数")
            sheet.cell(row, instruction_col, INSTRUCTION)
            sheet.row_dimensions[row].height = 44
            item["target"] = {"sheet_name": name, "cell": f"{get_column_letter(target_col)}{row}"}
            item["namespace"] = dataset
        if dataset == "f2" and name == "动力消防":
            sheet.merge_cells(start_row=header + 1, end_row=header + len(selected), start_column=category_col, end_column=category_col)
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 58 if col == question_col else (24 if col == target_col else 30)
        for row in sheet.iter_rows(min_row=header, max_row=header + len(selected)):
            for cell in row:
                cell.font = Font(name="Arial", size=11, bold=cell.row == header, color="FFFFFF" if cell.row == header else "000000")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row == header:
                    cell.fill = PatternFill("solid", fgColor="294766")
        sheet.freeze_panes = f"A{header + 1}"
    controls = []
    if dataset == "f3":
        sheet = book["改造核查"]
        sheet["A151"] = "公式保护控制"
        sheet["D151"] = "=SUM(1,2)"
        controls.append(
            {
                "control_id": "F3_FORMULA",
                "target": {"sheet_name": "改造核查", "cell": "D151"},
                "original_value": "=SUM(1,2)",
                "has_formula": True,
                "expected_parser_action": "formula_target",
                "expected_output_value": "=SUM(1,2)",
                "counted_as_challenge": False,
            }
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    book.close()
    fixed_zip(path)
    return {
        "local_path": path,
        "file_name": path.name,
        "artifact_relative_path": f"{dataset}/form/{path.name}",
        "expected_fields": 10,
        "expected_targets": [item["target"] for item in cases],
        "safety_controls": controls,
    }


def validate(pair: dict[str, Any], cases: list[dict[str, Any]]) -> dict[str, Any]:
    from nested_doc_rag.form.template_parser import parse_form_template
    from nested_doc_rag.ingestion import extract_file_chunks

    parsed = parse_form_template(pair["template"]["local_path"])
    assert len(parsed.items) == 10, parsed.report
    assert not parsed.report["ambiguous_rows"] and not parsed.report["errors"], parsed.report
    items_path = pair["template"]["local_path"].parent / "form_items.jsonl"
    items_path.write_text("".join(json.dumps(item, ensure_ascii=False, sort_keys=True) + "\n" for item in parsed.items))
    pair["form_items"] = {
        "local_path": items_path,
        "file_name": items_path.name,
        "artifact_relative_path": f"{pair['id']}/form/form_items.jsonl",
    }
    by_target = {(i["sheet_name"], i["target_cell"].rsplit("!", 1)[1]): i for i in parsed.items}
    for item in cases:
        field = by_target[item["target"]["sheet_name"], item["target"]["cell"]]
        assert field["question_text"] == item["question"]
        assert "heldout_answer" not in field
        assert item["original_value"] is None or item["original_value"] not in str(field)
        item["field_id"] = field["form_item_id"]
    chunks_by_document = {}
    for doc in pair["target_sources"] + pair["global_sources"]:
        chunks_by_document[doc["document_key"]] = list(extract_file_chunks(doc["local_path"], doc["local_path"].parent))
    ref_count = 0
    for item in cases:
        for ref in item["required_evidence"] + item["decoy_evidence"]:
            ref_count += 1
            matches = []
            for chunk in chunks_by_document[ref["document_key"]]:
                if any(
                    ref.get(key) != chunk.get(key)
                    for key in ("sheet_name", "cell_range", "paragraph_index", "table_index", "row_index")
                    if key in ref
                ):
                    continue
                if ref["source_text"] == chunk["raw_source_text"]:
                    matches.append(chunk)
            assert len(matches) == 1, (item["case_id"], ref, matches)
            assert ref["source_text_hash"] == digest(matches[0]["raw_source_text"].encode())
    book = load_workbook(pair["template"]["local_path"], data_only=False)
    for item in cases:
        assert book[item["target"]["sheet_name"]][item["target"]["cell"]].value == item["original_value"]
    book.close()
    assert all(
        control["target"]["cell"] in {row.get("target_cell") for row in parsed.report["skipped_rows"]}
        for control in pair["template"]["safety_controls"]
    )
    return {
        "id": pair["id"],
        "parser": parsed.report,
        "validated_gold_locators": ref_count,
        "declared_answers_source": "independent recipes; parser only checks physical source and target identities",
    }


def public_file(doc: dict[str, Any], declarations: Path) -> dict[str, Any]:
    return {
        **{key: value for key, value in doc.items() if key != "local_path"},
        "path": os.path.relpath(doc["local_path"], declarations),
        "sha256": digest(doc["local_path"].read_bytes()).removeprefix("sha256:"),
    }


def generate(output: Path, declarations: Path, check: bool) -> None:
    all_cases, pairs, validation = [], [], []
    for builder, room in ((f1, "北辰301"), (f2, "杉岚201"), (f3, "澄海南501")):
        target, global_docs, cases = builder(output)
        dataset = cases[0]["dataset_id"]
        pair = {
            "id": dataset,
            "room_context": room,
            "namespaces": {"target": dataset, "global": "global"},
            "collection": COLLECTION,
            "target_sources": target,
            "global_sources": global_docs,
            "template": form(output, dataset, cases),
            "gold": {"path": "gold.jsonl", "dataset_id": dataset, "challenge_count": 10},
        }
        validation.append(validate(pair, cases))
        pairs.append(
            {
                **pair,
                "target_sources": [public_file(d, declarations) for d in target],
                "global_sources": [public_file(d, declarations) for d in global_docs],
                "template": public_file(pair["template"], declarations),
                "form_items": public_file(pair["form_items"], declarations),
            }
        )
        all_cases.extend(cases)
    counts = dict(Counter(item["category"] for item in all_cases))
    assert counts == CATEGORIES, counts
    assert len({item["case_id"] for item in all_cases}) == 30
    assert len({item["field_id"] for item in all_cases}) == 30
    manifest = {
        "schema_version": SCHEMA,
        "dataset_version": "synthetic-native-1",
        "origin": "anonymous_synthetic",
        "challenge_count": 30,
        "category_counts": counts,
        "independent_pairs": 3,
        "pairs": pairs,
        "gold_origin": "declared recipes with explicit domain rules; no model or run artifact used",
        "path_base": "manifest directory",
        "namespace_rule": "each pair requires a separate workspace and independent target/global KBs",
        "generation": {
            "command": "python scripts/vnext_generate_datasets.py",
            "deterministic_ooxml_zip_time": "2026-01-01T00:00:00",
            "dependencies": ["production openpyxl", "production python-docx"],
            "dependency_versions": {"openpyxl": version("openpyxl"), "python-docx": version("python-docx")},
            "external_model_calls": 0,
        },
    }
    gold_data = b"".join((json.dumps(item, ensure_ascii=False, sort_keys=True) + "\n").encode() for item in all_cases)
    report = {
        "schema_version": SCHEMA,
        "pairs": validation,
        "challenge_count": 30,
        "category_counts": counts,
        "external_model_calls": 0,
        "qdrant_calls": 0,
        "gold_sha256": digest(gold_data),
        "check_only": check,
        "validation_scope": "native source parsing and submitted-form parsing; no retrieval/model/E2E acceptance",
    }
    output.mkdir(parents=True, exist_ok=True)
    (output / "validation.json").write_bytes(json_bytes(report))
    if check:
        expected = json.loads((declarations / "manifest.json").read_text())
        for actual_pair, expected_pair in zip(pairs, expected["pairs"], strict=True):
            for role in ("target_sources", "global_sources"):
                assert [(d["document_key"], d["sha256"]) for d in actual_pair[role]] == [
                    (d["document_key"], d["sha256"]) for d in expected_pair[role]
                ]
            assert actual_pair["template"]["sha256"] == expected_pair["template"]["sha256"]
            assert actual_pair["form_items"]["sha256"] == expected_pair["form_items"]["sha256"]
        assert (declarations / "gold.jsonl").read_bytes() == gold_data, "gold declarations drifted"
    else:
        declarations.mkdir(parents=True, exist_ok=True)
        (declarations / "manifest.json").write_bytes(json_bytes(manifest))
        (declarations / "gold.jsonl").write_bytes(gold_data)
    print(
        json.dumps(
            {
                "pairs": 3,
                "challenge_cases": 30,
                "category_counts": counts,
                "native_locators_validated": sum(r["validated_gold_locators"] for r in validation),
                "parser_fields": [r["parser"]["detected_fields"] for r in validation],
                "check_only": check,
                "external_model_calls": 0,
            },
            ensure_ascii=False,
        )
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-root", type=Path, default=PROJECT / "artifacts/vnext/phase5/datasets")
    parser.add_argument("--declarations-dir", type=Path, default=PROJECT / "docs/vnext/datasets")
    parser.add_argument("--check", action="store_true", help="regenerate and verify checked-in hashes/gold without changing declarations")
    args = parser.parse_args()
    generate(args.output_root.resolve(), args.declarations_dir.resolve(), args.check)


if __name__ == "__main__":
    main()
